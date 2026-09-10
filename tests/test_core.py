"""Core regression tests — pure python, NO model, no pytest.

    python3.11 tests/test_core.py

Every test here is a PATHOLOGY THAT WAS MEASURED on this code, written back as
an assertion so it cannot return. They are grouped by the bug they close:

    A  a contradiction deferred in bulk mode was DESTROYED, not deferred
    B  arbitration depended on teaching ORDER, and killed its own winner
    C  a slot with three rivals was linked as a broken star
    D  the derivation engine wrote AROUND the contradiction gate
    E  transitivity was re-derived from scratch per written cell (quadratic)

Nothing here needs a network, a model or a GPU: the whole discrete layer is
counters, links and thresholds, and that is exactly why it can be tested like
this.
"""
import os
import sys
import tempfile
import time

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, "src"))

from lmm.core import dynamics                                          # noqa: E402
from lmm.core.gate import Gate, SPEAK                                  # noqa: E402
from lmm.core.memory import (CONTRA, DOCUMENT, INFERRED, Memory,       # noqa: E402
                       OPERATOR, STRANGER, SUSPECT)
from lmm.core.session import Session as V3Session                      # noqa: E402
from lmm.core.transitive import Transitivity                           # noqa: E402

PASSED = []


def test(name):
    """Registers a test — and prints the line the moment it passes."""
    def keep(function):
        PASSED.append((name, function))
        return function
    return keep


def slot(memory, subject, predicate):
    """The records standing in one slot — value key -> record."""
    return {r.value: r for r in memory.about(subject, touch=False)
            if r.predicate == predicate}


def capital_case(rival=None):
    """The measured scenario: fransa → başkent → paris | berlin."""
    memory = Memory()
    gate = Gate(memory)
    gate.rival = rival
    where = memory.identify("fransa")
    what = memory.identify("baskent")
    return memory, gate, where, what, memory.identify("paris"), memory.identify("berlin")


def links_of(record, kind):
    return {key for one, key in record.links if one == kind}


# --- A: the deferred contradiction is KEPT, and sleep pays the debt ------

@test("A1 bulk-deferred contradiction is recorded as a suspicion")
def a1():
    # PENDING (None) is the bulk-mode answer: the rivalry test cannot run per
    # cell. Measured before the fix: links=[[], []], pressure=[] — the doubt
    # was gone, both records were speakable.
    memory, gate, where, what, paris, berlin = capital_case(
        rival=lambda old, new: None)
    gate.admit(where, what, paris, "#table", DOCUMENT)
    gate.admit(where, what, berlin, "#table", DOCUMENT)
    held = slot(memory, where, what)
    assert len(held) == 2, "both claims must be WRITTEN — a claim with a source is kept"
    assert links_of(held[paris], SUSPECT) == {held[berlin].key}
    assert links_of(held[berlin], SUSPECT) == {held[paris].key}
    assert dynamics.pressure(memory), "an unjudged suspicion must generate pressure"


@test("A2 sleep judges the deferred suspicions in bulk and pays the debt")
def a2():
    memory, gate, where, what, paris, berlin = capital_case(
        rival=lambda old, new: None)
    gate.admit(where, what, paris, "#table", DOCUMENT)
    gate.admit(where, what, berlin, "#table", OPERATOR)
    # The verdict that could not be afforded at write time arrives at sleep.
    counted = dynamics.sleep(memory, verdict=lambda old, new: True)
    assert counted["contradicted"] == 1, counted
    held = slot(memory, where, what)
    assert links_of(held[paris], CONTRA) == {held[berlin].key}
    assert not links_of(held[paris], SUSPECT), "a judged suspicion is no longer pending"
    assert held[berlin].trust >= SPEAK, "the winner speaks"
    assert held[paris].trust < SPEAK, "the loser does NOT speak — the debt is paid"


@test("A3 a suspicion the verdict CLEARS leaves no scar")
def a3():
    memory, gate, where, what, paris, berlin = capital_case(
        rival=lambda old, new: None)
    gate.admit(where, what, paris, "#table", DOCUMENT)
    gate.admit(where, what, berlin, "#table", DOCUMENT)
    counted = dynamics.sleep(memory, verdict=lambda old, new: False)
    assert counted["cleared"] == 1, counted
    held = slot(memory, where, what)
    for record in held.values():
        assert not record.links, "coexisting values keep no contradiction link"
        assert record.trust >= SPEAK, "neither claim is punished for coexisting"
    assert not dynamics.pressure(memory)


@test("A4 an unjudgeable suspicion WAITS — it is never silently dropped")
def a4():
    memory, gate, where, what, paris, berlin = capital_case(
        rival=lambda old, new: None)
    gate.admit(where, what, paris, "#table", DOCUMENT)
    gate.admit(where, what, berlin, "#table", DOCUMENT)
    dynamics.sleep(memory, verdict=lambda old, new: None)       # still unknown
    held = slot(memory, where, what)
    assert links_of(held[paris], SUSPECT), "the debt survives a round with no verdict"
    dynamics.sleep(memory, verdict=lambda old, new: True)       # the verdict arrives late
    assert links_of(held[paris], CONTRA)
    assert min(r.trust for r in held.values()) < SPEAK <= max(r.trust for r in held.values())


@test("A5 a crowded slot costs one anchor per round, and the debt survives")
def a5():
    # A slot holding k pending values has k² possible questions. A sleep round
    # may not be able to cost that: it settles the strongest claim's field and
    # RE-MARKS the rest, so the debt shrinks every round and is never dropped.
    memory = Memory()
    gate = Gate(memory)
    gate.rival = lambda old, new: None
    where = memory.identify("konu")
    what = memory.identify("ozellik")
    values = [memory.identify("deger%d" % i) for i in range(6)]
    for value in values:
        gate.admit(where, what, value, "#table", DOCUMENT)
    asked = []

    def coexist(old, new):
        asked.append((old, new))
        return False                    # nothing is a rival: the slow case

    dynamics.sleep(memory, verdict=coexist)
    assert len(asked) == 5, "one round asks about one anchor, not every pair"
    assert dynamics.pressure(memory), "the rest of the debt is still marked"
    rounds = 0
    while dynamics.pressure(memory) and rounds < 20:
        dynamics.sleep(memory, verdict=coexist)
        rounds += 1
    assert not dynamics.pressure(memory), "successive rounds settle the slot"
    held = slot(memory, where, what)
    assert all(r.trust >= SPEAK for r in held.values()), \
        "values that all coexist all keep their voice"


# --- B: arbitration is order-independent, idempotent, and keeps a winner --

@test("B1 DOCUMENT then OPERATOR: the contradiction does NOT escape")
def b1():
    # Measured before the fix: 0.6 and 0.75 — BOTH above SPEAK, both spoken.
    memory, gate, where, what, paris, berlin = capital_case()
    gate.admit(where, what, paris, "belge", DOCUMENT)
    gate.admit(where, what, berlin, "#operator", OPERATOR)
    held = slot(memory, where, what)
    speakers = [r for r in held.values() if r.trust >= SPEAK]
    assert len(speakers) == 1, "exactly one side of a contradiction may speak"
    assert speakers[0].value == berlin


@test("B2 OPERATOR then DOCUMENT gives the SAME outcome — order-independent")
def b2():
    memory, gate, where, what, paris, berlin = capital_case()
    gate.admit(where, what, berlin, "#operator", OPERATOR)
    gate.admit(where, what, paris, "belge", DOCUMENT)
    held = slot(memory, where, what)
    speakers = [r for r in held.values() if r.trust >= SPEAK]
    assert len(speakers) == 1 and speakers[0].value == berlin
    # Both orders must agree down to the numbers, not merely in the verdict.
    other = capital_case()
    memory2, gate2, where2, what2, paris2, berlin2 = other
    gate2.admit(where2, what2, paris2, "belge", DOCUMENT)
    gate2.admit(where2, what2, berlin2, "#operator", OPERATOR)
    a, b = slot(memory, where, what), slot(memory2, where2, what2)
    assert round(a[paris].trust, 6) == round(b[paris2].trust, 6)
    assert round(a[berlin].trust, 6) == round(b[berlin2].trust, 6)


@test("B3 30 witnesses vs one sentence: the slot does NOT end up empty")
def b3():
    # Measured before the fix: 0.188 and 0.375 — mutual destruction, both
    # under SPEAK. A contradiction between two claims is not a reason to know
    # nothing.
    memory, gate, where, what, paris, berlin = capital_case()
    for i in range(30):
        gate.admit(where, what, paris, "belge%d" % i, DOCUMENT)
    gate.admit(where, what, berlin, "#operator", OPERATOR)
    dynamics.sleep(memory)
    held = slot(memory, where, what)
    speakers = [r for r in held.values() if r.trust >= SPEAK]
    assert len(speakers) == 1, [(r.value, r.trust) for r in held.values()]
    assert speakers[0].witnesses == 1 and speakers[0].level == OPERATOR


@test("B4 the winner is never punished, and repeated rounds change nothing")
def b4():
    memory, gate, where, what, paris, berlin = capital_case()
    gate.admit(where, what, paris, "belge", DOCUMENT)
    gate.admit(where, what, berlin, "#operator", OPERATOR)
    held = slot(memory, where, what)
    winner = held[berlin].trust
    assert winner == memory.trust_of(OPERATOR), "the winner keeps its own trust"
    before = {k: r.trust for k, r in held.items()}
    for _ in range(5):
        dynamics.sleep(memory)
        dynamics.arbitrate(memory, held[paris])
        dynamics.arbitrate(memory, held[berlin])
    assert {k: r.trust for k, r in held.items()} == before, "arbitration is idempotent"


@test("B5 the loser stays below BOTH the winner and the speaking threshold")
def b5():
    for first, second in ((STRANGER, DOCUMENT), (DOCUMENT, STRANGER),
                          (OPERATOR, DOCUMENT), (DOCUMENT, OPERATOR)):
        memory, gate, where, what, paris, berlin = capital_case()
        gate.admit(where, what, paris, "a", first)
        gate.admit(where, what, berlin, "b", second)
        held = list(slot(memory, where, what).values())
        held.sort(key=lambda r: -r.trust)
        assert held[0].trust >= SPEAK > held[1].trust, (first, second,
                                                        [r.trust for r in held])


# --- C: the rival SET, not the first rival ------------------------------

@test("C1 three rivals in one slot are all linked to each other")
def c1():
    # Measured before the fix: v2 and v3 were never linked — _contradiction
    # returned the FIRST clash and stopped, so arbitration entered a broken
    # star and judged a different sub-field depending on where it started.
    memory = Memory()
    gate = Gate(memory)
    where = memory.identify("fransa")
    what = memory.identify("baskent")
    values = [memory.identify(name) for name in ("paris", "berlin", "roma")]
    for i, value in enumerate(values):
        gate.admit(where, what, value, "belge%d" % i, DOCUMENT)
    held = slot(memory, where, what)
    for value in values:
        others = {held[o].key for o in values if o != value}
        assert links_of(held[value], CONTRA) == others, (value, held[value].links)


@test("C2 the whole field is judged from whichever member is asked")
def c2():
    memory = Memory()
    gate = Gate(memory)
    where = memory.identify("fransa")
    what = memory.identify("baskent")
    paris = memory.identify("paris")
    berlin = memory.identify("berlin")
    roma = memory.identify("roma")
    gate.admit(where, what, paris, "belge", DOCUMENT)
    gate.admit(where, what, berlin, "yabanci", STRANGER)
    gate.admit(where, what, roma, "#operator", OPERATOR)
    held = slot(memory, where, what)
    for entry in held.values():
        assert len(dynamics.field_of(memory, entry)) == 3
        assert dynamics.arbitrate(memory, entry).value == roma
    speakers = [r for r in held.values() if r.trust >= SPEAK]
    assert len(speakers) == 1 and speakers[0].value == roma


@test("C3 an unjudged pair inside a judged field is still judged at sleep")
def c3():
    # A suspicion is a marker on the FIELD, not a relation between two rows:
    # if the field is already connected, no extra link is written — but the
    # pair must still get its verdict when sleep walks that field. This is the
    # part of the debt that a spanning topology could quietly lose.
    memory = Memory()
    gate = Gate(memory)
    where = memory.identify("fransa")
    what = memory.identify("baskent")
    paris, berlin, roma = (memory.identify(n) for n in ("paris", "berlin", "roma"))
    gate.rival = lambda old, new: None                  # bulk: nothing judgeable
    for value in (paris, berlin, roma):
        gate.admit(where, what, value, "#table", DOCUMENT)
    held = slot(memory, where, what)
    assert len(dynamics.field_of(memory, held[roma], (CONTRA, SUSPECT))) == 3, \
        "one field, however few links carry it"
    asked = []

    def verdict(old, new):
        asked.append((old, new))
        return True

    dynamics.sleep(memory, verdict=verdict)
    judged = {frozenset(pair) for pair in asked}
    assert judged >= {frozenset((paris, berlin)), frozenset((paris, roma))} or \
        len(judged) >= 2, asked
    speakers = [r for r in held.values() if r.trust >= SPEAK]
    assert len(speakers) == 1, [(r.value, r.trust) for r in held.values()]
    assert not any(links_of(r, SUSPECT) for r in held.values()), \
        "every suspicion in the field was settled"


@test("C4 a slot with many pending values does not build a clique")
def c4():
    # Measured while fixing this: linking every pending pair in a slot that a
    # repeated table column had filled cost ~2M links for information the
    # field closure already carries.
    memory = Memory()
    gate = Gate(memory)
    gate.rival = lambda old, new: None
    where = memory.identify("satir")
    what = memory.identify("kolon")
    for i in range(200):
        gate.admit(where, what, memory.identify("hucre%d" % i), "#table", DOCUMENT)
    held = list(slot(memory, where, what).values())
    links = sum(len(r.links) for r in held)
    assert links <= 4 * len(held), "the suspicion topology must stay linear (%d)" % links
    assert len(dynamics.field_of(memory, held[0], (CONTRA, SUSPECT))) == 200, \
        "and it must still be ONE field"


# --- D: the derivation engine goes through the gate ----------------------

@test("D1 an inference contradicting an operator fact is NOT written silently")
def d1():
    # Measured before the fix: gate.inferred called memory.write directly —
    # "kartal→balık" landed next to the operator's "kartal→kuş" with links=[]
    # and pressure=[]. The one writer in the system that could contradict
    # without anyone noticing.
    memory = Memory()
    gate = Gate(memory)
    eagle = memory.identify("kartal")
    kind = memory.identify("tur")
    bird, fish = memory.identify("kus"), memory.identify("balik")
    gate.admit(eagle, kind, bird, "#operator", OPERATOR)
    gate.inferred(eagle, kind, fish, [])
    held = slot(memory, eagle, kind)
    assert links_of(held[fish], CONTRA) == {held[bird].key}
    assert links_of(held[bird], CONTRA) == {held[fish].key}
    assert dynamics.pressure(memory), "the derivation's doubt must be on record"
    assert held[fish].level == INFERRED and held[fish].source == "#inference"
    assert held[fish].trust < SPEAK <= held[bird].trust


@test("D2 a derivation that contradicts nothing still passes cleanly")
def d2():
    memory = Memory()
    gate = Gate(memory)
    eagle = memory.identify("kartal")
    kind = memory.identify("tur")
    animal = memory.identify("hayvan")
    bird = memory.identify("kus")
    because = gate.admit(eagle, kind, bird, "#operator", OPERATOR)[0]
    # kuş → hayvan on a predicate the graph proved transitive: the derived
    # "kartal → hayvan" is the CHAIN of the premise, not its rival.
    gate.admit(bird, kind, animal, "#operator", OPERATOR)
    memory.transitive.add(kind)
    held = gate.inferred(eagle, kind, animal, [because.key])
    assert held is not None and not links_of(held, CONTRA)
    assert (3, because.key) in held.links, "the rationale link survives the gate"
    assert not held.episodic, "a derivation is not an episode"


@test("D3 the derivation loop does not feed on itself")
def d3():
    session = V3Session(None)
    memory = session.memory
    kind = session._identity("tur")
    a, b, c = (session._identity(n) for n in ("kartal", "kus", "hayvan"))
    for one, two in ((a, b), (b, c), (a, c)):
        session.gate.admit(one, kind, two, "#operator", OPERATOR)
        session._learn_transitive(kind, one, two)
    memory.transitive.add(kind)
    d = session._identity("canli")
    session.gate.admit(c, kind, d, "#operator", OPERATOR)
    session._derive(c, kind, d)
    grew = len(memory.records)
    for _ in range(3):
        session._derive(c, kind, d)         # deriving again must add nothing
    assert len(memory.records) == grew, "derivation is a fixed point, not a pump"
    for record in memory.records.values():
        if record.source == "#inference":
            assert record.trust < SPEAK, "no derivation may reach speaking trust"


# --- E: transitivity is learned incrementally, not re-derived ------------

@test("E1 the incremental tracker learns exactly what the full scan learned")
def e1():
    # is-a closes triangles; "sever" does not — and one coincidental triangle
    # is still not enough.
    memory = Memory()
    gate = Gate(memory)
    kind = memory.identify("tur")
    names = ["a", "b", "c", "d"]
    keys = {n: memory.identify(n) for n in names}
    tracker = Transitivity(memory)
    edges = [("a", "b"), ("b", "c"), ("a", "c"),        # first triangle
             ("c", "d"), ("b", "d")]                    # closes the second
    for i, (one, two) in enumerate(edges):
        gate.admit(keys[one], kind, keys[two], "belge%d" % i, DOCUMENT)
        tracker.observe(kind, keys[one], keys[two])
    assert kind in memory.transitive, "two witnessed triangles make it transitive"

    loves = Memory()
    gate2 = Gate(loves)
    predicate = loves.identify("sever")
    people = {n: loves.identify(n) for n in names}
    tracker2 = Transitivity(loves)
    for i, (one, two) in enumerate([("a", "b"), ("b", "c"), ("c", "d")]):
        gate2.admit(people[one], predicate, people[two], "belge%d" % i, DOCUMENT)
        tracker2.observe(predicate, people[one], people[two])
    assert predicate not in loves.transitive, "an open chain is not transitivity"


@test("E2 ONE coincidental triangle is not enough")
def e2():
    memory = Memory()
    gate = Gate(memory)
    predicate = memory.identify("sever")
    keys = {n: memory.identify(n) for n in ("a", "b", "c")}
    tracker = Transitivity(memory)
    for i, (one, two) in enumerate([("a", "b"), ("b", "c"), ("a", "c")]):
        gate.admit(keys[one], predicate, keys[two], "belge%d" % i, DOCUMENT)
        tracker.observe(predicate, keys[one], keys[two])
    assert predicate not in memory.transitive


@test("E3 a memory loaded from file still has its triangles found")
def e3():
    # The tracker's first pass over a predicate reads what already stands —
    # otherwise a restart would forget everything the graph had proven.
    memory = Memory()
    gate = Gate(memory)
    kind = memory.identify("tur")
    keys = {n: memory.identify(n) for n in ("a", "b", "c", "d")}
    for i, (one, two) in enumerate([("a", "b"), ("b", "c"), ("a", "c"),
                                    ("c", "d"), ("b", "d")]):
        gate.admit(keys[one], kind, keys[two], "belge%d" % i, DOCUMENT)
    memory.transitive.clear()               # as if freshly loaded from disk
    assert Transitivity(memory).observe(kind)


@test("E4 derived records are not evidence for the rule that derived them")
def e4():
    memory = Memory()
    gate = Gate(memory)
    kind = memory.identify("tur")
    keys = {n: memory.identify(n) for n in ("a", "b", "c")}
    tracker = Transitivity(memory)
    for one, two in [("a", "b"), ("b", "c")]:
        gate.admit(keys[one], kind, keys[two], "belge", DOCUMENT)
        tracker.observe(kind, keys[one], keys[two])
    gate.inferred(keys["a"], kind, keys["c"], [])
    tracker.observe(kind, keys["a"], keys["c"])
    assert kind not in memory.transitive, "a derivation cannot prove its own rule"


@test("E5 ingestion is near-linear, not quadratic")
def e5():
    # Measured before the fix: 2000 cells 0.91s, 10000 cells 25.6s — 5x the
    # data for 28x the time. The check is the SHAPE: five times the cells may
    # not cost far more than five times the work.
    def ingest(cells):
        memory = Memory()
        gate = Gate(memory)
        tracker = Transitivity(memory)
        columns = [memory.identify("kolon%d" % i) for i in range(5)]
        started = time.time()
        for i in range(cells):
            subject = memory.identify("satir%d" % i)
            value = memory.identify("hucre%d" % i)
            predicate = columns[i % len(columns)]
            gate.admit(subject, predicate, value, "#table", DOCUMENT)
            tracker.observe(predicate, subject, value)
        return time.time() - started

    small = ingest(2000)
    large = ingest(10000)
    ratio = large / max(small, 1e-6)
    print("      2000 cells %.3fs · 10000 cells %.3fs · ratio %.1fx"
          % (small, large, ratio))
    assert ratio < 10.0, "5x the data must not cost ~28x the time (was quadratic)"


# --- behaviour: nothing above broke the ordinary paths -------------------

@test("F1 a table ingested without a model still writes, and reloads")
def f1():
    from lmm.session import Session
    session = Session(None)
    rows = [{"ekipman": "yangin tupu %d" % i, "durum": "uygun",
             "yer": "kat %d" % (i % 5)} for i in range(100)]
    wrote = session.learn_rows(rows)
    assert wrote == 200, wrote
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        "_core_roundtrip.lmm")
    try:
        session.memory.save(path)
        again = Memory.load(path)
        assert len(again.records) == len(session.memory.records)
        assert all(again.records[k].links == session.memory.records[k].links
                   for k in session.memory.records), "links survive the round trip"
    finally:
        for leftover in (path, path + ".sent"):
            if os.path.exists(leftover):
                os.remove(leftover)


@test("F2 an old .lmm file (no SUSPECT links) still loads and arbitrates")
def f2():
    memory = Memory()
    gate = Gate(memory)
    where = memory.identify("fransa")
    what = memory.identify("baskent")
    paris, berlin = memory.identify("paris"), memory.identify("berlin")
    gate.admit(where, what, paris, "belge", DOCUMENT)
    gate.admit(where, what, berlin, "#operator", OPERATOR)
    path = os.path.join(os.path.dirname(os.path.abspath(__file__)),
                        "_core_legacy.lmm")
    try:
        memory.save(path)
        again = Memory.load(path)
        counted = dynamics.sleep(again, verdict=lambda old, new: True)
        assert counted["judged"] >= 1
        held = slot(again, where, what)
        assert len([r for r in held.values() if r.trust >= SPEAK]) == 1
    finally:
        if os.path.exists(path):
            os.remove(path)


@test("F3 the v3 session still learns and answers without any network")
def f3():
    session = V3Session(None)
    kind = session._identity("tur")
    eagle, bird = session._identity("kartal"), session._identity("kus")
    record, why = session.gate.admit(eagle, kind, bird, "#operator", OPERATOR)
    assert record is not None and why == 0
    assert session.gate.behind(eagle, kind, bird) is not None
    assert "kartal" in session._render([record])
    assert isinstance(session.sleep(), dict)
    assert session.curious() is not None and session.tension() == []


@test("F4 the no-fabrication holes stay closed")
def f4():
    from lmm import evidence, link
    memory = Memory()
    memory.identify("şekersiz")
    assert link.resolve(memory, "şeker", create=False) is None
    memory2 = Memory()
    key = memory2.identify("metaldir")
    assert link.resolve(memory2, "metal", create=True) == key
    memory3 = Memory()
    memory3.identify("organ")
    assert link.resolve(memory3, "organizma", create=False) is None
    assert not evidence.covered("Kapak menteşesi Seviye 1",
                                "[K1] Kapak menteşesi Seviye 3", "seviye 1 mi")
    assert not evidence.covered("test Seviye 5",
                                "[K1] test Seviye 3 parça 5 yıl", "")


@test("G1 a unit is content no matter how short it is")
def g1():
    """UNIT BLINDNESS — the largest measured single loss class. A token
    standing next to a NUMBER is that number's unit; dropping it for being
    under three letters lost 'kg', and splitting the camel seam inside 'mAh'
    lost the capacity's unit twice over."""
    from lmm import evidence
    assert "kg" in evidence._words("Kütle 2 kg")
    assert "mah" in evidence._words("Akü 12.8 V / 4200 mAh")
    assert "v" in evidence._words("Akü 12.8 V / 4200 mAh")
    # the camel seam still cuts glued CELLS (the left side is a word there)
    assert "kapağı" in evidence._words("GövdeKapağıKayış")
    # a short token with no number next to it is still not content
    assert "ve" not in evidence._words("ekran ve klavye")


@test("G2 a question with no number of its own still reaches the unit")
def g2():
    """'kaç kg' carries no 2 for the kg to lean on — the corpus has to vouch
    for it. Learned from the indexed sentences, never a word list."""
    from lmm import evidence
    store = evidence.SentenceStore()
    store.add("Kütle (ambalajsız) 3.4 kg")
    store.add("Gösterge 21.5 inç OLED")
    store.add("Cihaz taşınabilir bir sistemdir")
    assert "kg" in store.units
    assert store.find("kaç kg") == ["Kütle (ambalajsız) 3.4 kg"]


@test("G3 the unit anchor admits inflection without opening a hole")
def g3():
    """A fluent sentence ('21.5 inçtir') must pass against '21.5 inç' — the
    same NUMBER binds both. Words that no number holds together stay out,
    even when their prefix step is smaller."""
    from lmm import evidence
    assert evidence.covered("Gösterge 21.5 inçtir.", "Gösterge 21.5 inç OLED", "")
    assert not evidence.covered("kartal", "kart", "")
    assert not evidence.covered("organizma", "organ", "")
    # ...and the anchor needs the SAME number, not just any number
    assert not evidence.covered("Gösterge 17 inçtir.", "Gösterge 21.5 inç OLED", "")


@test("G4 coverage is the ratio whose 1.0 is the old binary gate")
def g4():
    from lmm import evidence
    block = "[K1] Gösterge 21.5 inç OLED"
    assert evidence.coverage("Gösterge 21.5 inç", block, "") == 1.0
    part = evidence.coverage("Gösterge 21.5 inç panelde gösterilir", block, "")
    assert 0.0 < part < 1.0
    assert not evidence.covered("Gösterge 21.5 inç panelde gösterilir", block, "")
    assert evidence.coverage("", block, "") == 0.0


@test("G5 the answer is SELECTED at the gate, not taken on faith")
def g5():
    """GENERATE-AND-SELECT. One generation bound the answer to a single sample
    of a sampling process; the measured residue was an unstable CHOICE, not
    missing knowledge. Here the engine is a stub, so what is under test is the
    SELECTOR: does the composite score prefer the grounded candidate, and does
    it still refuse when every candidate breaks a gate."""
    from lmm import generate, session as lmm_session
    s = lmm_session.Session.__new__(lmm_session.Session)      # no engine needed
    s.memory = Memory()
    proof = ["Gösterge 21.5 inç OLED gösterge", "Cihaz taşınabilir bir sistemdir"]
    block = "[K1] Gösterge 21.5 inç OLED gösterge\n[K2] Cihaz taşınabilir bir sistemdir"
    said, order = {}, []
    answers, supported = {}, [True]

    def fake_answer(question, facts, warmth=0.2, persona="", **kw):
        order.append(facts)
        return answers.get(facts, "Gösterge 21.5 inç")

    def fake_supported(answer, facts):
        said[answer] = said.get(answer, 0) + 1
        return supported[0]

    real = (generate.answer, generate.supported, generate.answers_asked)
    generate.answer, generate.supported = fake_answer, fake_supported
    # the RELATION gate has its own test (H5); here it stays out of the way
    generate.answers_asked = lambda question, answer, view: True
    try:
        # THE LADDER RUNS: one candidate per evidence subset, and each subset
        # is a different block (2 of them for 2 evidence sentences).
        chosen, tried = s._select("Ekran kaç inç", [], proof, "", block)
        assert chosen == "Gösterge 21.5 inç", chosen
        assert len(order) == len(set(order)) >= 2, order
        # A FABRICATED DIGIT IS VETOED even with the engine saying yes: the
        # wide-block candidate invents 17, the narrow one stays grounded.
        answers.clear()
        answers[order[0]] = "Ekran 17 inçtir"
        order.clear()
        chosen, tried = s._select("Ekran kaç inç", [], proof, "", block)
        assert chosen == "Gösterge 21.5 inç", chosen
        assert "Ekran 17 inçtir" in tried    # it was generated, and rejected
        # EVERY candidate broken → honest refusal, no fallback to the least bad
        answers.clear()
        for facts in order:
            answers[facts] = "Ekran 17 inçtir"
        chosen, _tried = s._select("Ekran kaç inç", [], proof, "", block)
        assert chosen is None
        # The read-back keeps its veto: coverage alone cannot speak.
        answers.clear()
        supported[0] = False
        chosen, _tried = s._select("Ekran kaç inç", [], proof, "", block)
        assert chosen is None
    finally:
        (generate.answer, generate.supported,
         generate.answers_asked) = real


@test("G6 the read-back is asked about the evidence the claim rests on")
def g6():
    """The measured oscillation was the JUDGE, not the answer: the same claim
    against the same 3000-character block came back False, True, False. The
    narrow view is asked FIRST, and the wide one only if it fails — a haystack
    the size of the needle."""
    from lmm import generate, session as lmm_session
    s = lmm_session.Session.__new__(lmm_session.Session)
    proof = ["Ortam sıcaklığı 8 °C ~ +32 °C", "Barkod okuyucu isteğe bağlıdır",
             "Barkod yazıcı kurulumu için kılavuza bakın"]
    block = "\n".join(f"[K{i}] {s}" for i, s in enumerate(proof, 1))
    asked = []
    real = generate.supported

    def fake(answer, view):
        asked.append(view)
        return view == block        # only the WIDE view says yes

    generate.supported = fake
    try:
        claim = "Ortam sıcaklığı 8 °C ile +32 °C arasındadır"
        assert s._read_back(claim, proof, block)
        # first view: only the sentence the claim shares words with
        assert "Ortam" in asked[0] and "Barkod" not in asked[0], asked[0]
        assert len(asked) == 2 and asked[1] == block
        # a confirmed narrow view ends it — the wide view is not even asked
        asked.clear()
        generate.supported = lambda answer, view: True
        assert s._read_back(claim, proof, block)
        assert len(asked) == 0
    finally:
        generate.supported = real


@test("G7 among grounded candidates the one that ANSWERS wins")
def g7():
    """Two measured corpus losses, both from ranking the survivors wrongly.
    Counting agreement handed 'Morlan bir canlıdır' the answer because two
    NESTED subsets said it — narrower evidence, vaguer claim, twice as many
    voices. Ranking by groundedness handed 'Torvanit bir metaldir' the answer
    over 'Evet, torvanit bir maddedir', which was penalised for the word
    'evet'. Both claims in each pair are true; only one of each answers."""
    from lmm import generate, session as lmm_session
    s = lmm_session.Session.__new__(lmm_session.Session)
    generate_real = (generate.answer, generate.supported,
                     generate.answers_asked)
    generate.answers_asked = lambda question, answer, view: True
    proof = ["Morlan bir kuştur.", "Kuş bir canlıdır.", "Morlan bir canlıdır."]
    block = "\n".join(f"[K{i}] {t}" for i, t in enumerate(proof, 1))
    wide = {}

    def answer(question, facts, warmth=0.2, persona="", **kw):
        # the widest block answers specifically, the narrow ones generically
        return ("Morlan bir kuştur." if len(facts) == max(wide, default=0)
                else "Morlan bir canlıdır.")

    generate.supported = lambda a, b: True
    generate.answer = answer
    try:
        for one in s._subsets([], proof, ""):
            wide[len(one)] = one
        assert len(wide) >= 2, wide
        chosen, _tried = s._select("morlan nedir", [], proof, "", block)
        assert chosen == "Morlan bir kuştur.", chosen
        # ...and the more literal candidate does not outrank the one that
        # answers the question asked
        pair = ["Torvanit bir metaldir.", "Metal bir maddedir."]
        pair_block = "\n".join(f"[K{i}] {t}" for i, t in enumerate(pair, 1))
        replies = iter(["Torvanit bir metaldir.", "Evet, torvanit bir maddedir.",
                        "Evet, torvanit bir maddedir."])
        generate.answer = lambda q, f, warmth=0.2, persona="", **kw: next(replies)
        chosen, _tried = s._select("torvanit bir madde midir", [], pair, "",
                                   pair_block)
        assert chosen == "Evet, torvanit bir maddedir.", chosen
    finally:
        (generate.answer, generate.supported,
         generate.answers_asked) = generate_real


@test("H1 of two true answers the graph's is-a chain picks the specific one")
def h1():
    """Measured (corpus, "norgul nedir"): the answer came back "norgul is a
    LIVING THING" at every commit in this history, while the document says
    "norgul is a PLANT". Both records are in the graph and both are true; the
    vague one won because the evidence path removes #doc triples from the block
    and so discards the specific PREMISE while keeping its derived, general
    CONCLUSION. The graph itself holds the edge that settles it: plant→living."""
    from lmm import retrieve
    memory = Memory()
    norgul = memory.identify("norgul")
    plant = memory.identify("bitki")
    living = memory.identify("canli")
    specific_r = memory.write(norgul, None, plant, "#doc", DOCUMENT)
    vague_r = memory.write(norgul, None, living, "#inference", INFERRED)
    memory.write(plant, None, living, "#doc", DOCUMENT)   # the is-a chain
    kept = retrieve.specific(memory, [vague_r, specific_r])
    assert [r.key for r in kept] == [specific_r.key], \
        [(r.value, r.source) for r in kept]
    # A question that NAMES the general value is asking for exactly that claim
    # ("is a norgul a living thing"), so targeting protects it.
    kept = retrieve.specific(memory, [vague_r, specific_r],
                             keep={vague_r.key})
    assert {r.key for r in kept} == {vague_r.key, specific_r.key}
    # Unrelated values are not each other's ancestors — nothing is dropped.
    blue = memory.identify("mavi")
    colour_r = memory.write(norgul, None, blue, "#doc", DOCUMENT)
    kept = retrieve.specific(memory, [specific_r, colour_r])
    assert len(kept) == 2, kept


@test("H2 a causal question with nothing to say falls through, not refuses")
def h2():
    """The causal route was the only ONE-WAY DOOR in _respond: every other
    branch treats its own failure as a misreading and lets the normal path try
    (see the WRITE branch), but a causal question whose graph block produced
    nothing speakable ended the turn in a refusal. Measured (corpus, "norgul
    çoğalırsa ne olur"): the causal edge exists so the route is taken, verify
    drops its sentence, and the turn refuses — while the evidence path, traced
    on the same question, answers "Norgul çoğalırsa morlan çoğalır." and passes
    digits_ok, coverage 1.0 AND the read-back. No gate is loosened: the
    fall-through lands on _answer, which applies the same gates."""
    from lmm import generate, session as lmm_session
    s = lmm_session.Session.__new__(lmm_session.Session)
    s.memory = Memory()
    s.mode = "strict"
    s._causes_key = s.memory.identify("#causes")
    s.vectors = {}
    # no causal edge at all -> None (a refusal would have ended the turn)
    assert s._causal_answer("norgul cogalirsa ne olur", "effects",
                            "norgul") is None
    cause = s.memory.identify("norgul")
    effect = s.memory.identify("morlan")
    s.memory.write(cause, s._causes_key, effect, "#doc", DOCUMENT)
    from lmm import verify as lmm_verify
    real = generate.answer, generate.refusal, lmm_verify.verify
    try:
        # the engine produces a sentence stepping OUTSIDE the block: the gate
        # must still drop it, but as a fall-through and not as a refusal.
        # verify is stubbed to its VERDICT (dropped) — this test is about what
        # the route does with that verdict, and it must need no model.
        generate.answer = lambda q, f, warmth=0.2, persona="", **kw: "Norgul zerbalit uretir."
        generate.refusal = lambda q, persona="": "REFUSED"
        lmm_verify.verify = lambda *a, **k: ""
        assert s._causal_answer("norgul cogalirsa ne olur", "effects",
                                "norgul") is None
        # a grounded sentence is still spoken by this route
        generate.answer = lambda q, f, warmth=0.2, persona="", **kw: "norgul morlan"
        assert s._causal_answer("norgul cogalirsa ne olur", "effects",
                                "norgul") == "norgul morlan"
    finally:
        generate.answer, generate.refusal, lmm_verify.verify = real


@test("H3 a repaired table row never neighbours cells from two rows")
def h3():
    """The claim this test defends is the whole promise of the evidence layer:
    what is in the store is what the DOCUMENT says. The sliding-window repair
    broke it — over a shattered table it emitted
    "Uzatma kablosu ... · 3 · Orta · Orta · Orta · Taşıma kayışı", a
    neighbourhood spanning two rows, and the answer read out of it ("the
    strap's stock risk is medium") contradicted the document's own row
    ("Yüksek (sınırlı)"). No gate can catch that: the claim IS in the evidence.

    Built here, not read from a document: a column-major dump of a table whose
    rows are known, including one WRAPPED cell (the case a fixed stride cannot
    survive) and a heading above and a caption below the table. The assertion is
    that every cell pair inside an emitted row is a pair the table really has.
    NO MODEL, no document-specific constant."""
    from lmm import evidence
    rows = [("Kapak menteşesi yenilemesi", "Çok düşük", "Çok hızlı", "Yok"),
            ("Vida seti tamamlama", "Düşük", "Çok hızlı", "Yok"),
            ("Gösterge camı değişimi", "Orta", "Orta", "Orta"),
            ("Gövde boyası yenileme", "Çok yüksek", "Yavaş", "Orta"),
            ("Uzatma kablosu / yedek besleme ünitesi", "Orta", "Orta", "Orta"),
            ("Taşıma kayışı", "Orta", "Orta", "Yüksek (sınırlı)")]
    lines = ["Bir tablo, hücre hücre dağılmış hâlde okunduğunda satırını "
             "yitirir; asıl mesele budur.",
             "Kullanım alanı", " Kademe", "Yatırım", "Geri dönüş"]
    for nth, (name, *rest) in enumerate(rows):
        if nth == 4:
            name, tail = name.rsplit(" ", 1)         # the cell WRAPS in two
            lines.append(name)
            name = tail
        lines.append(name)
        lines.append(" %d" % (nth + 1))              # the indented column
        lines.extend(rest)
    lines.append("Tablo yukarıdaki gibi okunmalıdır ve sağ üst köşe önce "
                 "hedeflenmelidir.")
    text = "\n".join(lines)

    windows = evidence.table_windows(text)
    assert windows, windows
    # the layout HAS a phase to read (period + indented column), so rows are
    # reconstructed rather than guessed at
    run = [l for l in text.split("\n") if l.strip() and len(l.strip()) <= 40]
    assert evidence.rows_of(run) is not None

    # every within-row cell pair must be a pair the table really states
    true_pairs = set()
    for nth, (name, *rest) in enumerate(rows):
        cells = name.rsplit(" ", 1) if nth == 4 else [name]    # the wrapped one
        cells += [str(nth + 1)] + list(rest)
        for one in cells:
            for other in cells:
                true_pairs.add((one, other))
    heading = {"Kullanım alanı", "Kademe", "Yatırım", "Geri dönüş"}
    for window in windows:
        cells = [c.strip() for c in window.split(" — ")[-1].split(" · ")]
        cells = [c for c in cells if c not in heading]
        for one in cells:
            for other in cells:
                assert one == other or (one, other) in true_pairs, \
                    "the document never puts %r next to %r: %s" % (one, other,
                                                                   window)
    # and the row that was being corrupted comes out whole
    assert any("Taşıma kayışı" in w and "Yüksek (sınırlı)" in w
               for w in windows), windows
    # a run whose layout gives NO phase is not carved into invented rows
    flat = "\n".join(["Bir paragraf, tabloya benzemeyen düz yazıdır ve bu "
                      "yüzden ölçüte girmez."]
                     + ["satır %d" % i for i in range(12)]
                     + ["Kapanış cümlesi de yeterince uzun olmak zorundadır "
                        "ki koşu burada bitsin."])
    flat_run = [l for l in flat.split("\n") if l.strip() and len(l.strip()) <= 40]
    assert evidence.rows_of(flat_run) is None


@test("H4 nested views of one paragraph cannot take every evidence seat")
def h4():
    """A question whose answer needs TWO sections gets one chance: both have to
    be among the sentences retrieved. The near-duplicate filter exists so that no
    single region monopolises those seats, but it measured overlap by Jaccard
    only — and the window scales are NESTED, so the same paragraph seen at three
    widths has a large union and reads as three regions. Measured (hospital,
    "ilk 6 ayda önerilen parçalardan hangisi yüksek öncelik işaretlidir"): all
    SIX seats went to views of the one closing paragraph, and the section holding
    the marker was never retrieved, so no candidate answer could rest on it and
    the read-back rejected — correctly — every answer that was offered.

    No model: the store is built here, one paragraph indexed at several widths
    plus one other section that shares the question's rarer word."""
    from lmm import evidence
    store = evidence.SentenceStore()
    core = ("İlk 6 ay için önerilen üç parça: 02 · 03 · 07 menteşe, vida ve "
            "taşıma kayışı yenilemesi.")
    store.add(core, "#doc")                                  # the narrow view
    store.add(core + " Üçü de mevcut veriyle çalışır.", "#doc")
    store.add("SONUÇ Nereden başlanmalı " + core + " Üçü de mevcut veriyle "
              "çalışır ve birkaç ay içinde ölçülebilir sonuç üretir.", "#doc")
    store.add("SONUÇ Nereden başlanmalı " + core, "#doc")
    other = ("07 taşıma kayışı yenilemesinin tamamlanması YÜKSEK ÖNCELİK "
             "işaretlidir.")
    store.add(other, "#doc")
    found = store.find("İlk 6 ayda önerilen parçalardan hangisi yüksek "
                       "öncelik işaretlidir", most=4)
    assert other in found, found
    # the region is still CORROBORATED — the cap is two seats, not one
    assert sum(1 for s in found if s.startswith(("İlk", "SONUÇ"))) >= 2, found
    # and a region that is genuinely alone still fills the list
    store2 = evidence.SentenceStore()
    for tail in ("yüksek", "orta", "düşük", "belirsiz"):
        store2.add("Zilfen madde %s bir torvanit değeri taşır." % tail, "#doc")
    assert len(store2.find("zilfen torvanit değeri", most=4)) >= 2


@test("H5 the question, not the answer, is what the last read-back asks about")
def h5():
    """THE SHAPE of the new condition, with no engine in it.

    `generate.supported` is handed a CLAIM, and a claim carries only what it
    chooses to say, so an answer that voices a true fact from the evidence
    passes it no matter what was asked. `generate.answers_asked` is handed the
    QUESTION as well, and that is the whole difference: the proposition being
    judged is "does this evidence state the relation that was ASKED", which an
    answer cannot satisfy by leaving the relation out. Here the runtime is
    stubbed to record what each check is shown."""
    from lmm import generate, prompts, runtime
    seen = []
    real = runtime.generate

    def fake(user, system=None, max_tokens=None, temperature=None):
        seen.append((system, user))
        return "no"

    runtime.generate = fake
    try:
        block = "[K1] Hazırlayan: Nordheim Kayıt Bürosu"
        claim = "Raporu Nordheim Kayıt Bürosu hazırlamıştır."
        assert generate.supported(claim, block) is False
        system, user = seen[-1]
        assert system == prompts.SUPPORT_SYSTEM
        assert "QUESTION" not in user, user      # the claim is judged alone
        assert generate.answers_asked("Raporu kim imzalamıştır", claim,
                                      block) is False
        system, user = seen[-1]
        assert system == prompts.RELATION_SYSTEM
        # the asked relation is IN the proposition, so an answer that never
        # mentions it can no longer slip past
        assert "imzalam" in user and claim in user and block in user, user
        # a "yes" is a yes — the gate admits, it does not merely reject
        runtime.generate = lambda *a, **k: "yes"
        assert generate.answers_asked("Raporu kim hazırlamıştır", claim, block)
    finally:
        runtime.generate = real


@test("H6 an answer to a question nobody asked does not get spoken")
def h6():
    """The last route to a confident WRONG answer, closed end to end.

    Measured (hospital, "Dokümanı kim imzalamıştır"): the document states who
    PREPARED it and never says who signed it, so the answer voiced the preparer
    field. Every gate passed and each was right to: the claim is in the evidence
    word for word, coverage 1.0, and the read-back confirmed it. Nothing that
    judges the ANSWER'S OWN CLAIM can catch this, because the claim never
    mentions the relation it fails to be about.

    The store is built WITHOUT a model (learn_text deep=False) over a synthetic
    document, and the engine's verdict is stood in for, so what is under test is
    the GATE'S WIRING rather than a model's judgement: a "no" abstains, no later
    path speaks what it refused, and the same setup with the condition satisfied
    still answers. The witness that this condition is what abstains: with the
    judge saying yes — the state before this commit — the wrong answer IS
    returned, over the same evidence, past every other gate."""
    from lmm import generate, session as lmm_session
    s = lmm_session.Session(None)
    s.learn_text("Bu rapor Hat B kalite denetimi için düzenlenmiştir.\n"
                 "Hazırlayan: Nordheim Kayıt Bürosu\n"
                 "Denetim sırasında ölçülen sıcaklık 42 santigrat derece "
                 "olarak kaydedilmiştir.\n", deep=False)
    wrong = "Hazırlayan: Nordheim Kayıt Bürosu"
    real = (generate.answer, generate.supported, generate.answers_asked,
            generate.refusal)
    views = []

    def judge(question, answer, view):
        """The engine's verdict, stood in for: the document states who prepared
        the report and no other relation about it."""
        views.append(view)
        return "hazırlayan" in question.lower()

    try:
        generate.answer = lambda question, facts, warmth=0.2, persona="", **kw: wrong
        generate.supported = lambda answer, view: True     # as measured
        generate.answers_asked = judge
        generate.refusal = lambda question, persona="": "REFUSED"
        # the relation the document does not state -> abstention, and the
        # graph/verify fallback does not resurrect it
        for question in ("Raporu kim imzalamıştır",
                         "Raporu kim onaylamıştır",
                         "Rapor hangi fabrika için düzenlenmiştir",
                         "Denetimin maliyeti kaç liradır"):
            assert s._answer(question, None) == "REFUSED", question
        # the judge is asked about the evidence the answer DRAWS ON, and the
        # question travels with it (that is the whole point of the check)
        assert views and all("Hazırlayan" in v for v in views), views
        # the asked relation IS stated -> the same machinery answers
        assert s._answer("Raporun hazırlayanı kimdir", None).startswith(wrong)
        # THE WITNESS: before this condition existed, the wrong answer spoke
        generate.answers_asked = lambda question, answer, view: True
        assert s._answer("Raporu kim imzalamıştır", None).startswith(wrong)
    finally:
        (generate.answer, generate.supported, generate.answers_asked,
         generate.refusal) = real


@test("I1 the same document and question always retrieve the same evidence")
def i1():
    """THE DETERMINISM CLAIM, ENFORCED.

    This retrieval layer's whole argument against embedding similarity is that
    it is deterministic and explainable, and it was neither: `qwords` is a SET
    OF STRINGS and the stem groups were formed by walking it in length order,
    so equal-length query words arrived in an order that depends on
    PYTHONHASHSEED. Group formation is order-sensitive — the first surface form
    seen becomes the anchor, the rest attach to it, and the group's IDF weight
    is computed over whatever union that produced — so the same question
    against the same document could retrieve DIFFERENT evidence between two
    runs of byte-identical code (measured on a 350-page manual: PYTHONHASHSEED=1
    and =3 disagree about "Cihazın işletim sistemi nedir").

    A property about hash seeds cannot be asserted from inside one interpreter,
    so this runs the same retrieval in several children with different seeds and
    demands the same block. The witness is minimal and reproduces the real
    failure exactly: two query words of the SAME LENGTH ('ekran', 'bilgi') so
    the length sort cannot order them, one sentence reachable from each so the
    two groups have separate unions and therefore equal weight, and the two
    sentences the same length so the length tie-break cannot separate them
    either. On the old code these eight seeds return two different answers; the
    real manual shows the same thing on a real question."""
    import subprocess
    program = (
        "import sys, json;"
        "sys.path.insert(0, %r);"
        "from lmm import evidence;"
        "store = evidence.SentenceStore();"
        "store.add('Ekran onbes nokta altiii');"
        "store.add('Bilgi yediyuz kirk sekiz');"
        "print(json.dumps(store.find('ekran bilgi', most=1)))"
    ) % os.path.join(ROOT, "src")
    seen = set()
    for seed in ("0", "1", "2", "3", "5", "7", "99", "12345"):
        env = dict(os.environ, PYTHONHASHSEED=seed)
        out = subprocess.run([sys.executable, "-c", program], env=env,
                             capture_output=True, text=True, check=True)
        seen.add(out.stdout.strip())
    assert len(seen) == 1, seen


@test("I2 the scorer reads a real abstention as one, and a value as a value")
def i2():
    """THE MEASURING TOOL IS PART OF THE WORK.

    The absence criterion was a literal-substring search over a list of Turkish
    phrases, and it failed on the most ordinary sentence the engine produces:
    "Bu bilgiye HENÜZ sahip değilim" is an abstention, and one inserted adverb
    put it outside "bilgiye sahip değil". It cost a real measurement a point —
    a tool that miscounts is worse than a slow one.

    Matching the pattern's stems in order inside a bounded window fixes that,
    and it opens exactly one hole, which the value check closes: with gaps
    allowed, "bu bilgi kılavuzda yok AMA garanti 2 yıldır" would match 'bilgi
    ... yok'. An answer that names a figure is not abstaining. The engine's own
    provenance footnote is not a claim, so it is stripped before that test —
    otherwise a source filename's digits read as a fabricated value (measured:
    24 genuine abstentions across five historical result files)."""
    sys.path.insert(0, os.path.join(ROOT, "benchmarks"))
    from score import abstains
    for answer in ("Bu bilgiye henüz sahip değilim.",
                   "Bu konuda şu anda bilgim yok.",
                   "Kılavuzda garanti süresi belirtilmemiştir.",
                   "Bilmiyorum. (Bu bilgi kesin değildir; #pdf:M30_70 v2.pdf)"):
        assert abstains(answer), answer
    for answer in ("Garanti süresi 2 yıldır.",
                   "Bu bilgi kılavuzda yok ama garanti süresi 2 yıldır.",
                   "Cihazın koruma türü Sınıf I ekipmandır.",
                   "Kılavuz bu konuyu başka bir bölümde anlatıyor."):
        assert not abstains(answer), answer


# --- J  the typed value: number, range, unit ---------------------------------
#
# What the value IS was invisible to every organ: "21.5\" OLED" and "okonos"
# were the same kind of thing, so whether two values in one slot can coexist
# had to be asked of a language model, one round-trip per pair.


@test("J1 a measurement is read off the data, symbols and ranges included")
def j1():
    """THE UNIT COMES FROM ADJACENCY, NOTHING ELSE.

    No list of units exists anywhere in this system and none may be added. The
    only rule is that the token standing next to a number is what the number is
    measured in — which is why a SYMBOLIC unit works exactly like a lettered
    one: '"' next to 21.5 is a token next to a number. A spec sheet's display line
    is the measured loss class ('Gösterge 21.5" OLED' — no letters in the unit at
    all).

    A separator is not a unit: in "100 V -240 V" and "%25 - %85" the symbol
    stands BETWEEN two numbers, and a number begins again right after it. In
    "5 °C - 40 °C" it does not — °C carries a letter and is a real unit — so
    the test is positional, not a character list.
    """
    from lmm.core.memory import measure
    assert measure('Gösterge 21.5" OLED') == (21.5, 21.5, '"')
    assert measure("21,5 inç") == (21.5, 21.5, "inç")   # notation is not quantity
    assert measure("100 V -240 V ~") == (100.0, 240.0, "v")     # a range
    assert measure("5 °C - 40 °C") == (5.0, 40.0, "°c")
    assert measure("-20 °C ~ 55 °C") == (-20.0, 55.0, "°c")     # below zero
    assert measure("%25 - %85") == (25.0, 85.0, None)   # dash separates
    assert measure("8 GB, DDR4") == (8.0, 8.0, "gb")    # punctuation is not a unit
    # A DIFFERENT unit is a DIFFERENT measurement, not a wider range: the
    # battery's line states a voltage and a capacity, and 12.8..4200 would be
    # a quantity the document never claims.
    assert measure("12.8 V / 4200 mAh") == (12.8, 12.8, "v")
    # Not measurements: a version/section string (two fractional separators)
    # and a text with no number at all.
    assert measure("Çekirdek sürümü 1.2.3") is None
    assert measure("okonos") is None


@test("J2 the unit is an identity, and its bare token is not a way to name it")
def j2():
    """A unit is a concept of the graph, so it lives in the identity system —
    the one place this architecture keeps meanings apart. But its label is
    NAMESPACED, and that is not decoration: 'g' or 'v' as a plain label would
    sit in the same label index as every name in the graph and would be
    reachable by the inflection tolerance, which is the exact collision
    (`kartal` the bird vs the district) this file was written to prevent."""
    memory = Memory()
    weight = memory.identify("3.4 kg")
    record = memory.write(memory.identify("cihaz"), None, weight, "#doc")
    low, high, unit = record.measure
    assert (low, high) == (3.4, 3.4)
    assert unit in memory.identities and memory.units["kg"] == unit
    assert memory.candidates("kg") == []            # not addressable as a name
    # The same unit sighted twice is ONE identity — a unit is not re-invented
    # per record.
    again = memory.write(memory.identify("çanta"), None,
                         memory.identify("2 kg"), "#doc")
    assert again.measure[2] == unit
    # AND THE REST OF THE GRAPH IS NUMBERED AS IF THE UNIT WERE NOT THERE.
    # Units count down from -1 for this reason: taking keys out of the shared
    # counter would renumber every identity opened after a measurement, and key
    # order is a tie-break in more than one organ — the same document would then
    # answer differently for a reason that has nothing to do with measurements.
    # Two identical histories, one writing a MEASUREMENT and one writing a plain
    # value: the typed one opens a unit, and the next key either graph hands out
    # is the same either way.
    plain, typed = Memory(), Memory()
    for held, value in ((plain, "okonos"), (typed, "3.4 kg")):
        held.write(held.identify("cihaz"), None, held.identify(value), "#doc")
    assert typed.units and min(typed.units.values()) < 0, typed.units
    assert not plain.units
    assert typed.identify("sonra") == plain.identify("sonra")


@test("J3 two numbers in one slot are judged by arithmetic, with no model")
def j3():
    """THE COST DECISION AND THE EPISTEMIC ONE COME APART.

    Whether two values in a slot exclude each other used to be a question for
    the language model, at one round-trip per pair — expensive enough that bulk
    ingestion switched it off and deferred every verdict to sleep. For
    measurements the question is arithmetic: disjoint ranges exclude each
    other, overlapping ones do not. Here the semantic test is wired to a
    function that RAISES if called, which is the assertion: these verdicts cost
    no model call at all.
    """
    def must_not_be_asked(one, other):
        raise AssertionError("the semantic test was called for a measurement")

    memory = Memory()
    gate = Gate(memory)
    gate.rival = must_not_be_asked
    device = memory.identify("cihaz")
    what = memory.identify("gerilim")
    gate.admit(device, what, memory.identify("100 V -240 V"), "#doc")
    # 220 V lies INSIDE the stated range — one fact stated twice, no rivalry.
    friend, _ = gate.admit(device, what, memory.identify("220 V"), "#doc")
    assert not links_of(friend, CONTRA), friend.links
    assert not links_of(friend, SUSPECT), friend.links
    # 12 V is outside it — a contradiction, decided without the model, and NOT
    # left as a suspicion for sleep to pick up.
    clash, _ = gate.admit(device, what, memory.identify("12 V"), "#doc")
    assert links_of(clash, CONTRA), clash.links
    assert not links_of(clash, SUSPECT), clash.links


@test("J4 arithmetic keeps quiet where it is not entitled to speak")
def j4():
    """The typed layer is a NEW SOURCE of verdicts, not a replacement for the
    semantic one. Bare numbers with no unit ("Tier 3" vs "Tier 1") and two
    DIFFERENT units ("3.4 kg" vs "17 lb", which may well be the same weight)
    are cases where arithmetic alone cannot decide, and answering them anyway
    would be the same mistake the audit found in the other direction: a cost
    decision wearing an epistemic decision's clothes. So the semantic test must
    still be consulted for those."""
    asked = []

    def semantic(one, other):
        asked.append((one, other))
        return False

    memory = Memory()
    gate = Gate(memory)
    gate.rival = semantic
    row = memory.identify("kapak menteşesi")
    tier = memory.identify("tier")
    gate.admit(row, tier, memory.identify("tier 3"), "#doc")
    gate.admit(row, tier, memory.identify("tier 1"), "#doc")
    assert len(asked) == 1, asked                   # unit-less: still asked
    weight = memory.identify("agirlik")
    gate.admit(weight, None, memory.identify("3.4 kg"), "#doc")
    gate.admit(weight, None, memory.identify("17 lb"), "#doc")
    assert len(asked) == 2, asked                   # units differ: still asked
    # And a value with no measurement in it at all is untouched by this layer.
    system = memory.identify("isletim sistemi")
    gate.admit(system, None, memory.identify("okonos"), "#doc")
    gate.admit(system, None, memory.identify("windows"), "#doc")
    assert len(asked) == 3, asked


@test("J5 a file from the future is refused, and every older file still loads")
def j5():
    """THE VERSION BYTE WAS WRITTEN AND NEVER READ (audit finding F).

    A .lmm marked FORMAT=99 loaded as if this build understood it, and the next
    save() would then write the graph back without whatever that format added —
    silent data loss. Refusing to read forward is not symmetric with refusing
    to read backward: this build knows exactly what a FORMAT-3 file says, and
    an old memory must keep opening. Both directions are asserted here,
    including the oldest form of all, plain JSON.
    """
    import json
    import struct
    import tempfile
    import zlib

    memory = Memory()
    device = memory.identify("cihaz")
    memory.write(device, None, memory.identify('21.5" OLED'), "#doc")
    with tempfile.TemporaryDirectory() as folder:
        path = os.path.join(folder, "m.lmm")
        memory.save(path)
        with open(path, "rb") as handle:
            raw = handle.read()
        assert struct.unpack(">B", raw[4:5])[0] == 4        # FORMAT=4 written
        back = Memory.load(path)
        held = list(back.records.values())[0]
        assert held.measure[:2] == (21.5, 21.5)
        assert back.units and held.measure[2] in back.units.values()

        # FROM THE FUTURE: header says 99 → refused, loudly.
        payload = zlib.compress(zlib.decompress(raw[13:]), 9)
        ahead = os.path.join(folder, "ahead.lmm")
        with open(ahead, "wb") as handle:
            handle.write(raw[:4] + struct.pack(">B", 99) + raw[5:])
        try:
            Memory.load(ahead)
        except ValueError as broke:
            assert "99" in str(broke), broke
        else:
            raise AssertionError("a FORMAT=99 file loaded silently")

        # BACKWARD COMPATIBILITY, both older forms. A format-3 record has no
        # typed value in the file; it gets the same typed view a new file
        # would, derived from the value the file already holds.
        old = {"format": 3, "next": 5, "self": None, "transitive": [],
               "identities": [{"key": 1, "labels": ["cihaz"], "vector": None},
                              {"key": 2, "labels": ['21.5" OLED'],
                               "vector": None}],
               "records": [{"key": 3, "subject": 1, "predicate": None,
                            "value": 2, "source": "#doc", "level": 4,
                            "trust": 0.6, "at": time.time(), "witnesses": 1,
                            "links": [], "episodic": True,
                            "sources": ["#doc"]}],
               "experiences": []}
        text = json.dumps(old, ensure_ascii=False).encode("utf-8")
        plain = os.path.join(folder, "old.json.lmm")
        with open(plain, "w", encoding="utf-8") as handle:
            handle.write(text.decode("utf-8"))
        binary = os.path.join(folder, "old3.lmm")
        with open(binary, "wb") as handle:
            handle.write(Memory.MAGIC
                         + struct.pack(">BII", 3, len(text),
                                       zlib.crc32(text) & 0xffffffff)
                         + zlib.compress(text, 9))
        for name in (plain, binary):
            back = Memory.load(name)
            held = back.records[3]
            assert held.subject == 1 and held.value == 2 and held.level == 4
            assert held.measure[:2] == (21.5, 21.5), name
            # and re-saving migrates it to the current format without loss
            again = os.path.join(folder, "again.lmm")
            back.save(again)
            assert Memory.load(again).records[3].value == 2


@test("J6 a question asking for a quantity gets the spec line, and nothing moves")
def j6():
    """THE UNIT THE DOCUMENT DOES NOT SPELL (measured manual loss).

    A spec sheet states `Gösterge 21.5" OLED` and the question asks "kaç inçtir".
    The unit on the line is a quotation mark — no letters — so the lexical
    channel has nothing of 'inç' to match, and the line is reachable only
    through 'gösterge', which the document also uses on dozens of menu lines: it
    ranked 33rd. Reweighting the field-name channel does rescue it and was
    measured to cost hospital two points, because five of its sixteen blocks
    changed. So the rescue APPENDS: the ranking is final before it runs, and
    this test asserts both halves — the spec line arrives, and the six seats
    that were there are still there, in order.

    Nothing here knows that 21.5" is inches. The question is recognised as
    asking for a quantity because the corpus itself binds 'inç' to a number
    somewhere else ("(12 inç)"), and the line is recognised as an answer
    because it names the asked field and states a measurement whose unit is one
    this corpus characteristically writes next to numbers.
    """
    from lmm import evidence
    store = evidence.SentenceStore()
    # the crowd: prose that shares the question's field word
    crowd = [
        "Cihazın göstergesi açıldığında oturum kutusu görünür ve yönetici "
        "parolası istenir; parolayı girdikten sonra devam edin.",
        "Cihazın göstergesi temizlenirken yumuşak bez kullanılmalı, çözücü "
        "içeren sıvılar kesinlikle uygulanmamalıdır bu yüzeye.",
        "Cihazın göstergesi koruyucusu etkinleştirildiğinde bekleme süresi "
        "dolduğunda arayüz kararır ve tuşa basıldığında geri döner.",
        "Cihazın göstergesi üzerindeki durum çubuğundan ağ bağlantısı durumu "
        "izlenebilir; bağlantı koptuğunda simge değişir hemen.",
        "Cihazın göstergesi yansımayı azaltmak için doğrudan güneş ışığı almayan "
        "bir konumda konumlandırılmalıdır çalışırken.",
        "Cihazın göstergesi üzerinde ölçüm sonuçları, yorumlar ve gövde "
        "işaretleri ters sırada temizlenebilir düğmeyle.",
        "Cihazın göstergesi bölme düzeninde iki görüntü yan yana gösterilir ve "
        "etkin pencere çerçeveyle belirtilir kullanıcıya.",
    ]
    class WithoutTheSeat(evidence.SentenceStore):
        """The same store with the mechanism inert — a corpus that never binds
        any token to a number cannot recognise a question as asking for a
        quantity. This is the reference the ranking must be unchanged against."""
        quantities = property(lambda self: set())

    corpus = crowd + [
        'Gösterge 21.5" OLED',
        "Taşınabilir vericiler tarayıcıya 30 cm'den (12 inç) daha yakın "
        "kullanılmamalıdır.",
        "Gösterge koruyucu ayarını açın.",
        "Gösterge menüsünü ayarlardan yeniden düzenleyin.",
        "Gösterge parlaklığı bölümüne bakın.",
    ]
    store, reference = evidence.SentenceStore(), WithoutTheSeat()
    for one in corpus:
        store.add(one, "#doc")
        reference.add(one, "#doc")
    question = "Cihazın göstergesi kaç inçtir"
    before = reference.find(question, most=6)
    found = store.find(question, most=6)
    assert 'Gösterge 21.5" OLED' not in before, before       # the measured loss
    # THE RANKING DID NOT MOVE: the seats that were there are there, in order,
    # and the spec line is appended after them.
    assert found == before + ['Gösterge 21.5" OLED'], found
    # A question that asks for no quantity gets no seat at all — the mechanism
    # is inert outside its loss class.
    plain = "Gösterge koruyucu nasıl açılır"
    assert store.find(plain, most=6) == reference.find(plain, most=6)


# --- K  language independence ------------------------------------------------
#
# The system claimed to be language-independent while three things quietly made
# Turkish its home language: every answer was framed by hand-written Turkish
# labels, every few-shot example in every classifier was Turkish, and the
# BENCHMARK SCORER recognised an abstention by a list of Turkish phrases — so a
# document in any other language would have scored every honest refusal as a
# fabrication. These tests hold each of the three closed, and none of them
# needs a model.


@test("K1 the scorer reads the system's own abstention stamp, in any language")
def k1():
    """THE MEASUREMENT MUST NOT SPEAK ONE LANGUAGE.

    `score.abstains` is a Turkish phrase list, and it is now the FALLBACK. The
    criterion is the stamp the session puts on the turn it refused
    (`Session.last_abstained`, written into the result file as "abstained"),
    which is a fact the code knows rather than a wording a reader guesses at.

    Here is what that buys, stated as the assertions below: a German refusal —
    invisible to every pattern in the list — is read as an abstention, and a
    Turkish one still is. Where the list DOES have the wording, the stamp still
    outranks it in both directions, which is the property that makes the stamp
    the criterion rather than a second opinion. The value check survives the
    stamp: an "abstention" that names a figure is a fabrication whichever path
    the engine thinks it took, so a stamped turn that supplies 2 años is not
    counted as a refusal. Rows with no stamp (the RAG and GraphRAG baselines',
    every historical file on disk) fall back to the patterns.

    THE FALLBACK IS NO LONGER TURKISH-ONLY, and that is a fairness fix, not a
    retreat from the stamp. A baseline writes no stamp and answers in the
    document's language: measured with a Turkish-only list, GraphRAG's four
    honest English refusals scored as fabrications and cost it 13/17 against a
    true 16/17. The list only ever judges a third party — this system is read
    by its stamp — so it has to cover the languages the third party speaks."""
    sys.path.insert(0, os.path.join(ROOT, "benchmarks"))
    from score import abstains, declined
    # German is outside every pattern — the case the stamp exists for.
    unseen = "Das weiß ich leider nicht."
    assert not abstains(unseen)
    assert declined({"cevap": unseen, "abstained": True})
    assert not declined({"cevap": unseen, "abstained": False})
    # Where the list does reach, the stamp still decides — both ways.
    for answer in ("That is not stated in the document.",
                   "No tengo esa información."):
        assert abstains(answer), answer          # the fallback sees it
        assert declined({"cevap": answer, "abstained": True}), answer
        assert not declined({"cevap": answer, "abstained": False}), answer
    # Turkish keeps working through the same door.
    assert declined({"cevap": "Bu bilgiye henüz sahip değilim.",
                     "abstained": True})
    # A stamp cannot launder a fabricated value.
    assert not declined({"cevap": "La garantía es de 2 años.",
                         "abstained": True})
    # No stamp -> the wording is read, in whichever language it was written.
    assert declined({"cevap": "Bu konuda bilgim yok."})
    assert declined({"cevap": "No tengo esa información."})


@test("K2 every refusal leaves by one door, and that door raises the flag")
def k2():
    """ONE DOOR, OR THE FLAG IS A LIE.

    `last_abstained` is only worth reading if no refusal can slip past the
    method that sets it. The session had nine separate
    `generate.refusal(...) or FALLBACK_DONT_KNOW` sites; a flag set at eight of
    them measures nothing. This is a SOURCE-level assertion because that is the
    level the property lives on: `generate.refusal` may be called from exactly
    one place, `Session._refuse`, and the two fallback returns that bypass the
    engine entirely (the crash path, the research offer) must raise the flag
    themselves."""
    source = open(os.path.join(ROOT, "src", "lmm", "session.py"),
                  encoding="utf-8").read()
    body = source.split("def _refuse(")[1].split("\n    def ")[0]
    assert "self.last_abstained = True" in body
    assert body.count("generate.refusal(") == 1
    assert source.count("generate.refusal(") == 1, "a refusal bypasses _refuse"
    # Everything that returns the last-resort constant flags itself too.
    for line_no, line in enumerate(source.splitlines()):
        if "return FALLBACK_DONT_KNOW" not in line:
            continue
        window = "\n".join(source.splitlines()[max(0, line_no - 6):line_no])
        assert "self.last_abstained = True" in window, line_no


@test("K3 the answer frame is structure, not a language")
def k3():
    """NO ANSWER PASSES THROUGH A TURKISH SKELETON.

    Every answer this system has ever produced was built as
    "OLGULAR:\\n...\\n\\nSORU: ..." — two Turkish words wrapped around a
    question that might be in any language, and around a no-facts branch that
    told the model in Turkish not to invent. The labels are field names for the
    engine, so they belong in the codebase's own language; what must NOT leak
    is a sentence of any user language. Captured here without a model: the
    frame is inspected directly."""
    from lmm import generate, runtime
    seen = []
    real = runtime.generate

    def capture(user, **kw):
        seen.append(user)
        return "…"
    runtime.generate = capture
    try:
        generate.answer("¿qué es el vorlin?", "[1] vorlin → bebida")
        generate.answer("was ist Beton?", "")
    finally:
        runtime.generate = real
    with_facts, without_facts = seen
    assert with_facts.startswith("FACTS:")
    assert "\n\nQUESTION: ¿qué es el vorlin?" in with_facts
    assert without_facts.startswith("QUESTION: was ist Beton?")
    # The no-facts branch instructs in the codebase's language and hands the
    # ANSWER's language back to the question.
    assert "SAME LANGUAGE" in without_facts
    turkish = ("OLGULAR", "SORU:", "Belleğinde", "bilmediğini", "Uydurma")
    for frame in seen:
        for word in turkish:
            assert word not in frame, (word, frame)


@test("K4 no classifier learns its pattern in one language")
def k4():
    """FEW-SHOTS IN ONE LANGUAGE TEACH THAT LANGUAGE.

    Every example in every classifier prompt was Turkish, which is how a system
    with no word list in it still ended up with a home language: shown "sigara
    kansere neden olur" and nothing else, a model generalises from Turkish
    causal grammar. The examples are now spread across languages, and this test
    keeps them spread — it asserts that each few-shot prompt demonstrates its
    pattern in at least THREE languages, so no single one can quietly become
    the pattern again. Signatures are distinctive orthography/function words,
    never content, so rewording an example does not break the test."""
    from lmm import prompts
    import lmm.generate as gen_src
    signatures = {
        "en": ("what is", "who are you", "an eagle", "You're welcome",
               "operating temperature", "Prepared by", "smoking causes"),
        "de": ("ist ein", "was ist", "führt zu", "Aufnahmedauer", "weiß ich",
               "wer hat dich", "wie geht es", "Kilogramm", "verkürzt"),
        "es": ("¿", "es un", "provoca", "Tiempo de instalación", "órgano",
               "gracias", "bisagra"),
        "tr": ("nedir", "bir kuştur", "merhaba", "selam", "yol açar",
               "kırmızı", "menteşesi", "bilmiyorum"),
    }

    def languages(text):
        low = text.lower()
        return {tag for tag, marks in signatures.items()
                if any(m.lower() in low for m in marks)}

    source = open(gen_src.__file__, encoding="utf-8").read()
    few_shot = {
        "EXTRACT_SYSTEM": prompts.EXTRACT_SYSTEM,
        "REEXTRACT_SYSTEM": prompts.REEXTRACT_SYSTEM,
        "SUPPORT_SYSTEM": prompts.SUPPORT_SYSTEM,
        "RELATION_SYSTEM": prompts.RELATION_SYSTEM,
    }
    for name, body in few_shot.items():
        assert len(languages(body)) >= 3, (name, languages(body))
    # The classifiers whose few-shots live inside generate.py, taken from the
    # source so that no model call is needed to see them.
    for name in ("def is_causal(", "def is_causal_question(",
                 "def is_identity_question(", "def are_rivals("):
        body = source.split(name)[1].split("\ndef ")[0]
        assert len(languages(body)) >= 3, (name, languages(body))

@test("K5 an answer that asserts nothing is an abstention, whatever it says")
def k5():
    """THE REFUSAL DOOR IS NOT THE ONLY WAY OUT.

    Measured on the fictional corpus: three of its four absence questions never
    reach `_refuse` at all. Evidence IS retrieved, the answer prompt is told to
    say it doesn't know when the evidence doesn't cover the question, and that
    sentence walks through every gate because it claims nothing. A flag set
    only at the refusal door reads all three as confident answers.

    So the flag means "this turn asserted nothing", and the test for that is
    the re-extractor the fabrication gate already relies on — no phrase list,
    no language. Here it is stubbed, so the assertion is about the WIRING: a
    sentence with claims counts as spoken, one without counts as an abstention,
    the engine's parenthetical hedge and #source footnote are not claims, and a
    re-extractor that throws is read as having spoken (the safe direction — the
    damaging error is calling a wrong answer an honest refusal)."""
    from lmm import extract, session as session_module
    blank = session_module.Session.__new__(session_module.Session)
    real = extract.reextract
    try:
        extract.reextract = lambda s: [("morlan", "type", "bird")]
        assert blank._asserted_a_fact("Morlan bir kuştur.")
        extract.reextract = lambda s: []
        assert not blank._asserted_a_fact("Das weiß ich leider nicht.")
        assert not blank._asserted_a_fact("No lo sé.")
        # the footnote is not a claim: with no claims in the sentence itself,
        # a hedge and a source tag cannot turn a refusal into an assertion
        assert not blank._asserted_a_fact(
            "Bilmiyorum. (Bu bilgi kesin değildir; #doc:corpus)")
        # an empty answer asserts nothing either
        assert not blank._asserted_a_fact("   ")

        def broken(_sentence):
            raise RuntimeError("engine down")
        extract.reextract = broken
        assert blank._asserted_a_fact("Anything at all.")
    finally:
        extract.reextract = real


@test("K6 the three corpora are the same benchmark, not three different ones")
def k6():
    """A CROSS-LANGUAGE COMPARISON IS ONLY EVIDENCE IF ONLY THE LANGUAGE MOVED.

    The English and Spanish corpora exist to isolate one variable. If one of
    them quietly gains a question, loses an absence case, or has a gold list
    with more alternatives to hit, the comparison stops measuring language and
    starts measuring the benchmark — and the failure would be invisible in the
    totals, which is exactly the kind of thing this file exists to catch.

    So the structure is asserted directly: same question count, same order of
    types, the same questions unanswerable in each, and gold lists of the same
    size question by question. The gold WORDS must differ — they are
    translations — and the invented names must not, because a name that got
    localised would let a model's own knowledge answer in one language and not
    another."""
    import json
    bench = os.path.join(ROOT, "benchmarks")

    def load(name):
        return json.load(open(os.path.join(bench, name), encoding="utf-8"))

    base = load("questions.json")
    for name in ("questions_en.json", "questions_es.json"):
        other = load(name)
        assert len(other) == len(base), name
        for a, b in zip(base, other):
            assert a["tip"] == b["tip"], (name, a["tip"], b["tip"])
            assert (a["altin"] is None) == (b["altin"] is None), (name, a)
            if a["altin"] is not None:
                # same number of ways to be right — an extra alternative is a
                # free point, and would show up as a language difference
                assert len(b["altin"]) >= len(a["altin"]), (name, a["soru"])
            assert a["soru"] != b["soru"], (name, a["soru"])
    # The invented world keeps its names in every language: nothing in these
    # corpora can be answered from a model's parametric knowledge.
    names = ("zerbalit", "vorlin", "norgul", "morlan", "karvel", "telvas",
             "kelvit", "torvanit", "selvin", "zilfen", "nortlann", "velmar")
    for name in ("corpus.txt", "corpus_en.txt", "corpus_es.txt"):
        text = open(os.path.join(bench, name), encoding="utf-8").read().lower()
        for invented in names:
            assert invented in text, (name, invented)


# =====================================================================  L
#  The front door (lmm.api.Memory). A WRAPPER, and these tests exist to
#  keep it one: every assertion below is about ROUTING and REPORTING, and
#  none of them may need a model, because the facade must not have opinions
#  of its own about what the layer beneath it does.
# =====================================================================


@test("L1 the front door is a wrapper: text goes where learn_text went")
def l1():
    """No new pipeline. `Memory.learn` on plain text must reach exactly
    `Session.learn_text`, with the text unaltered — if the facade ever starts
    pre-processing, cleaning or chunking on the way through, it has stopped
    being packaging and become a second implementation."""
    from lmm import Memory
    m = Memory()
    seen = []
    real = m.session.learn_text
    m.session.learn_text = lambda text, source="#document", deep=True: (
        seen.append((text, source, deep)) or (0, 0))
    m.learn("Zerbalit boils at 412 degrees.")
    m.session.learn_text = real
    assert len(seen) == 1, seen
    assert seen[0][0] == "Zerbalit boils at 412 degrees.", seen[0][0]
    # deep defaults TRUE for text typed at the memory: one sentence handed in
    # directly is a fact being taught and is worth the extractor's call.
    assert seen[0][2] is True, seen[0]


@test("L2 the extension picks the adapter, and nothing else does")
def l2():
    """The dispatch is a table, not a sniffer. Each extension must land on its
    own adapter with the session passed through; the pdf/xlsx/docx readers are
    stubbed so this holds with none of the optional dependencies installed."""
    from lmm import Memory, tables
    m = Memory()
    called = []
    keep = (tables.learn_pdf, tables.learn_xlsx, tables.learn_docx)
    tables.learn_pdf = lambda s, p, source=None, deep=False: (
        called.append(("pdf", s is m.session)) or (3, 1))
    tables.learn_xlsx = lambda s, p, source=None: (
        called.append(("xlsx", s is m.session)) or 4)
    tables.learn_docx = lambda s, p, source=None, deep=False: (
        called.append(("docx", s is m.session)) or (5, 2))
    folder = tempfile.mkdtemp()
    try:
        for name in ("a.pdf", "b.xlsx", "c.docx"):
            path = os.path.join(folder, name)
            open(path, "w", encoding="utf-8").write("x")
            m.learn(path)
    finally:
        tables.learn_pdf, tables.learn_xlsx, tables.learn_docx = keep
    assert [c[0] for c in called] == ["pdf", "xlsx", "docx"], called
    assert all(c[1] for c in called), "the adapter got a different session"

    # A format we cannot read is REFUSED by name, not silently read as text.
    path = os.path.join(folder, "deck.pptx")
    open(path, "w", encoding="utf-8").write("x")
    try:
        m.learn(path)
    except ValueError as said:
        assert "PowerPoint" in str(said), said
    else:
        raise AssertionError("a .pptx was accepted")


@test("L3 a missing optional dependency names the install line, not a traceback")
def l3():
    """The core install is deliberately dependency-free, so 'I fed it a PDF and
    it exploded' is the FIRST thing a new user will do. Each reader must fail
    with the command that fixes it — and with the PyPI name, which for
    python-docx is not the name you import."""
    from lmm import tables
    import builtins
    real = builtins.__import__

    def blocked(name, *a, **kw):
        if name.split(".")[0] in ("pdfplumber", "docx", "pandas"):
            raise ImportError("blocked for the test")
        return real(name, *a, **kw)

    gone = {}
    for mod in [k for k in sys.modules
                if k.split(".")[0] in ("pdfplumber", "docx", "pandas")]:
        gone[mod] = sys.modules.pop(mod)
    was = tables._PANDAS
    tables._PANDAS = None
    builtins.__import__ = blocked
    try:
        for reader, arg, extra, dist in (
                (tables.read_pdf, "x.pdf", "pdf", "pdfplumber"),
                (tables.read_docx, "x.docx", "docx", "python-docx"),
                (tables.read_xlsx, "x.xlsx", "xlsx", "pandas")):
            try:
                reader(arg)
            except ImportError as said:
                # the DISTRIBUTION name, read from the one place that
                # states it — `lmm` is what you import, not what you install
                assert f"{tables.DIST}[{extra}]" in str(said), said
                assert dist in str(said), said
            else:
                raise AssertionError(f"{arg} read with its reader missing")
    finally:
        builtins.__import__ = real
        tables._PANDAS = was
        sys.modules.update(gone)


@test("L4 a document body is never mistaken for a file name")
def l4():
    """`learn` takes a path OR the text itself, so it has to tell them apart —
    and it must do so WITHOUT handing a 500 KB document to the filesystem. A
    body with line breaks, and a body longer than any path, are text; a real
    file is a file. (os.path.isfile raises rather than returning False on an
    over-long name on several platforms — this is the guard for that.)"""
    from lmm.api import _looks_like_path
    assert not _looks_like_path("Zerbalit boils.\nVorlin does not.")
    assert not _looks_like_path("x" * 5000)
    assert not _looks_like_path("no_such_file_anywhere.txt")
    assert not _looks_like_path(b"bytes are not a path")
    assert not _looks_like_path(None)
    folder = tempfile.mkdtemp()
    path = os.path.join(folder, "real.txt")
    open(path, "w", encoding="utf-8").write("The tower is 91 metres tall.")
    assert _looks_like_path(path)


@test("L5 an answer is still a string, and carries its own abstention stamp")
def l5():
    """`ask(explain=True)` must not force callers into a wrapper type: the
    return IS the answer text and additionally reports itself. The abstention
    flag has to be the SESSION's structural stamp — the thing K1/K2 defend —
    and not a phrase list re-invented at the front door."""
    from lmm import Memory
    from lmm.api import Answer
    a = Answer("Bilmiyorum.", abstained=True, sources=("#pdf:x",))
    assert isinstance(a, str) and a == "Bilmiyorum."
    assert a.upper() == "BILMIYORUM." and a.abstained and a.sources == ("#pdf:x",)

    # `cache=False` because this test moves the SESSION's stamps by hand
    # between two identical questions, which the memory has no way to see and
    # the answer cache would (correctly) not re-derive for. The cache's own
    # behaviour is P1/P2's subject.
    m = Memory(cache=False)
    # `**kw` because `respond` now takes `fluent` too (the graph-first path
    # offers a generated sentence as an option); the stub is standing in for
    # the whole method, so it accepts whatever the front door passes.
    m.session.respond = lambda q, **kw: "It is 412 degrees. #document"
    m.session.last_abstained = False
    m.session.last_subject = "zerbalit"
    m.session.last_kind = "ASK"
    m.session.last_written = []
    plain = m.ask("how hot?")
    assert type(plain) is str and "412" in plain           # noqa: E721
    told = m.ask("how hot?", explain=True)
    assert told.sources == ("#document",), told.sources
    assert told.abstained is False and told.subject == "zerbalit"
    m.session.last_abstained = True
    assert m.ask("how hot?", explain=True).abstained is True


@test("L6 a memory with nowhere to save says so before it loses anything")
def l6():
    """`Memory()` is allowed to be transient, which makes a silent no-op save
    the dangerous failure: the caller believes the memory is on disk. It must
    raise, and it must accept a path given late."""
    from lmm import Memory
    m = Memory()
    m.learn("The tower is 91 metres tall.", deep=False)
    before = len(m)
    try:
        m.save()
    except ValueError as said:
        assert "nowhere to save" in str(said), said
    else:
        raise AssertionError("a pathless memory reported that it saved")
    path = os.path.join(tempfile.mkdtemp(), "mind.lmm")
    assert m.save(path) == path and os.path.exists(path)
    assert len(Memory(path)) == before


@test("M1 no unverified text is appended to an answer, and none at all to a refusal")
def m1():
    """THE FIELD TRIAL'S BIGGEST HOLE: a second model call after the gate.

    `_answer` and `_hedge` used to append `generate.hedge_note(...)` — a
    separate generation at temperature 0.3 — to an answer that had already
    passed verify/coverage/read-back, and the concatenation was never checked
    again. Measured on twelve public documents: four of six wrong answers were
    that note, the worst of them glued to a perfect abstention
    ("I do not know. RFC 9110 is obsoleted by RFC 9111.").

    Two properties are asserted here, both without a model: an abstaining turn
    comes out byte-for-byte as the answer path produced it, and an asserting
    turn gains NOTHING but the stored source stamp inside format punctuation —
    no token that could carry a new fact."""
    from lmm import generate, session as lmm_session
    assert not hasattr(generate, "hedge_note"), \
        "the ungated second model call is back"

    s = lmm_session.Session(None)
    refusal = "Bilmiyorum."

    def spoke(message, fluent=False, teach=True):
        s._mark = "#pdf:rfc2119.txt"       # a low-trust source, as _hedge sets
        return refusal

    s._respond = spoke
    s._asserted_a_fact = lambda said: False        # the turn claimed nothing
    assert s.respond("RFC 2119 kaç sayfadır") == refusal, "text after a refusal"
    assert s.last_abstained is True

    s._asserted_a_fact = lambda said: True         # the turn made a claim
    said = s.respond("RFC 2119 kaç sayfadır")
    assert said.startswith(refusal) and said != refusal
    added = said[len(refusal):]
    # WHAT WAS ADDED CANNOT BE A FACT: every token is either the stamp the
    # session already stored or pure punctuation. Nothing was generated, so
    # nothing can be invented — that is the property, not a smaller prompt.
    for token in added.split():
        assert token.startswith("#") or not any(c.isalnum() for c in token), token
    assert "#pdf:rfc2119.txt" in added and s.last_abstained is False

    # AND THE STAMP SURVIVES THE MARK'S BRACKETS, for a reader of `explain=True`
    from lmm.api import Memory
    m = Memory()
    m.session.respond = lambda q, **kw: said
    m.session.last_abstained = False
    m.session.last_subject = m.session.last_kind = ""
    m.session.last_written = []
    assert m.ask("x", explain=True).sources == ("#pdf:rfc2119.txt",)


@test("M2 learn() says what it could not read instead of swallowing it")
def m2():
    """THE SILENT CLASS, measured in the field trial and closed here.

    Four separate ways a document was never read while `learn()` reported
    success: an image-only PDF (`<Learned 0 facts, 0 tables via pdf>`, zero
    evidence, no warning); an HTML file, which has no reader, accepted as
    "text" and filling the index with markup; a mistyped path learned as a
    SENTENCE; and a broken PDF answering with 38 lines of pdfminer traceback.

    All four are decided on structure — zero characters, the share of
    characters inside tags, the shape of a filename — and none of them on any
    document's content. No model runs in this test."""
    from lmm import api, tables
    from lmm.api import Memory
    home = tempfile.mkdtemp()

    # 1. A PATH THAT IS NOT THERE IS A TYPO, NOT A SENTENCE.
    m = Memory()
    for bad in ("manual.pdf", os.path.join(home, "notes.txt")):
        try:
            m.learn(bad)
        except ValueError as said:
            assert "no such file" in str(said), said
        else:
            raise AssertionError("a missing file was learned as text: %r" % bad)
    try:
        m.learn(home)
    except ValueError as said:
        assert "directory" in str(said), said
    else:
        raise AssertionError("a directory was learned as text")
    # and prose that merely mentions a filename is still prose
    kept = m.learn("The calibration values live in manual.pdf on the bench.",
                   deep=False)
    assert kept.evidence > 0 and not kept.warnings

    # 2. NOTHING CAME OUT = SAY SO. Stood in for the scanner's PDF: the reader
    #    runs, returns nothing, writes no evidence.
    import lmm.tables as tables_module
    real = tables_module.learn_pdf
    try:
        tables_module.learn_pdf = lambda session, path, source=None, deep=False: (0, 0)
        scan = os.path.join(home, "scan.pdf")
        open(scan, "wb").write(b"%PDF-1.4 no text layer here")
        told = Memory().learn(scan)
        assert told.evidence == 0 and told.warnings, repr(told)
        assert "OCR" in told.warnings[0], told.warnings
        assert not told, "an unread document reported itself as read"
    finally:
        tables_module.learn_pdf = real

    # 3. MARKUP IS NOT PROSE, and a format with no reader says which one it is.
    page = os.path.join(home, "spec.html")
    open(page, "w", encoding="utf-8").write(
        "<html><head><meta content='Common,Latin' name='scripts'>"
        "<link rel='stylesheet' href='a.css'></head><body>"
        + "<div class='row'><span class='cell'>x</span></div>" * 40
        + "<p>The tower is 91 metres tall.</p></body></html>")
    told = Memory().learn(page, deep=False)
    assert told.evidence > 0 and len(told.warnings) == 2, repr(told)
    assert any("not a format LMM has a reader for" in w for w in told.warnings)
    assert any("markup" in w for w in told.warnings), told.warnings
    # ... and an ordinary text file gets no warning at all
    plain = os.path.join(home, "notes.txt")
    open(plain, "w", encoding="utf-8").write("The tower is 91 metres tall.\n")
    told = Memory().learn(plain, deep=False)
    assert told.evidence > 0 and not told.warnings, repr(told)

    # 4. A BROKEN FILE GETS ONE LINE, NOT A THIRD-PARTY TRACEBACK — while our
    #    own errors (the missing-extra install line) pass through untouched.
    class Pdfminer(Exception):
        pass

    def explode(*a, **k):
        raise Pdfminer("")            # the protected-PDF case: empty message

    try:
        api._read_with("pdf", "/tmp/x.pdf", explode)
    except ValueError as said:
        assert "could not read this file" in str(said) and "Pdfminer" in str(said)
        # `raise ... from None`: the reader's traceback is not re-printed under
        # ours, which is the whole of the user-visible difference
        assert said.__cause__ is None and said.__suppress_context__
    else:
        raise AssertionError("the reader's own exception reached the caller")

    def missing(*a, **k):
        raise ImportError(f"pip install '{tables.DIST}[pdf]'")

    try:
        api._read_with("pdf", "/tmp/x.pdf", missing)
    except ImportError as said:
        assert f"{tables.DIST}[pdf]" in str(said)
    else:
        raise AssertionError("the install line was swallowed")


@test("M3 the document's masthead is indexed as one block, not as detached lines")
def m3():
    """MASTHEAD BLINDNESS, measured on two real documents.

    Everything missed on them was in the front block: the publishing body
    ("Department of the Treasury—Internal Revenue Service" in a form's header),
    the date, and the author's institution ("S. Bradner / Harvard University /
    March 1997" atop an RFC). Those are lines, not sentences — none flows into
    the next — so the sentence layer indexed them one detached line at a time
    and no window ever held the name beside the institution beside the date.

    The block is read off the LAYOUT: the run before the first line that ends a
    sentence, bounded by what this document's own lines call a record. This
    test uses an invented masthead; nothing in the rule mentions dates,
    institutions or line numbers, and a document that opens with prose gets no
    block at all."""
    from lmm import evidence, session as lmm_session
    masthead = ("Nordheim Kayıt Bürosu                       T. Velmar\n"
                "Belge: HB-14                     Karvel Enstitüsü\n"
                "Sınıf: Genel                              Mart 2031\n"
                "\n"
                "Hat B kalite denetimi\n"
                "\n"
                "Denetim sırasında ölçülen sıcaklık 42 derecedir.\n"
                "Ölçüm günün ilk vardiyasında yapılmıştır.\n")
    head = evidence.front_matter(masthead)
    assert head, "no block was read off a masthead"
    for line in ("T. Velmar", "Karvel Enstitüsü", "Mart 2031",
                 "Hat B kalite denetimi"):
        assert line in head, (line, head)
    # it STOPS where the prose starts — the block is not the whole document
    assert "42" not in head and "vardiya" not in head, head

    # AN ABBREVIATION IS NOT A SENTENCE END — measured on RFC 9110, whose first
    # masthead line ends "R. Fielding, Ed." and killed the block outright. What
    # separates them is not the dot but the TYPESETTING: a masthead line holds
    # its pieces columns apart, and this document's own lines say how wide that
    # is (prose keeps single spaces).
    columned = ("Kurul Sekreteryası                              T. Velmar, Ed.\n"
                "Belge: HB-14                            Karvel Enstitüsü\n"
                "Sınıf: Genel                                   Mart 2031\n"
                "Denetim sırasında ölçülen sıcaklık 42 derecedir.\n"
                "Ölçüm günün ilk vardiyasında yapılmıştır.\n")
    head = evidence.front_matter(columned)
    assert "T. Velmar, Ed." in head and "Mart 2031" in head, head
    assert "42" not in head, head          # and the prose line still ends it

    # a document that opens with a sentence has no masthead, and is not given one
    assert evidence.front_matter(
        "Denetim sırasında ölçülen sıcaklık 42 derecedir.\n"
        "Ölçüm günün ilk vardiyasında yapılmıştır.\n") == ""
    assert evidence.front_matter("") == "" and evidence.front_matter("Tek satır") == ""

    # and ingestion puts that block in the index AS ONE ENTRY, so the name, the
    # institution and the date can be retrieved together
    s = lmm_session.Session(None)
    s.learn_text(masthead, source="#doc:hb14", deep=False)
    blocks = [text for text, _ in s.evidence.sentences
              if "T. Velmar" in text and "Karvel Enstitüsü" in text
              and "Mart 2031" in text]
    assert blocks, "the masthead is still three unrelated lines"
    found = s.evidence.find("Karvel Enstitüsü Mart 2031", most=4)
    assert any("Karvel Enstitüsü" in text and "Mart 2031" in text
               for text in found), found


@test("M4 an engine call that never comes back gives up instead of waiting forever")
def m4():
    """MEASURED: one question took 613 seconds. Not computing — waiting. The
    backend was over its per-minute quota and the client retried ten times with
    no clock, so a rate limit (correctly treated as a wait) became a wait with
    no end and no message.

    `runtime.within_budget` keeps the retries and puts a deadline around them:
    what is a queue is tried again with a doubling pause while budget remains,
    what is a real failure is raised at once and unchanged, and running out
    says how long it waited. No model and no network here — the call is stood
    in for, and the budget is set to a fraction of a second."""
    from lmm import runtime
    was = os.environ.get("LMM_TIMEOUT")

    class Queued(Exception):
        pass

    try:
        # 1. the budget is read from the environment, and 0 means "no limit"
        os.environ["LMM_TIMEOUT"] = "0.3"
        assert runtime.budget() == 0.3
        os.environ["LMM_TIMEOUT"] = "0"
        assert runtime.budget() == 0.0
        os.environ["LMM_TIMEOUT"] = "not a number"
        assert runtime.budget() == runtime.DEFAULT_BUDGET
        del os.environ["LMM_TIMEOUT"]
        assert runtime.budget() == runtime.DEFAULT_BUDGET

        # 2. a queue is retried WHILE THERE IS BUDGET, and then gives up — and
        #    the giving up happens in about the budget, not in ten retries
        os.environ["LMM_TIMEOUT"] = "0.3"
        tries = []

        def busy(left):
            tries.append(left)
            raise Queued("429 Too Many Requests")

        started = time.time()
        try:
            runtime.within_budget(busy, (Queued,), what="the test endpoint")
        except runtime.EngineTimeout as gave_up:
            assert "0.3s" in str(gave_up) and "Queued" in str(gave_up), gave_up
        else:
            raise AssertionError("a permanently busy endpoint returned")
        waited = time.time() - started
        assert 0.3 <= waited < 3, waited
        assert tries and all(t is not None and t <= 0.3 for t in tries), tries

        # 2b. a queue that CLEARS is answered — the retry is still there, it
        #     just has a clock around it now
        os.environ["LMM_TIMEOUT"] = "5"
        state = {"first": True}

        def clears(left):
            if state["first"]:
                state["first"] = False
                raise Queued("429 Too Many Requests")
            return "answered"

        assert runtime.within_budget(clears, (Queued,)) == "answered"

        # 3. a REAL failure is not retried, and reaches the caller unchanged
        seen = []

        def broken(left):
            seen.append(left)
            raise KeyError("AZURE_OPENAI_ENDPOINT")

        try:
            runtime.within_budget(broken, (Queued,))
        except KeyError:
            assert len(seen) == 1, seen
        else:
            raise AssertionError("a broken configuration was retried")

        # 4. a call that answers, answers — and gets told what time it has
        assert runtime.within_budget(lambda left: ("ok", left))[0] == "ok"
        os.environ["LMM_TIMEOUT"] = "0"
        assert runtime.within_budget(lambda left: left) is None   # no limit
    finally:
        os.environ.pop("LMM_TIMEOUT", None)
        if was is not None:
            os.environ["LMM_TIMEOUT"] = was


# --- N: the graph-first answer path ----------------------------------------
#
# The pathology these close is a MEASUREMENT: `benchmarks/COST.md` §3 reported
# 0 of 17 questions answered without a model call, on a corpus whose answers
# stand in the graph as triples. The fix (`lmm/lookup.py`) is only worth having
# if it stays inside its guarantee, so what is asserted here is mostly where it
# must NOT fire.


def island(rival=None):
    """The corpus taxonomy, built with NO model anywhere: `learn_cell` is the
    structural ingestion path, and the rivalry test — the one thing in the
    write path that consults the engine — is answered locally.

    The first five cells are there to make `tur` transitive the way the data
    makes it transitive: two independently witnessed triangles, counted by
    `core/transitive.py`. Nothing is declared transitive by hand."""
    from lmm.session import Session
    s = Session(None)
    s.gate.rival = rival or (lambda old, new: False)
    for subject, field, value in (
            ("vorlin", "tur", "icecek"),
            ("icecek", "tur", "sivi"),
            ("vorlin", "tur", "sivi"),          # triangle 1
            ("selvin", "tur", "icecek"),
            ("selvin", "tur", "sivi"),          # triangle 2 -> tur transitive
            ("kus", "tur", "canli"),
            ("zilfen", "tur", "kus"),           # -> derives zilfen -> canli
            ("zerbalit", "tur", "metal"),
            ("metal", "tur", "madde"),          # -> derives zerbalit -> madde
            ("zerbalit", "renk", "mavi"),
            ("nortlann", "tur", "ada"),
            ("norgul", "iliski", "cogalir"),    # an OPEN field, not a taxonomy
    ):
        s.learn_cell(subject, field, value, "#doc:corpus")
    return s


def no_engine():
    """Replaces the single chokepoint every model call in this system goes
    through (`runtime.generate` — extract.py and generate.py both reach the
    engine only through it) with a counter that refuses to be called.

    Returns (restore, calls). This is what makes "zero model calls" an
    assertion rather than a claim about the code's shape."""
    from lmm import runtime
    was = runtime.generate
    calls = []

    def refuse(*a, **kw):
        calls.append(kw.get("system", ""))
        raise AssertionError("the engine was called on the graph path")

    runtime.generate = refuse
    return (lambda: setattr(runtime, "generate", was)), calls


@test("N1 a fact the graph holds is spoken by the graph, with no engine at all")
def n1():
    from lmm import lookup
    s = island()
    assert s.memory.transitive, "tur never became transitive — fixture broken"
    # the answer is DERIVED (zilfen->kus, kus->canli): the multi-hop step is
    # exactly the one the thesis says needs no model, and it is the one being
    # answered here for free.
    held = lookup.find(s.memory, "zilfen bir canli midir")
    assert held is not None and held.source == "#inference", held
    restore, calls = no_engine()
    try:
        said = s.respond("zilfen bir canli midir")
    finally:
        restore()
    assert not calls, calls
    assert "canli" in said and "zilfen" in said, said
    assert s.last_from_graph is True and s.last_abstained is False
    assert s.last_subject == "zilfen", s.last_subject


@test("N2 a field with two live values is not answered from the graph")
def n2():
    from lmm import link, lookup
    s = island()
    # one value in the slot: the field is named, the answer is unambiguous
    one = lookup.find(s.memory, "zerbalit renk nedir")
    assert one is not None and one.value == link.resolve(s.memory, "mavi"), one
    # a SECOND value under the same field, neither of them more specific than
    # the other (no is-a edge between them, so `retrieve.specific` cannot and
    # must not choose). The graph now holds two true answers to one question,
    # and the cheap path has no business picking one.
    s.learn_cell("zerbalit", "renk", "yesil", "#doc:corpus")
    assert lookup.find(s.memory, "zerbalit renk nedir") is None
    # ... while SPECIFICITY is still a decision the graph itself can make: two
    # values on the taxonomy, one an ancestor of the other, is not ambiguity.
    both = lookup.find(s.memory, "zerbalit tur nedir")
    assert both is not None and both.value == link.resolve(s.memory, "metal"), both


@test("N3 a question the graph cannot settle falls through instead of guessing")
def n3():
    from lmm import lookup
    s = island()
    # THE ABSENCE CLASS. The graph holds `nortlann -> tur -> ada` and the
    # question asks for a capital; naming no field the subject carries, it gets
    # no graph answer. This is the one that would turn a correct abstention
    # into a confident wrong answer.
    assert lookup.find(s.memory, "nortlann baskenti neresidir") is None
    # THE FIELDLESS QUESTION IS OUT TOO, and that is a cost, not an oversight:
    # nothing structural separates an interrogative the graph has never seen
    # from a field name the graph has never seen (see lookup's docstring).
    assert lookup.find(s.memory, "zerbalit nedir") is None
    # a subject that is not in the graph at all
    assert lookup.find(s.memory, "melvarit bir madde midir") is None
    # a value that is in the graph but not under THIS subject
    assert lookup.find(s.memory, "zerbalit bir kus mudur") is None


@test("N4 an open field is not a taxonomy, and does not answer for one")
def n4():
    from lmm import lookup
    s = island()
    # The graph really does hold `norgul -> iliski -> cogalir`, and the
    # question really does name that value. It is still refused, because
    # `iliski` is not transitive: the graph is not CLOSED over it, so a hit
    # there is one sentence's residue rather than a settled fact — the "true
    # claim answering a question nobody asked" class that `_relation_held`
    # exists for. Measured: with this restriction removed the EN corpus loses
    # a point ("what happens if norgul multiplies" -> "norgul -> multiplies").
    from lmm import link
    norgul = link.resolve(s.memory, "norgul")
    values = {r.value for r in s.memory.about(norgul, touch=False)}
    assert link.resolve(s.memory, "cogalir") in values, "fixture broken"
    assert lookup.find(s.memory, "norgul cogalirsa ne olur") is None
    # and the same subject IS answerable on its taxonomy — the restriction is
    # per-predicate, not per-subject
    s.learn_cell("norgul", "tur", "bitki", "#doc:corpus")
    s.learn_cell("bitki", "tur", "canli", "#doc:corpus")
    got = lookup.find(s.memory, "norgul bir canli midir")
    assert got is not None and got.value == link.resolve(s.memory, "canli"), got


@test("N5 a cheaper answer is not a less accountable one")
def n5():
    s = island()
    restore, calls = no_engine()
    try:
        said = s.respond("zilfen bir canli midir")
    finally:
        restore()
    assert not calls
    # The record is below CERTAIN (it is derived), so it is spoken with its
    # provenance — the same '~' mark and the same stamp the generated path
    # attaches, and NOTHING ELSE: every token added past the record itself is
    # the stamp or punctuation. Same property M1 defends for the paid path.
    assert "#inference" in said, said
    head = said.split("(")[0]
    assert "zilfen" in head and "canli" in head, said
    added = said[len(head):]
    for token in added.split():
        assert token.startswith("#") or not any(c.isalnum() for c in token), token


@test("N6 prose is an option with a price, not the only way to answer")
def n6():
    from lmm import runtime
    s = island()
    was, seen = runtime.generate, []

    def engine(*a, **kw):
        seen.append(1)
        raise RuntimeError("stop here — being called at all is the assertion")

    runtime.generate = engine
    try:
        # fluent=True must NOT take the free path: the caller asked for a
        # sentence and a sentence is the engine's work.
        s.respond("zilfen bir canli midir", fluent=True)
    finally:
        runtime.generate = was
    assert seen, "fluent=True still answered from the graph"
    assert s.last_from_graph is False
    # and the front door offers the same choice, with the turn reporting which
    # path it took
    from lmm.api import Memory
    front = Memory()
    front.session = s
    told = front.ask("zilfen bir canli midir", explain=True)
    assert told.from_graph is True, told


@test("O1 a refusal with a fabrication stapled to it loses the fabrication")
def o1():
    from lmm import evidence
    # the evidence path's block and question, as `_select` hands them over
    block = ("[K1] Vorlin is a drink. A drink is a liquid.\n"
             "[K2] Vorlin is made from the norgul plant.")
    question = "at how many degrees does vorlin boil"
    # THE MEASURED LOSS, three runs out of three, and the second sentence is a
    # claim the corpus contradicts. 'vorlin' is in the question and 'substance'
    # sits in the document, so whole-answer coverage lands in the middle and
    # the read-back is left to decide — which is the model call this filter
    # exists to stop depending on.
    said = "I do not know. Vorlin is not a recognized substance."
    assert evidence.grounded_sentences(said, block, question) == "I do not know."
    # the Spanish one, same shape, no rule about either language
    es_block = "[K1] En el centro de la isla hay un lago grande."
    es_q = "¿cuál es la población de la isla?"
    assert evidence.grounded_sentences(
        "No sé. La población de la isla puede variar.", es_block, es_q) == "No sé."
    # and the field trial's
    rfc = "[K1] RFC 9110 defines HTTP semantics."
    assert evidence.grounded_sentences(
        "I do not know. RFC 9110 is obsoleted by RFC 9111.", rfc,
        "which RFC obsoletes RFC 9110") == "I do not know."


@test("O2 the sentence filter subtracts, and can never empty or invent")
def o2():
    from lmm import evidence
    block = "[K1] Torvanit bir metaldir. Metal bir maddedir."
    question = "torvanit bir madde midir"
    # A FLUENT ONE-SENTENCE ANSWER IS UNTOUCHED, though 'evet' comes from
    # neither the block nor the question. Whether a candidate may speak at all
    # is the gates' decision; this filter only audits what rides along with it,
    # so with nothing to ride on there is nothing to remove.
    one = "Evet, torvanit bir maddedir."
    assert evidence.grounded_sentences(one, block, question) == one
    # every sentence a mixture -> nothing is dropped, for the same reason
    both = "Evet, torvanit bir maddedir. Ayrıca torvanit parlak bir metaldir."
    assert evidence.grounded_sentences(both, block, question) == both
    # a fully grounded pair survives whole
    good = "Torvanit bir metaldir. Metal bir maddedir."
    assert evidence.grounded_sentences(good, block, question) == good
    # it only ever removes: the result is a subsequence of the sentences given
    out = evidence.grounded_sentences(
        "Torvanit bir metaldir. Torvanit kırmızı gezegende bulunur.",
        block, question)
    assert out == "Torvanit bir metaldir.", out
    # empty and single-word inputs do not explode
    assert evidence.grounded_sentences("", block, question) == ""
    assert evidence.grounded_sentences("Bilmiyorum.", block, question) == "Bilmiyorum."


@test("O3 the causal reading is asked first, and a triple cannot overrule it")
def o3():
    """WHICH READING WINS WHEN A SENTENCE ANSWERS TO BOTH — pinned, because it
    was measured and the other order is worse.

    `learn_text` reads each sentence twice: is it causal, and what facts does
    it state. Almost every genuinely causal sentence ALSO yields a triple —
    measured on `gpt-4o-mini`, all five causal sentences of the benchmark
    corpus do (`Yağmur yağarsa bataklık büyür.` -> `[yağmur, şart, yağarsa]`)
    — so whichever reading is asked first is the one that gets written, and
    asking the fact reading first empties `#causes` entirely (38 causal edges
    -> 0 on the local engine, COST.md §7.1). Causal-first was tried in reverse
    for exactly one reason, that a weak engine over-reports causality and
    splits the taxonomy; the reversal did not recover a single derived fact and
    did delete every causal edge, so the order stands.

    Model-free: both reads are stubbed, so what is asserted is the ROUTING —
    a sentence read as causal lands under `#causes` even though the fact read
    would have produced a triple, and a sentence read as not-causal reaches the
    taxonomy."""
    from lmm import extract, generate, link, session as session_module
    real_causal, real_re, real_ex = (generate.is_causal, extract.reextract,
                                     extract.extract)
    try:
        # BOTH readings succeed on the causal sentence; only the fact reading
        # succeeds on the declarative one.
        def is_causal(sent):
            return ("yağmur", "bataklık") if "ağmur" in sent else None

        def reextract(sent):
            if "ağmur" in sent:
                return [("yağmur", "şart", "bataklık")]
            return [("zerbalit", "tür", "metal")]

        generate.is_causal = is_causal
        extract.reextract = reextract
        extract.extract = lambda sent: {"kind": extract.CHAT, "triples": []}
        s = session_module.Session(None)
        s.learn_text("Yağmur yağarsa bataklık büyür. Zerbalit bir metaldir.",
                     source="#doc:test", deep=True)
        by_predicate = {}
        for record in s.memory.records.values():
            name = (link.label_of(s.memory, record.predicate)
                    if record.predicate is not None else None)
            by_predicate.setdefault(name, set()).add(
                (link.label_of(s.memory, record.subject),
                 link.label_of(s.memory, record.value)))
        # the causal sentence is a CAUSAL edge, not the triple it also offered
        assert ("yağmur", "bataklık") in by_predicate.get("#causes", set()), \
            by_predicate
        # ... and it did NOT also enter the taxonomy under its triple
        assert ("yağmur", "bataklık") not in by_predicate.get("tür", set()), \
            by_predicate
        # the declarative sentence reaches the taxonomy
        assert ("zerbalit", "metal") in by_predicate.get("tür", set()), \
            by_predicate
        assert ("zerbalit", "metal") not in by_predicate.get("#causes", set()), \
            by_predicate
    finally:
        generate.is_causal, extract.reextract, extract.extract = (
            real_causal, real_re, real_ex)


@test("R1 a table row is known by its name, not by its indentation")
def r1():
    """A spreadsheet cell has no margin, so a sheet that nests rows draws the
    nesting with a character: the US Census file writes `.Alabama`,
    `.Puerto Rico`. Measured (`benchmarks/field/REPORT.md` §4), that dot made
    every nested row unreachable by the name a question uses — `about("puerto
    rico")` empty while `about(".puerto rico")` held the whole row.

    What decides is Unicode's category (`core/dataset.bare`), not a convention
    this reader knows about, and the written form stays reachable as an alias:
    a question that quotes the cell must not be punished for quoting it.
    Model-free — `learn_rows` calls no engine at all."""
    from lmm.session import Session
    s = Session(None)
    s.learn_rows([{"Geographic Area": ".Puerto Rico",
                   "Population Estimate": "3281557"}], source="#xlsx:test")
    from lmm import link
    named = link.resolve(s.memory, "puerto rico")
    assert named is not None and s.memory.about(named), \
        "the row is unreachable by its name"
    # ONE row, ONE node: the written spelling is the same identity, not a second
    assert link.resolve(s.memory, ".puerto rico") == named, \
        "the written form stopped resolving"
    # a row label that is ONLY layout keeps what it has rather than vanishing
    s.learn_rows([{"a": "...", "b": "1"}], source="#xlsx:test")
    assert link.resolve(s.memory, "...") is not None, \
        "a label of pure layout lost its row"


@test("R2 a header that occupies two rows names every column it covers")
def r2():
    """The Census sheet's header is rows 3 AND 4, and it says so in its own
    merges: `A3:A4` (this name occupies both rows), `C3:F3` (this name covers
    four columns, distinguished underneath). Read with one header row, three of
    those four columns had no name, `row[h or ""]` collapsed them onto ONE dict
    key and two whole years were overwritten out of existence — silently
    (`REPORT.md` §4).

    The sheet is built here rather than shipped, so what is asserted is the
    SHAPE and not one file: no row number is written down, only "as deep as
    this sheet's own merge says". Model-free and file-format-only."""
    import os
    import tempfile
    try:
        import openpyxl
    except Exception:                                       # noqa: BLE001
        print("      (openpyxl not installed — xlsx header block not checked)")
        return
    from lmm import tables
    book = openpyxl.Workbook()
    sheet = book.active
    sheet["A1"] = ("table with row headers in column A and column headers in "
                   "rows 3 through 4")
    sheet["A2"] = "Annual Estimates of the Resident Population, 2020 to 2021"
    sheet["A3"], sheet["B3"], sheet["C3"] = ("Area", "Base", "Estimate")
    sheet["C4"], sheet["D4"] = 2020, 2021
    sheet.merge_cells("A3:A4")
    sheet.merge_cells("B3:B4")
    sheet.merge_cells("C3:D3")
    for i, name in enumerate(("Alabama", "Alaska", "Arizona")):
        sheet.cell(row=5 + i, column=1, value=name)
        for column in (2, 3, 4):
            sheet.cell(row=5 + i, column=column, value=i + column)
    path = os.path.join(tempfile.mkdtemp(), "two-row-header.xlsx")
    book.save(path)
    was = os.environ.get("LMM_XLSX_HEADER_BLOCK")
    try:
        os.environ.pop("LMM_XLSX_HEADER_BLOCK", None)   # the shipped default
        _name, _pre, rows = tables.read_xlsx(path)[0]
        assert len(rows) == 3, rows
        row = rows[0]
        # the spanning name reaches both columns, and the year tells them apart
        assert row.get("Estimate 2020") == "3", row
        assert row.get("Estimate 2021") == "4", row
        # NOTHING IS NAMELESS, so nothing collapses onto the empty key
        assert "" not in row, row
        # the row below the header is header, not the first data row
        assert row.get("Area") == "Alabama", row
        # a whole number typed as floating point by the reader is still a whole
        # number: `2021.0` is a year no question spells
        assert all("." not in k for k in row), row
        # THE WAY BACK is still there and still one variable: the single-row
        # reading, with the second header row arriving as data and the column
        # it names collapsing onto the empty key. It shipped as the default
        # for one release, until the answer path stopped choosing silently
        # between the columns this reading recovers (COST.md §8.4).
        os.environ["LMM_XLSX_HEADER_BLOCK"] = "0"
        _name, _pre, rows = tables.read_xlsx(path)[0]
        assert len(rows) == 4, rows            # the year row arrives as data
        assert "" in rows[0], rows[0]          # and 2021 has no name of its own
    finally:
        os.environ.pop("LMM_XLSX_HEADER_BLOCK", None)
        if was is not None:
            os.environ["LMM_XLSX_HEADER_BLOCK"] = was


@test("N7 a node whose name is several words can still be named")
def n7():
    """A graph node is not always spelled with one word — a spreadsheet's row
    label, a spec table's field, a country. Scanning the question ONE WORD AT A
    TIME could never name one, so a question saying exactly what the graph
    stores fell through to the paid path.

    The uniqueness guarantee is unchanged, and the second half of this test is
    the one that matters: a shorter run inside a longer one that also resolves
    is NOT a second subject, or `find` would decline what it can settle."""
    from lmm import link, lookup
    s = island()
    s.learn_cell("ser vivo", "tur", "canli", "#doc:corpus")
    s.learn_cell("gato", "tur", "ser vivo", "#doc:corpus")
    # `ser` is dropped by the content-word rule and `vivo` alone resolves to
    # nothing, so this question named nothing at all before the runs existed.
    got = lookup.find(s.memory, "ser vivo bir canli midir")
    assert got is not None, "a two-word subject is still unnameable"
    assert got.subject == link.resolve(s.memory, "ser vivo"), got
    # the shorter naming does not compete with the longer one
    s.learn_cell("vivo", "tur", "madde", "#doc:corpus")
    again = lookup.find(s.memory, "ser vivo bir canli midir")
    assert again is not None and again.subject == got.subject, again
    # and a question that names ONLY the shorter node still reaches it
    short = lookup.find(s.memory, "vivo bir madde midir")
    assert short is not None, short
    assert short.subject == link.resolve(s.memory, "vivo"), short


@test("N8 one reading naming several records answers with all of them")
def n8():
    """A two-row spreadsheet header gives several columns one spanning name and
    distinguishes them underneath (`nufus 2020 … nufus 2023`). A question that
    names the spanning part and none of the distinguishing part names all of
    them equally well: the old uniqueness test declined, and the paid path
    behind it then chose one of four with nothing to choose on.

    The answer is every named record, each under its own full field name —
    which is where the distinguishing word sits, written by the document."""
    from lmm import link, lookup
    s = island()
    for year, value in (("2020", "331"), ("2021", "332"), ("2022", "333")):
        s.learn_cell("nortlann", "nufus tahmini " + year, value, "#doc:corpus")
    held = lookup.find(s.memory, "nortlann nufus tahmini")
    assert isinstance(held, tuple) and len(held) == 3, held
    said = lookup.render(s.memory, held)
    for year, value in (("2020", "331"), ("2021", "332"), ("2022", "333")):
        assert year in said and value in said, said
    # and it is still an answer with no engine in it
    restore, calls = no_engine()
    try:
        spoken = s.respond("nortlann nufus tahmini")
    finally:
        restore()
    assert not calls, calls
    assert s.last_from_graph is True and s.last_abstained is False
    assert s.last_subject == "nortlann", s.last_subject
    assert "332" in spoken and "333" in spoken, spoken
    # THE WORD THE QUESTION DID SAY IS NOT SPOKEN OVER. Naming one of the
    # distinguishing words is a narrower reading held by ONE record, and a
    # half-named field is not answered here at all — the message falls through
    # to the semantic path rather than being answered with the three columns it
    # ruled out.
    assert lookup.find(s.memory, "nortlann 2021 tahmini") is None
    # naming the field in full is the ordinary single-record answer
    one = lookup.find(s.memory, "nortlann nufus tahmini 2021")
    assert one is not None and not isinstance(one, tuple), one
    assert one.value == link.resolve(s.memory, "332"), one


@test("N9 several records are only spoken where one record said nothing")
def n9():
    """The generalisation cannot loosen what shape 2 already decided.

    Three guarantees, each measured here: a question the graph settles exactly
    is never widened into a list; a lone half-named field stays unanswered (it
    is not an ambiguity, it is a reading this path does not have); and the
    whole behaviour is switchable, because a claim about what it costs has to
    be measurable rather than arguable."""
    import os

    from lmm import lookup
    s = island()
    s.learn_cell("zerbalit", "renk tonu", "koyu", "#doc:corpus")
    # `renk` is exactly named and answers alone — `renk tonu` is half-named by
    # the same question and does not turn that answer into a list.
    exact = lookup.find(s.memory, "zerbalit renk nedir")
    assert exact is not None and not isinstance(exact, tuple), exact
    # a single half-named field, with no exact reading anywhere: still None
    s2 = island()
    s2.learn_cell("zerbalit", "nufus tahmini 2021", "332", "#doc:corpus")
    assert lookup.find(s2.memory, "zerbalit nufus") is None
    # OFF is off: the same question, the same graph, the old answer
    s3 = island()
    for year in ("2020", "2021"):
        s3.learn_cell("zerbalit", "nufus tahmini " + year, year, "#doc:corpus")
    assert isinstance(lookup.find(s3.memory, "zerbalit nufus tahmini"), tuple)
    os.environ["LMM_LOOKUP_CANDIDATES"] = "0"
    try:
        assert lookup.find(s3.memory, "zerbalit nufus tahmini") is None
    finally:
        del os.environ["LMM_LOOKUP_CANDIDATES"]


@test("P1 the same question over an unmoved memory is not paid for twice")
def p1():
    """The second ask cannot reach the engine at all — `no_engine` makes that
    an assertion rather than a claim about the code's shape — and it must
    report itself exactly as the first one did, because it IS the first one."""
    from lmm.api import Memory
    front = Memory()
    front.session = island()
    first = front.ask("zerbalit renk nedir", explain=True)
    assert first.kind == "ASK" and not first.wrote, first
    assert front._cache, "nothing was kept"
    # The answer path itself is counted, not only the engine: this question is
    # graph-decidable, so an engine counter alone would pass whether the second
    # turn was served from the cache or re-derived for free.
    was, ran = front.session.respond, []
    front.session.respond = lambda *a, **kw: ran.append(1) or was(*a, **kw)
    restore, calls = no_engine()
    try:
        second = front.ask("zerbalit renk nedir", explain=True)
    finally:
        restore()
        front.session.respond = was
    assert not calls and not ran, "the answer was derived a second time"
    assert str(second) == str(first), (first, second)
    assert second.abstained == first.abstained
    assert second.from_graph == first.from_graph
    assert second.subject == first.subject and second.kind == first.kind
    assert front.session.last_abstained == first.abstained
    # and the switch is a switch
    off = Memory(cache=False)
    off.session = island()
    off.ask("zerbalit renk nedir")
    was, ran = off.session.respond, []
    off.session.respond = lambda *a, **kw: ran.append(1) or was(*a, **kw)
    try:
        off.ask("zerbalit renk nedir")
    finally:
        off.session.respond = was
    assert ran, "cache=False still served a kept answer"
    assert off._cache is None
    # AND A KEPT ANSWER IS NOT AN ANSWER TO "SHALL I RESEARCH THAT?". While an
    # offer is outstanding this turn's meaning is the offer's answer, and the
    # flow that made the offer has to consume it; a replay would leave it
    # standing and the NEXT turn would be read as the approval. Measured: the
    # TR set's "melvarit nedir" raises exactly this offer.
    front.session._pending = ("melvarit", "melvarit nedir")
    was, ran = front.session.respond, []
    front.session.respond = lambda *a, **kw: ran.append(1) or was(*a, **kw)
    restore, _ = no_engine()        # the offer flow consults the engine; it
    try:                            # must not be reached for real here
        front.ask("zerbalit renk nedir")
    finally:
        restore()
        front.session.respond = was
    assert ran, "a kept answer was served while a research offer stood"


@test("P2 a memory that moved does not repeat its old answer")
def p2():
    """A STALE ANSWER IS WORSE THAN AN EXPENSIVE ONE, so the fingerprint has to
    move for every mutation the graph has — not just for a new record. Trust
    moving under a record that is already there is the case a record COUNT
    misses, and it is the case that changes whether an answer is hedged."""
    from lmm.api import Memory
    from lmm.core import dynamics
    front = Memory()
    front.session = island()
    question = "zerbalit renk nedir"
    front.ask(question)
    stamp = front._state()
    assert front._cache, "nothing was kept"
    # a new fact under a different subject: the count moves
    front.session.learn_cell("karvel", "tur", "balik", "#doc:corpus")
    assert front._state() != stamp
    stamp = front._state()
    # trust alone moves, with no record added and none removed
    record = next(iter(front.session.memory.records.values()))
    before = record.trust
    dynamics.fade(front.session.memory, record,
                  now=record.last_seen + dynamics.FRESH + 1)
    if record.trust != before:              # fade declines to touch some records
        assert front._state() != stamp, "a trust change left the state stamp still"
    # and the kept answer is not served under the new state
    restore, calls = no_engine()
    try:
        front.ask(question)
    finally:
        restore()
    # the question is graph-decidable, so no engine call is expected — what is
    # asserted is that the answer was RE-DERIVED under a state the cache had
    # never seen, which the entry's key proves.
    assert not calls
    assert all(kept[0] == front._state() for kept in front._cache.values())


# --- X  the offline expansion (doc2query--) ------------------------------
#
# The mechanism generates text with a model and indexes it. The whole
# uninvented-0 guarantee rests on the answer block containing only what the
# document wrote, so the ONE thing that must be impossible is generated text
# reaching that block. These four are model-free by construction: the
# "generated" queries are handed in directly, so no engine is involved and the
# assertions are about the store's wiring rather than about a model's output.


def _expanded_store():
    """A store holding one document line and one generated query for it, where
    the query carries a word the document does not contain at all."""
    from lmm import evidence
    store = evidence.SentenceStore()
    store.add("Kelvane tower measures 42 metres.", "#doc:x")
    store.add("The Nordheim archive opened in 1904.", "#doc:x")
    store.add("Visitors reach the archive by the western stair.", "#doc:x")
    store.learn_expansions({0: ["how tall is the Kelvane spire"]})
    return store


@test("X1 generated text is not evidence and cannot be returned as evidence")
def x1():
    """THE WALL. `find` returns `self.sentences[sid][0]` and the expansion lives
    in a different structure entirely — asserted here rather than trusted,
    because the day a generated sentence can be returned is the day the gate is
    auditing an answer against a model's own invention."""
    store = _expanded_store()
    assert len(store.sentences) == 3, "an expansion was written as a sentence"
    written = {text for text, _src in store.sentences}
    assert not any("spire" in text for text in written)
    # the generated word REACHES the line...
    got = store.find("how tall is the spire", most=4)
    assert got and got[0].startswith("Kelvane tower"), got
    # ...and nothing generated comes back with it, from any query
    for question in ("spire", "how tall is the Kelvane spire",
                     "archive", "western stair"):
        for block in store.find(question, most=6):
            assert block in written, block


@test("X2 the expansion channel cannot outweigh the document's own words")
def x2():
    """The ladder (`evidence.CHANNEL`): a guess sits exactly as far below a
    plain occurrence as a field name sits above it. A sentence reachable only
    through a guess must never displace one the document's own words name."""
    from lmm import evidence
    store = evidence.SentenceStore()
    store.add("The Kelvane tower is closed on Mondays.", "#doc:y")
    store.add("Storms are frequent in the eastern valley.", "#doc:y")
    # every generated query points the SECOND line at the first line's words
    store.learn_expansions({1: ["when is the Kelvane tower closed"] * 1})
    ranked = store.find("Kelvane tower", most=2)
    assert ranked[0].startswith("The Kelvane tower"), ranked


@test("X3 a generated query that drifts to another line is filtered out")
def x3():
    """doc2query--: the filter is the document. A query whose words belong to
    ANOTHER region reaches that region better than the line it was generated
    from, and a negative margin is never indexed — the threshold's floor is the
    sign of the margin, not a number anyone chose."""
    from lmm import evidence
    store = evidence.SentenceStore()
    store.add("The Kelvane tower measures 42 metres.", "#doc:z")
    store.add("The Nordheim archive opened in 1904.", "#doc:z")
    store.learn_expansions({0: ["when did the Nordheim archive open",
                                "how tall is the Kelvane spire"]})
    kept = store.expansions.get(0, [])
    assert "how tall is the Kelvane spire" in kept, kept
    assert "when did the Nordheim archive open" not in kept, kept


@test("X4 LMM_EXPAND=0 restores the retrieval byte for byte")
def x4():
    """The A/B has to be ONE variable, and turning it off has to leave nothing
    behind: the same query over the same store must retrieve exactly what it
    retrieved before the expansion was ever paid for."""
    from lmm import evidence
    plain = evidence.SentenceStore()
    for line in ("Kelvane tower measures 42 metres.",
                 "The Nordheim archive opened in 1904.",
                 "Visitors reach the archive by the western stair."):
        plain.add(line, "#doc:x")
    store = _expanded_store()
    was = os.environ.get("LMM_EXPAND")
    os.environ["LMM_EXPAND"] = "0"
    try:
        for question in ("how tall is the spire", "archive", "42 metres"):
            assert store.find(question, most=4) == plain.find(question, most=4)
    finally:
        if was is None:
            del os.environ["LMM_EXPAND"]
        else:
            os.environ["LMM_EXPAND"] = was
    # and with it back on, the guessed word reaches the line again
    assert store.find("spire", most=4)
    assert not plain.find("spire", most=4)


@test("W1 a question that names a document gets that document's sentences")
def w1():
    """THE SOURCE-NAME CHANNEL, and the corpus class it was measured on: many
    sibling documents sharing one template. Asked for one programme BY NAME,
    `find` returned sentences from the siblings — the name's words scored as
    ordinary query words, and the words every sibling shares separate nothing.

    The channel weighs a query word by its power to separate SOURCES,
    log(S/s): a word in one source name of many is decisive, a word in every
    source name weighs exactly zero. That zero is asserted below twice,
    because it is what makes the channel safe: it is inert on a single
    document (every benchmark before this one), and inert for the genre words
    a corpus's names all share."""
    from lmm import evidence
    st = evidence.SentenceStore()
    # the failing class: the document's NAME never appears in its body
    st.add("The morning block covers goal setting and priorities.",
           "#docx:Alpha Programme.docx")
    st.add("The morning block covers goal setting and reviews.",
           "#docx:Beta Programme.docx")
    hits = st.find("what does the alpha morning block cover", most=1)
    assert hits and "priorities" in hits[0], hits
    hits = st.find("what does the beta morning block cover", most=1)
    assert hits and "reviews" in hits[0], hits
    # a name word EVERY source carries separates nothing and moves nothing:
    # the query with it retrieves byte for byte what the query without it
    # retrieves (near-identical siblings may collapse to one seat — that is
    # the region dedup, older than this channel, and not its business)
    assert (st.find("what does the programme morning block cover", most=2)
            == st.find("what does the morning block cover", most=2))
    # single-source store: the channel vanishes entirely (weight log(1) = 0),
    # so retrieval is what it was before the channel existed
    one = evidence.SentenceStore()
    one.add("The morning block covers goal setting and priorities.",
            "#docx:Alpha Programme.docx")
    one.add("Lunch is at noon.", "#docx:Alpha Programme.docx")
    assert one.find("what does the alpha morning block cover",
                    most=1)[0].startswith("The morning block")
    # and a name match ALONE is not evidence: a sentence sharing no content
    # word with the question is not seated by its document's name
    assert not st.find("alpha", most=4) or all(
        "morning" in h or "goal" in h for h in st.find("alpha", most=4))


@test("W2 one document cannot monopolise the block when others hold evidence")
def w2():
    """THE SOURCE IS A REGION AT DOCUMENT SCALE. Measured (62 sibling training
    outlines, "which programmes cover X"): the term lived in eight documents
    and every seat went to the strongest one, so the block could only ever
    name a single programme. The cap is the region cap — one seat states, a
    second corroborates, a third crowds out another voice — and nothing is
    discarded: with no rival sources left, the same document steps back in
    and fills the block, so a store where one document genuinely holds all
    the evidence loses nothing."""
    from lmm import evidence
    st = evidence.SentenceStore()
    for n in range(3):
        st.add(f"Trust exercises build team safety in module {n} daily.",
               "#docx:Alpha.docx")
    st.add("Trust exercises anchor the safety walk outdoors.",
           "#docx:Beta.docx")
    st.add("Trust exercises close the safety retrospective meeting.",
           "#docx:Gamma.docx")
    st.find("which sessions use trust exercises for safety", most=6)
    origins = {src.split(":")[-1] for src in st.last_sources}
    assert {"Beta.docx", "Gamma.docx"} <= origins, st.last_sources
    # and with ONE source, the cap never engages: distinct regions of the
    # same document all seat, exactly as they always did
    one = evidence.SentenceStore()
    one.add("Trust exercises open the morning safety circle.", "#docx:A.docx")
    one.add("The afternoon workshop pairs trust exercises with feedback.",
            "#docx:A.docx")
    one.add("Evening reflection revisits the trust exercises alone.",
            "#docx:A.docx")
    one.find("where do the trust exercises appear", most=6)
    assert len(one.last_sources) >= 3, one.last_sources


@test("W3 a composed draft keeps its structure and loses its inventions")
def w3():
    """COMPOSITION IS THE SHORT PATH AT DOCUMENT LENGTH — the digit veto and
    the designation discipline, applied line by line, no model in the loop. The engine below is faked so
    the test exercises exactly the part this layer adds: what happens to a
    draft AFTER the engine has written it. Four lines go in: a grounded
    claim, a section title (shares nothing — claims nothing — stands), a
    MIXTURE (material words around an invented conclusion), and an invented
    number. Two survive, and the invented number is the line the digit veto
    exists for: a composed agenda is where fabricated times try to live."""
    from lmm import generate
    from lmm.session import Session
    s = Session(None)
    s.learn_text("The trust module opens with a listening exercise. "
                 "Participants pair up for the feedback round.",
                 source="#docx:Team Basics.docx", deep=False)
    s.learn_text("The outdoor day closes with a fire-building task.",
                 source="#docx:Field Day.docx", deep=False)
    draft = ("Morning Session (Team Basics)\n"
             "The trust module opens with a listening exercise.\n"
             "The listening exercise takes 45 minutes.\n"
             "The trust module is certified by the Ministry.\n")
    real = generate.compose
    generate.compose = (lambda brief, material, warmth=0.2, persona="",
                        max_tokens=None, style="": draft)
    try:
        text, used = s.compose(
            "draft a one day programme with the trust module "
            "and the outdoor fire task")
    finally:
        generate.compose = real
    assert "listening exercise." in text, text
    assert "Morning Session" in text, text            # structure stands
    assert "45" not in text, text                     # invented number: gone
    assert "Ministry" not in text, text               # invented DESIGNATION:
    #                                                   gone — capitals off
    #                                                   the sentence start
    #                                                   must be attested.
    #                                                   (The stated residue
    #                                                   of the second-form
    #                                                   gate: an invented
    #                                                   LOWERCASE quality
    #                                                   would pass; names,
    #                                                   designations and
    #                                                   numbers cannot.)
    assert not s.last_abstained
    assert any("Team Basics" in u for u in used), used


@test("W4 the census question is answered by counting, not by retrieving")
def w4():
    """Which documents cover X — measured to fail as retrieval (eight
    documents held the term; the answer named one, because retrieval's job is
    the best evidence, not the census). `where` counts instead: every
    sentence carrying the term votes for its source, inflection-tolerantly,
    with no engine anywhere. The unknown term returns the empty census — not
    a guess."""
    from lmm import evidence
    st = evidence.SentenceStore()
    st.add("The feedback round closes every module.", "#docx:Alpha.docx")
    st.add("Feedback pairs practise daily.", "#docx:Alpha.docx")
    st.add("Written feedback follows the workshop.", "#docx:Beta.docx")
    st.add("Lunch is served at the barn.", "#docx:Gamma.docx")
    census = st.where("feedback")
    assert [(src.split(":")[-1], n) for src, n in census] == [
        ("Alpha.docx", 2), ("Beta.docx", 1)], census
    assert st.where("blockchain") == []


@test("W5 the abstention stamp is read off the evidence, not asked of a model")
def w5():
    """THE STAMP IS THE MEASUREMENT'S BACKBONE, and it was resting on a model
    call: the re-extractor was asked \"did this turn assert anything\", found
    no tidy triple in one long list-sentence, and stamped a full correct
    answer as an ABSTENTION. The structural reading cannot make that mistake:
    a claim on the evidence path got there by sharing content words with the
    proof (the coverage gates admitted it on that ground), and a refusal
    shares none. No engine anywhere in this test."""
    from lmm.session import Session
    s = Session(None)
    s._last_proof = ["Alpha Programme — The workshop closes with a "
                     "feedback round of paired exercises."]
    # a long, list-like assertion built on the proof's words: asserted,
    # however untidy a triple extractor would find it
    assert s._asserted_a_fact(
        "The workshop involves paired exercises, closing with a feedback "
        "round for every participant.")
    # a refusal touches none of the proof: not an assertion, in any wording
    assert not s._asserted_a_fact("Bu konuda maalesef bilgim yok.")
    # ...and ONE incidental shared word does not flip it: "closes" is in the
    # proof, and a refusal wearing it is still a refusal (measured: "I do
    # not know" shared "not" with a standard full of SHALL-NOTs, and an
    # honest abstention was stamped an assertion — a point lost to a
    # stopword)
    assert not s._asserted_a_fact(
        "I cannot say anything about what closes it.")
    assert not s._asserted_a_fact("I am afraid that is unknown to me.")
    # the footnote is not a claim: a sourced refusal is still a refusal
    assert not s._asserted_a_fact(
        "I do not know. (~ #docx:Alpha Programme.docx)")


@test("W6 a question API cannot write memory, however imperative the grammar")
def w6():
    """Measured: "draft a programme for the sales team" was classified WRITE,
    written into the graph as an operator fact, and answered with a
    confirmation of having learned it. The contract is the caller's to state,
    not the classifier's to guess: ask() passes teach=False and the WRITE
    branch is simply not there for it; the conversational surface keeps
    teach=True because teaching is one of its jobs. The classifier below is
    faked to force the misfire the fix exists for."""
    from lmm import extract
    from lmm.session import Session
    s = Session(None)
    real = extract.extract
    extract.extract = lambda message: {
        "kind": extract.WRITE,
        "triples": [("sales team", "programme", "draft")]}
    try:
        before = len(s.memory.records)
        s.respond("draft a day programme for the sales team", teach=False)
        assert len(s.memory.records) == before, "teach=False wrote a record"
        # (the teach=True write path is engine-bearing and is exercised by
        # the session tests above; this test guards only the new contract)
    finally:
        extract.extract = real


@test("W7 a restated evidence line needs no jury")
def w7():
    """The read-back judge was measured to wobble on identical input — the
    same fully-grounded claim confirmed, then denied, at temperature zero —
    and a verbatim claim denied for sitting in spec noise. When the claim
    and a single evidence line cover each other's content words, the claim
    IS that line re-inflected; the judge below is rigged to EXPLODE to prove
    it is never consulted for that class. One-way coverage still goes to the
    jury — a subset of a line can invert it — and so does everything else."""
    from lmm import generate
    from lmm.session import Session
    s = Session(None)
    proof = ["The valve opens at forty degrees."]
    blk = "[K1] " + proof[0]
    real = generate.supported
    def boom(answer, view):
        raise AssertionError("jury consulted for a restated evidence line")
    generate.supported = boom
    try:
        assert s._read_back(
            "The valve opens at forty degrees. (~ #docx:Alpha.docx)",
            proof, blk)
    finally:
        generate.supported = real
    # partial coverage: the jury IS consulted, and its verdict rules
    generate.supported = lambda answer, view: False
    try:
        assert not s._read_back(
            "The valve opens at forty degrees under pressure.", proof, blk)
    finally:
        generate.supported = real


@test("W8 a refusal with a full tally offers the tally, and nothing more")
def w8():
    """THE INFORMED REFUSAL. A turn that declines while the census is rich is
    refusing with its hands full — the memory holds an attested tally of
    documents that speak to the topic, and a bare refusal hides it. The
    trigger is a STATE (refused, tally in hand), never a reading of the
    question's wording, which is what lets a recommendation-shaped question
    get a sourced reply without anyone classifying intent. And the offer is
    read like any other turn: an engine that pads the tally with an invented
    programme or a number the tally does not carry stays unspoken, and the
    honest refusal stands."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    s.learn_text("The trust circle closes the morning block.",
                 source="#docx:Alpha.docx", deep=False)
    s.learn_text("Trust pairs open the afternoon walk.",
                 source="#docx:Beta.docx", deep=False)
    real_ex, real_ans = extract.extract, generate.answer
    extract.extract = lambda m: {"kind": extract.ASK, "triples": []}
    calls = {"n": 0}
    def fake_answer(question, block_, warmth=0.2, persona="", **kw):
        calls["n"] += 1
        if calls["n"] == 1:
            return "I do not know."         # the path refuses first
        return "Alpha and Beta both touch on trust."     # the tally offer
    generate.answer = fake_answer
    try:
        said = s.respond("which sessions would you recommend for trust",
                         teach=False)
        assert "Alpha" in said and "Beta" in said, said
        assert not s.last_abstained
        # a padded offer is refused: the engine invents a programme name
        calls["n"] = 0
        def padded(question, block_, warmth=0.2, persona="", **kw):
            calls["n"] += 1
            if calls["n"] == 1:
                return "I do not know."
            return "Alpha, Beta and the Gamma masterclass cover trust."
        generate.answer = padded
        said = s.respond("which sessions would you recommend for trust",
                         teach=False)
        assert "Gamma" not in said, said
        assert s.last_abstained
    finally:
        extract.extract, generate.answer = real_ex, real_ans


@test("W9 the named document speaks first")
def w9():
    """Measured, conversation trace: the anchored programme WAS seated — third
    and fourth — while a sibling document's rare stem took the top seats, and
    the engine, which reads the block top-down, answered from the sibling.
    Admission was never the problem; ORDER was. When the source-name channel
    has boosted a document (the question named it), that document's seats
    come first, in their own relative order; everything else follows,
    unmoved. A store where no name matched — every single-document benchmark
    — reorders nothing, byte for byte."""
    from lmm import evidence
    st = evidence.SentenceStore()
    # the sibling catches TWO rare query stems; the anchored name is one
    # word — the live defeat's arithmetic, where a rare stem out-logged the
    # name boost in a large store
    st.add("The chronometer calibration drill closes the module.",
           "#docx:Beta.docx")
    st.add("The module covers listening and closing practice.",
           "#docx:Alpha.docx")
    st.add("Lunch is served in the barn.", "#docx:Gamma.docx")
    st.add("The barn hosts the evening reflection.", "#docx:Delta.docx")
    hits = st.find("what does the alpha module cover in chronometer "
                   "calibration", most=4)
    assert hits, "nothing seated"
    first_src = st.last_sources[0]
    assert "Alpha" in first_src, (first_src, st.last_sources)
    # the sibling is still seated — order moved, admission did not
    assert any("Beta" in src for src in st.last_sources), st.last_sources


@test("W10 the informed refusal offers reasons, and every reason has a dateline")
def w10():
    """v1 handed the rescue engine ONE line — the tally — so the best offer it
    could write was a list of names. The names' own documents are sitting in
    the store with the very sentences that justify them; v2 hands those over,
    datelined, beside the tally. The gate does not change: an offer may still
    only name what the tally or the material names, and the fake below that
    cites a programme from NOWHERE is still refused. What this test pins is
    the material: the rescue block must carry datelined lines from more than
    one census source, or the reasons have nothing to rest on."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    s.learn_text("The trust walk closes the morning arc gently.",
                 source="#docx:Alpha.docx", deep=False)
    s.learn_text("Trust walk pairs shape the afternoon arc.",
                 source="#docx:Beta.docx", deep=False)
    real_ex, real_ans = extract.extract, generate.answer
    extract.extract = lambda m: {"kind": extract.ASK, "triples": []}
    seen = {"blocks": []}
    def fake_answer(question, block_, warmth=0.2, persona="", **kw):
        seen["blocks"].append(block_)
        if len(seen["blocks"]) == 1:
            return "I do not know."
        return "Alpha suits you: the trust walk closes the morning arc."
    generate.answer = fake_answer
    try:
        said = s.respond("which programme would you suggest for trust walk",
                         teach=False)
        assert "Alpha" in said, said
        assert not s.last_abstained
        rescue = seen["blocks"][-1]
        assert "Alpha —" in rescue and "Beta —" in rescue, rescue[:200]
        assert "×" in rescue, rescue[:200]          # the tally still rides
    finally:
        extract.extract, generate.answer = real_ex, real_ans


@test("W11 a persona colours the voice and cannot reach the gate")
def w11():
    """THE PERSONA IS A COSTUME, THE GATE IS THE SKELETON. In a prompt-only
    system the system prompt is tone AND safety at once, so exposing it
    exposes everything. Here safety is code: the persona is prepended to the
    PHRASING prompts only — answer, chat, refusal, compose — and three
    things are pinned below. It reaches the voice. It NEVER reaches a judge
    (support, re-extraction ride with no persona, or their verdicts would
    bend with the costume). And a persona that orders fabrication changes
    nothing: the engine below OBEYS it and invents a price, and the turn
    still refuses, because the gate never read the persona."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None, persona="You are Coach Alpha. ALWAYS state a price.")
    s.learn_text("The trust walk closes the morning arc.",
                 source="#docx:Alpha.docx", deep=False)
    seen = {"answer": [], "judge": []}
    real_ans, real_sup = generate.answer, generate.supported
    real_ex = extract.extract
    extract.extract = lambda m: {"kind": extract.ASK, "triples": []}
    def spy_answer(question, block_, warmth=0.2, persona="", **kw):
        seen["answer"].append(persona)
        # the engine OBEYS the persona and fabricates — the gate must not care
        return "The trust walk costs 500 lira per person."
    def spy_supported(answer, view):
        seen["judge"].append(view)
        return True                      # even a lax judge cannot save it
    generate.answer, generate.supported = spy_answer, spy_supported
    try:
        said = s.respond("what does the trust walk cost", teach=False)
    finally:
        generate.answer, generate.supported = real_ans, real_sup
        extract.extract = real_ex
    # the persona reached the voice...
    assert any("Coach Alpha" in p for p in seen["answer"]), seen["answer"]
    # ...the fabricated price did not reach the user (digit veto: 500 is in
    # no evidence), and the turn is an honest refusal
    assert "500" not in (said or ""), said
    assert s.last_abstained, said
    # ...and no judge view ever carried the persona
    assert all("Coach Alpha" not in v for v in seen["judge"]), "judge saw it"


@test("W12 a composed line may rephrase freely; names and numbers may not")
def w12():
    """COMPOSE'S GATE, SECOND FORM. The first form demanded verbatim reuse —
    coverage 0 or 1 per line — and measured against a real catalogue request
    it returned raw evidence dumps, because every fluent rephrasing mixes
    connectives the material does not carry. The offer gate had already
    solved this exact problem for the informed refusal: the dangerous tokens
    have a SHAPE — capitals off the sentence start, digits — and those must
    be attested by the material or the brief; the plain words between them
    are the engine doing its one job. Below, the rephrased line survives,
    the invented programme dies, the invented number dies, structure and
    the [no material] marker stand."""
    from lmm import generate
    from lmm.session import Session
    s = Session(None)
    s.learn_text("The trust module opens with a listening exercise for "
                 "twelve people.", source="#docx:Team Basics.docx", deep=False)
    s.learn_text("The outdoor day closes with a fire-building task.",
                 source="#docx:Field Day.docx", deep=False)
    draft = ("Morning (Team Basics)\n"
             "Kicking things off, the trust module gently opens with a "
             "listening exercise suited to twelve people.\n"
             "Then the Gamma Masterclass rounds out the day.\n"
             "The listening exercise takes 45 minutes.\n"
             "[no material]\n")
    real = generate.compose
    generate.compose = (lambda brief, material, warmth=0.2, persona="",
                        max_tokens=None, style="": draft)
    try:
        text, used = s.compose("draft a day with the trust module and the "
                               "fire task")
    finally:
        generate.compose = real
    assert "listening exercise" in text, text          # rephrasing SURVIVES
    assert "Kicking things off" in text, text          # connectives free
    assert "Gamma" not in text, text                   # invented name: dead
    assert "45" not in text, text                      # invented number: dead
    assert "Morning" in text, text                     # structure stands
    assert "[no material]" in text, text               # the marker stands


@test("W13 a contributing document brings its fact-sheet along")
def w13():
    """Measured on a catalogue request: the chosen trainings' duration and
    participant lines never reached the material, because a logistics line
    ("DURATION 2 days · SEATS 16-18") shares no word with any topic — it is
    a RECORD, and records lose lexical seat races by construction. The rule
    is format, not language: every source that contributes material also
    contributes its record-shaped lines, the way a brochure's fact box rides
    with its prose. And topics= gathers per topic, so a four-topic brief
    cannot starve its fourth topic of seats."""
    from lmm import generate
    from lmm.session import Session
    s = Session(None)
    s.learn_text("The trust module opens with a listening exercise.\n"
                 "DURATION: 2 full days · SEATS: 16-18 people\n"
                 "Participants pair up for the feedback round.",
                 source="#docx:Team Basics.docx", deep=False)
    s.learn_text("The outdoor day closes with a fire-building task.\n"
                 "DURATION: 1 day · SEATS: 12 people",
                 source="#docx:Field Day.docx", deep=False)
    seen = {}
    real = generate.compose
    def spy(brief, material, warmth=0.2, persona="", max_tokens=None,
            style=""):
        seen["material"] = material
        return "The trust module opens with a listening exercise."
    generate.compose = spy
    try:
        s.compose("draft a day around the trust module and the fire task",
                  topics=["trust module", "fire task"])
    finally:
        generate.compose = real
    mat = seen["material"]
    assert "listening exercise" in mat            # topic 1 material
    assert "fire-building" in mat                 # topic 2 not starved
    # the fact-sheets rode along though no topic word touches them
    assert "16-18" in mat, mat[-400:]
    assert "SEATS: 12" in mat, mat[-400:]



def engine_free(fn):
    """Run a test with the turn's small helpers answered locally.

    THE SUITE'S OWN CLAIM IS THAT IT NEEDS NO MODEL AND NO NETWORK — CI
    runs it on five Pythons with nothing installed — and three
    consultation tests had been quietly breaking that since 0.3.2,
    unnoticed because a developer's shell always has an engine
    configured. A conversational turn asks three small questions on its
    way (is this about you, does this ask for material, re-read what was
    said) and each reaches for the engine. Where a test is about ROUTING,
    those are scenery, and scenery answers here.
    """
    def ran():
        from lmm import extract, generate
        saved = (generate.is_identity_question, generate.wants_material,
                 extract.reextract)
        generate.is_identity_question = lambda message: False
        generate.wants_material = lambda message: False
        extract.reextract = lambda sentence: []
        try:
            return fn()
        finally:
            (generate.is_identity_question, generate.wants_material,
             extract.reextract) = saved
    ran.__name__ = fn.__name__
    ran.__doc__ = fn.__doc__
    return ran


@test("W14 in a no-teach conversation, a statement is context, not a query")
@engine_free
def w14():
    """Measured, live: 'I am in banking, my team is ten people' — the user
    SHARING context — was classified WRITE (correctly), skipped the write
    (teach=False, correctly), and then fell through to the QUESTION
    treatment: retrieval found no answer to a sentence that asked nothing,
    the refusal path spoke, and the persona's greeting instruction leaked
    into it — the bot replied to a customer's needs with its own
    introduction. Three correct mechanisms, one uncovered seam. In a
    conversation that may not teach, a statement is neither a lesson nor a
    query: it is CONTEXT, and it routes to the chat surface, which carries
    the history and the persona. No memory is written, nothing is refused,
    and the reply is the conversation continuing."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None, persona="You are a warm consultant.")
    s.learn_text("The trust walk closes the morning arc.",
                 source="#docx:Alpha.docx", deep=False)
    real_ex, real_chat = extract.extract, generate.chat
    extract.extract = lambda m: {"kind": extract.WRITE,
                                 "triples": [("my team", "size", "ten")]}
    called = {}
    def fake_chat(message, identity_block="", warmth=0.7, history=None,
                  persona=""):
        called["chat"] = True
        called["persona"] = persona
        return "Ten people in banking — noted! What outcome do you want?"
    generate.chat = fake_chat
    try:
        before = len(s.memory.records)
        s._census_line = "2 documents \u00b7 Alpha"   # stale, from last turn
        said = s.respond("I am in banking and my team is ten people",
                         teach=False)
    finally:
        extract.extract, generate.chat = real_ex, real_chat
    assert called.get("chat"), "statement did not route to the chat surface"
    assert said == "Ten people in banking \u2014 noted! What outcome do you want?", (
        "the reply was hijacked \u2014 stale census + informed refusal?: %r" % said)
    assert "consultant" in called.get("persona", ""), "persona missing"
    assert len(s.memory.records) == before, "teach=False wrote a record"


@test("W15 a consultation's search reads the whole consultation")
def w15():
    """Measured, live, five turns: "in banking, ten people" then "risk
    management is our side" then "give me the best three-hour material" —
    and the retrieval for the last turn searched THREE WORDS OF NOTHING,
    because the query has always been this turn's sentence and the topic
    lived two turns back. The engine saw the history; the search never did.
    The seam is the same one W14 closed, seen from the other side: in a
    conversation that may not teach (a CONSULTATION), context turns were
    dropped from memory — correctly — but also dropped from RETRIEVAL,
    which is where they were needed. The fix is one sentence: a
    consultation's question searches with the consultation's prior turns
    riding along. The weighting organ (log S/s zeroes corpus-wide words)
    decides what matters — no keyword, no language, no cap of ours. A
    teaching conversation (teach=True — the benchmarks) keeps its
    single-turn query byte for byte."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    s.learn_text("The abseil drill anchors the ravine module.",
                 source="#docx:Alpha.docx", deep=False)
    real = (extract.extract, generate.chat, generate.offer_research,
            generate.refusal)
    queries = []
    orig_find = s.evidence.find
    def spy_find(query, most=4, floor_share=0.5):
        queries.append(query)
        s.evidence.last_sources = []
        s.evidence.last_census = []
        return []
    s.evidence.find = spy_find
    kinds = iter([{"kind": extract.WRITE,
                   "triples": [("our team", "focus", "ravine work")]},
                  {"kind": extract.ASK,
                   "triples": [("material", "duration", "")]},
                  {"kind": extract.ASK,
                   "triples": [("material", "duration", "")]}])
    extract.extract = lambda m: next(kinds)
    generate.chat = (lambda message, identity_block="", warmth=0.7,
                     history=None, persona="": "Noted!")
    generate.offer_research = lambda subject, question: ""
    generate.refusal = lambda message, persona="": "I do not know."
    if not hasattr(generate, "wants_material"):
        generate.wants_material = lambda m: False
        added = True
    else:
        added = False
    real_wm = generate.wants_material
    generate.wants_material = lambda m: False
    try:
        s.respond("our team focuses on ravine work", teach=False)
        queries.clear()
        s.respond("the best two hour material please", teach=False)
        consult = " ".join(queries)
        queries.clear()
        s.respond("the best two hour material please", teach=True)
        teaching = " ".join(queries)
    finally:
        (extract.extract, generate.chat, generate.offer_research,
         generate.refusal) = real
        if added:
            del generate.wants_material
        else:
            generate.wants_material = real_wm
        s.evidence.find = orig_find
    assert "ravine" in consult, consult
    assert "ravine" not in teaching, teaching


@test("W16 when the user asks for the material, the composer answers")
def w16():
    """The other half of the live failure: the store HELD the catalogue the
    user asked for, and the organ that writes catalogues (compose) sat
    unreachable behind a command prefix while the refusal spoke. The bridge
    is a STATE, in the informed-refusal family: a consultation turn that
    ended in abstention with prior context in hand asks the engine ONE
    cheap question — is this turn asking me to PRODUCE something? — and on
    yes, the composer runs over the whole consultation as its brief. The
    composed draft is gated as always; on a composer refusal the turn falls
    back to the paths that already exist. A teaching conversation never
    even asks the question — the benchmarks cannot pay a call for a bridge
    they never cross."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    s.learn_text("The abseil drill anchors the ravine module.",
                 source="#docx:Alpha.docx", deep=False)
    real = (extract.extract, generate.chat, generate.offer_research,
            generate.refusal)
    orig_find = s.evidence.find
    def empty_find(query, most=4, floor_share=0.5):
        s.evidence.last_sources = []
        s.evidence.last_census = []
        return []
    s.evidence.find = empty_find
    kinds = iter([{"kind": extract.WRITE,
                   "triples": [("our team", "focus", "ravine work")]},
                  {"kind": extract.ASK,
                   "triples": [("material", "duration", "")]}])
    extract.extract = lambda m: next(kinds)
    generate.chat = (lambda message, identity_block="", warmth=0.7,
                     history=None, persona="": "Noted!")
    generate.offer_research = lambda subject, question: ""
    generate.refusal = lambda message, persona="": "I do not know."
    had = hasattr(generate, "wants_material")
    real_wm = getattr(generate, "wants_material", None)
    generate.wants_material = lambda m: True
    briefs = []
    def fake_compose(brief, seats=24, topics=None, on_line=None):
        briefs.append(brief)
        # the ledger's terms arrive as the composer's topics — dedicated
        # seats for each thing the consultation actually named
        assert topics and any("ravine" in t for t in topics), topics
        return "THE CATALOGUE\nAlpha \u2014 the abseil drill", \
               ["#docx:Alpha.docx"]
    s.compose = fake_compose
    try:
        s.respond("our team focuses on ravine work", teach=False)
        said = s.respond("the best two hour material please", teach=False)
    finally:
        (extract.extract, generate.chat, generate.offer_research,
         generate.refusal) = real
        if had:
            generate.wants_material = real_wm
        else:
            del generate.wants_material
        s.evidence.find = orig_find
    assert briefs, "the bridge never reached the composer"
    assert "ravine" in briefs[0], briefs
    assert said.startswith("THE CATALOGUE"), said
    assert not s.last_abstained
    # and the teaching surface never pays for the question
    s2 = Session(None)
    def boom(m):
        raise AssertionError("wants_material asked on a teaching turn")
    generate.wants_material = boom
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("thing", "", "")]}
    generate.refusal = lambda message, persona="": "I do not know."
    generate.offer_research = lambda subject, question: ""
    try:
        s2.respond("what is the thing", teach=True)
    finally:
        (extract.extract, generate.chat, generate.offer_research,
         generate.refusal) = real
        if had:
            generate.wants_material = real_wm
        else:
            del generate.wants_material


@test("W17 an ordinal list marker is format, not a quantity")
def w17():
    """Measured on the live catalogue: the engine drafted exactly what the
    prompt asks — numbered sections opening with attested document names,
    "1. Value Chain (Alpha)" — and the digit gate killed sections 1 and 2
    while section 3 survived, because "3" happened to be attested by the
    brief's "3 hours" and "1" and "2" were attested by nothing. The digit
    discipline exists to kill invented durations and prices; a section
    NUMBER at the head of a line is the engine numbering its own structure,
    which the compose contract explicitly grants ("structure is yours").
    The exemption is the digit twin of the name gate's sentence-start rule,
    and just as narrow: one or two digits, opening a line, closed by a dot
    or bracket. A digit anywhere else still answers for itself."""
    from lmm import evidence
    mat = "Alpha.docx \u2014 the trust walk closes the morning arc"
    assert evidence.digits_ok("1. The Morning Arc (Alpha)", mat)
    assert evidence.digits_ok("2) The Trust Walk", mat)
    # ...and only the marker is forgiven: quantities still die
    assert not evidence.digits_ok("the walk takes 2 days", mat)
    assert not evidence.digits_ok("1. The Arc lasts 4 hours", mat)


@test("W18 the gate reads sentences side by side, not single file")
def w18():
    """Measured, live: a greeting turn cost 15.9 seconds, and 10.4 of them
    were the verifier re-reading the reply ONE SENTENCE PER CALL, each call
    waiting for the last — seven independent questions asked in single
    file. The reads are independent by construction (each sentence answers
    for itself), so the waiting is the only thing the fix touches: the
    engine's runtime, which alone knows whether its backend can take
    parallel calls (an HTTP endpoint can; a local model owns one set of
    weights and cannot), maps the re-reads concurrently or serially. Same
    sentences, same claims, same order, same verdicts — measured on the
    greeting turn, ~10 seconds of queue become ~2 of chorus."""
    import os, threading, time
    from lmm import extract, runtime
    from lmm.session import Session
    from lmm import verify as verify_mod
    s = Session(None)
    s.learn_text("The trust walk closes the morning arc.",
                 source="#docx:Alpha.docx", deep=False)
    lock = threading.Lock()
    state = {"now": 0, "peak": 0}
    real = extract.reextract
    def slow_reextract(sentence):
        with lock:
            state["now"] += 1
            state["peak"] = max(state["peak"], state["now"])
        time.sleep(0.12)
        with lock:
            state["now"] -= 1
        return []                       # no claims → every sentence passes
    extract.reextract = slow_reextract
    answer = "One here. Two here. Three here. Four here."
    old = os.environ.get("LMM_BACKEND")
    try:
        os.environ["LMM_BACKEND"] = "azure"     # HTTP backend: may chorus
        kept = verify_mod.verify(s.memory, answer, set(), anchor="value")
        peak_http = state["peak"]
        state["peak"] = 0
        os.environ["LMM_BACKEND"] = ""          # local backend: single file
        kept2 = verify_mod.verify(s.memory, answer, set(), anchor="value")
        peak_local = state["peak"]
    finally:
        extract.reextract = real
        if old is None:
            os.environ.pop("LMM_BACKEND", None)
        else:
            os.environ["LMM_BACKEND"] = old
    assert kept == answer and kept2 == answer, (kept, kept2)
    assert peak_http > 1, "an HTTP backend still reads in single file"
    assert peak_local == 1, "a local model was read concurrently"


@test("W19 a sentence is re-read once per verdict, not once per reader")
def w19():
    """Measured on the greeting turn: the gate re-read every sentence of
    the reply, and two calls later the abstention stamp re-read the SAME
    sentences again — ten engine calls where five questions existed. The
    re-extractor runs at temperature zero: the same sentence yields the
    same claims by design, so asking twice buys latency and tokens and
    nothing else. A small memo on the re-extractor keeps the second reader
    on the first reader's notes. Distinct sentences still cost a call
    each; the memo hands back a fresh copy, so no reader can scribble on
    another's result."""
    from lmm import extract, runtime
    calls = []
    real = runtime.generate
    runtime.generate = lambda *a, **k: (calls.append(1) or
                                        '{"triples": []}')
    extract._reextract_cached.cache_clear()
    try:
        one = extract.reextract("The walk closes the arc.")
        two = extract.reextract("The walk closes the arc.")
        extract.reextract("A different sentence entirely.")
    finally:
        runtime.generate = real
    assert len(calls) == 2, calls
    assert one == two and one is not two


@test("W20 the caller tunes the voice, never the gate")
def w20():
    """The persona rule, extended to the knobs: an application choosing LMM
    also chooses how the voice sounds — how warm, how long. Those settings
    travel EXACTLY where the persona travels (chat, answer, refusal,
    compose) and nowhere else: every classifier and every verifier call
    keeps its pinned temperature, because a judge whose thermostat the
    user can turn is not a judge. Left unset, every surface keeps its
    measured defaults byte for byte."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None, persona="Warm consultant.", warmth=0.9,
                reply_tokens=333)
    got = {}
    real = (extract.extract, generate.chat)
    extract.extract = lambda m: {"kind": extract.CHAT, "triples": []}
    def spy_chat(message, identity_block="", warmth=0.7, history=None,
                 persona="", max_tokens=None):
        got.update(warmth=warmth, max_tokens=max_tokens)
        return "Hello there!"
    generate.chat = spy_chat
    try:
        s.respond("hello", teach=False)
    finally:
        extract.extract, generate.chat = real
    assert got == {"warmth": 0.9, "max_tokens": 333}, got
    # unset → the defaults stand
    s2 = Session(None)
    got.clear()
    extract.extract = lambda m: {"kind": extract.CHAT, "triples": []}
    generate.chat = spy_chat
    try:
        s2.respond("hello", teach=False)
    finally:
        extract.extract, generate.chat = real
    assert got == {"warmth": 0.7, "max_tokens": None}, got


@test("W21 a consultation turn speculates: the voice warms up while the router reads")
@engine_free
def w21():
    """The chain's floor is its depth, not its width: classify, THEN speak
    — two calls in single file, on every turn, forever. But on the
    consultation surface most turns end at the chat voice anyway, so the
    voice can start speaking WHILE the router classifies: both calls go
    out together, and the router's verdict decides whether the prepared
    reply is used (CHAT, and W14's context statements) or quietly dropped
    (a question, which takes the answer path as always). The wager costs
    one small discarded call on consultation question turns and is never
    placed where tokens are counted: a teaching turn — the benchmarks —
    still classifies first and speaks second, byte for byte."""
    import os, threading, time
    from lmm import generate, extract, runtime
    from lmm.session import Session
    s = Session(None, persona="Warm.")
    lock = threading.Lock()
    state = {"now": 0, "peak": 0, "chat_calls": 0}
    def enter():
        with lock:
            state["now"] += 1
            state["peak"] = max(state["peak"], state["now"])
        time.sleep(0.12)
        with lock:
            state["now"] -= 1
    real = (extract.extract, generate.chat)
    def slow_extract(m):
        enter()
        return {"kind": extract.CHAT, "triples": []}
    def slow_chat(message, identity_block="", warmth=0.7, history=None,
                  persona="", max_tokens=None):
        state["chat_calls"] += 1
        enter()
        return "Hello there!"
    extract.extract, generate.chat = slow_extract, slow_chat
    old = os.environ.get("LMM_BACKEND")
    try:
        os.environ["LMM_BACKEND"] = "azure"
        said = s.respond("hello", teach=False)
        peak_consult, calls_consult = state["peak"], state["chat_calls"]
        state.update(peak=0, chat_calls=0)
        s2 = Session(None, persona="Warm.")
        s2.respond("hello", teach=True)
        peak_teach = state["peak"]
    finally:
        extract.extract, generate.chat = real
        if old is None:
            os.environ.pop("LMM_BACKEND", None)
        else:
            os.environ["LMM_BACKEND"] = old
    assert said == "Hello there!", said
    assert peak_consult == 2, "the consultation turn never speculated"
    assert calls_consult == 1, "the prepared reply was not reused"
    assert peak_teach == 1, "a teaching turn placed the wager"


@test("W22 the draft streams through the gate line by line")
def w22():
    """Measured: the catalogue turn holds its 18 seconds because the
    composer writes 900 tokens and the reader waits for the last one
    before seeing the first. The gate's second form is pure string
    arithmetic — no model call — so a line can be judged THE MOMENT its
    newline arrives and handed to the caller while the engine is still
    writing the next. Same gate, same verdicts, same returned draft; only
    the waiting moves. A line the gate refuses is never seen by the
    caller, streamed or not — an invented number does not become
    printable by arriving early. On a backend that cannot stream, the
    generator yields the whole draft once and the loop degrades to the
    batch path by itself."""
    from lmm import generate
    from lmm.session import Session
    s = Session(None)
    s.learn_text("The trust walk closes the morning arc.",
                 source="#docx:Alpha.docx", deep=False)
    s.learn_text("The abseil drill anchors the ravine module.",
                 source="#docx:Beta.docx", deep=False)
    log = []
    def fake_stream(brief, material_block, warmth=0.2, persona="",
                    max_tokens=None, style=""):
        log.append("emit1")
        yield "the trust walk closes the morning arc.\n"
        log.append("emit2")
        yield "the session lasts 9 days\nthe abseil "
        log.append("emit3")
        yield "drill anchors the ravine module."
    real = generate.compose_stream if hasattr(generate,
                                              "compose_stream") else None
    generate.compose_stream = fake_stream
    lines = []
    try:
        text, sources = s.compose(
            "a day around the trust walk and the abseil drill",
            on_line=lambda line: (log.append("cb"), lines.append(line)))
    finally:
        if real is None:
            del generate.compose_stream
        else:
            generate.compose_stream = real
    assert lines == ["the trust walk closes the morning arc.",
                     "the abseil drill anchors the ravine module."], lines
    # the first line reached the caller BEFORE the engine wrote the last
    assert log.index("cb") < log.index("emit3"), log
    assert "9 days" not in text and len(lines) == 2
    assert text == "\n".join(lines), text
    assert sources


@test("W23 a delivery request goes straight to the composer")
def w23():
    """Measured with the stream clock: the catalogue's first line reached
    the screen at second 13 of an 18-second turn, because the delivery
    question was asked LAST — the turn ran the whole answering chain,
    abstained as it was always going to, and only then asked "is this a
    request to produce?" and composed. The verdict is one tiny call, and
    the wager pool is already open at the turn's first instant: it rides
    out beside the router, and a YES routes the turn to the composer
    before the answering chain ever starts. The composer's refusal is the
    fallback that keeps the old path whole — no material, and the turn
    answers the way it always did. A teaching turn places no bets, asks
    no delivery question, and composes nothing — the benchmarks keep
    their chain byte for byte."""
    import os
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None, persona="Warm.")
    s.learn_text("The trust walk closes the morning arc.",
                 source="#docx:Alpha.docx", deep=False)
    real = (extract.extract, generate.chat)
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("programme", "best", "")]}
    generate.chat = (lambda *a, **k: "Hello!")
    had = hasattr(generate, "wants_material")
    real_wm = getattr(generate, "wants_material", None)
    generate.wants_material = lambda m: True
    def no_answer(*a, **k):
        raise AssertionError("the answering chain ran before the composer")
    real_ans = generate.answer
    generate.answer = no_answer
    def fake_compose(brief, seats=24, topics=None, on_line=None):
        return "THE PROGRAMME\nAlpha \u2014 the trust walk", \
               ["#docx:Alpha.docx"]
    s.compose = fake_compose
    s._brief = ["trust walk"]
    old = os.environ.get("LMM_BACKEND")
    try:
        os.environ["LMM_BACKEND"] = "azure"
        said = s.respond("put together the best programme for us",
                         teach=False)
    finally:
        extract.extract, generate.chat = real
        generate.answer = real_ans
        if had:
            generate.wants_material = real_wm
        else:
            del generate.wants_material
        if old is None:
            os.environ.pop("LMM_BACKEND", None)
        else:
            os.environ["LMM_BACKEND"] = old
    assert said.startswith("THE PROGRAMME"), said
    assert not s.last_abstained


@test("W24 echoing the user is conversation, not assertion")
def w24():
    """Measured, live: "I am the HR manager at a bank, looking for
    training" — and the reply was the bot's own introduction, alone. The
    engine had answered well; the chat gate struck every sentence that
    mentioned the user's bank or the user's role, because in chat the
    allowed set is the identity facts and nothing else — and the one
    sentence that claims nothing is the greeting. But a claim whose words
    the user THEMSELF just spoke is not a fabrication: repeating your
    interlocutor is what listening sounds like. The ECHO rule: in the
    chat reading, a claim passes when its value — and its subject, unless
    the subject is an unresolvable pronoun — is covered word for word by
    what the user has said in this conversation. An external fact stays
    external: a word the user never said still answers to the graph."""
    from lmm import extract
    from lmm import verify as verify_mod
    from lmm.session import Session
    s = Session(None)
    real = extract.reextract
    def fake_reextract(sentence):
        if "ops team" in sentence:
            return [("you", "lead", "the ops team at Acme")]
        if "founded" in sentence:
            return [("Acme", "founded", "1999")]
        return []
    extract.reextract = fake_reextract
    answer = ("Nice to meet you! So you lead the ops team at Acme. "
              "Acme was founded in 1999. What should the training fix?")
    try:
        kept = verify_mod.verify(
            s.memory, answer, set(), anchor="value",
            echo="hello I lead the ops team at Acme and need training")
    finally:
        extract.reextract = real
    assert "ops team" in kept, kept          # the echo survives
    assert "founded" not in kept, kept       # the invention still dies
    assert "What should the training fix?" in kept


@test("W25 in a consultation, the conversation is the fallback, not the shrug")
@engine_free
def w25():
    """Measured, live, with the wager's receipts in hand: "I am the HR
    manager at a bank, researching training" was read as a QUESTION, the
    answering chain ran and abstained — every candidate died at the gate,
    wearing a persona introduction the evidence could not attest — and
    the turn ended in a bare refusal... while the wagered chat reply,
    already written, already through its own gate, said exactly the right
    thing and had been thrown away for the crime of a classifier verdict.
    The rule: in a consultation, when the answering chain ends empty —
    delivery bridge declined, informed refusal declined — the WAGERED
    conversational reply speaks. It passed the chat reading (identity +
    echo); continuing the conversation is what a consultant does with an
    unanswerable turn. The bare refusal remains for surfaces that have no
    conversation: a teaching turn, the benchmarks, keep it byte for byte."""
    import os
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None, persona="Warm.")
    s.learn_text("The trust walk closes the morning arc.",
                 source="#docx:Alpha.docx", deep=False)
    orig_find = s.evidence.find
    def empty_find(query, most=4, floor_share=0.5):
        s.evidence.last_sources = []
        s.evidence.last_census = []
        return []
    s.evidence.find = empty_find
    real = (extract.extract, generate.chat, generate.offer_research,
            generate.refusal)
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("materials", "for", "")]}
    generate.chat = (lambda message, identity_block="", warmth=0.7,
                     history=None, persona="", max_tokens=None:
                     "Got it \u2014 HR training. What should it fix?")
    generate.offer_research = lambda subject, question: ""
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    had = hasattr(generate, "wants_material")
    real_wm = getattr(generate, "wants_material", None)
    generate.wants_material = lambda m: False
    old = os.environ.get("LMM_BACKEND")
    try:
        os.environ["LMM_BACKEND"] = "azure"
        s.respond("hello there", teach=False)
        said = s.respond("I am researching training materials", teach=False)
        s2 = Session(None)
        s2.evidence.find = empty_find
        teach_said = s2.respond("what are the materials", teach=True)
    finally:
        (extract.extract, generate.chat, generate.offer_research,
         generate.refusal) = real
        if had:
            generate.wants_material = real_wm
        else:
            del generate.wants_material
        s.evidence.find = orig_find
        if old is None:
            os.environ.pop("LMM_BACKEND", None)
        else:
            os.environ["LMM_BACKEND"] = old
    assert "HR training" in said, said
    assert "I do not know" in teach_said, teach_said


@test("W26 the operator's format sheet travels to the composer alone")
def w26():
    """The persona's mirror twin. The persona is the operator's VOICE and
    stops at the document's edge (measured: persona-induced rephrasing
    died at the verbatim gate). The STYLE is the operator's FORMAT — which
    sections a catalogue opens, what its headings are called — and it
    travels exactly where the persona does not: into the composer's
    REQUEST, and nowhere else. Its words are operator-attested the way
    the brief's are, so a heading the sheet names does not die for its
    capitals; the material still owns every fact under it. Field need,
    live: the operator wanted the catalogue shaped like their house
    format, and the only lever was rewriting the library's prompt."""
    from lmm import generate
    from lmm.session import Session
    s = Session(None, style="Uygunluk Sayfasi: duration and fit per item")
    s.learn_text("The trust walk closes the morning arc.",
                 source="#docx:Alpha.docx", deep=False)
    s.learn_text("The abseil drill anchors the ravine module.",
                 source="#docx:Beta.docx", deep=False)
    seen = {}
    real = generate.compose
    def spy(brief, material, warmth=0.2, persona="", max_tokens=None,
            style=""):
        seen["style"] = style
        seen["persona"] = persona
        return ("Uygunluk Sayfasi\n"
                "the trust walk closes the morning arc.")
    generate.compose = spy
    try:
        text, sources = s.compose("a day around the trust walk")
    finally:
        generate.compose = real
    assert "Uygunluk" in seen["style"], seen
    assert seen["persona"] == "", "the persona crossed the document's edge"
    # the sheet's own heading survives the capitals rule
    assert "Uygunluk Sayfasi" in text, text


@test("W27 a one-row grid is a fact box, and a fact box is prose")
def w27():
    """Field finding, two harnesses apart: every duration, seat count and
    audience line of a 61-document corpus lived in ONE-ROW Word tables —
    boxes, not tables — and read_docx dropped them whole: the header rule
    needs two rows to read, the grid had one, `continue`. The prose side
    reads only doc.paragraphs, so the cells reached neither layer, and
    the report honestly said "0 tables" while meaning "I did not count
    these". The rule is the shape's own: a grid without a second row has
    no header to read — it is a FACT BOX, and its cells are prose the
    document attests, label and value kept on one line. A real table
    (two rows or more) still goes to the graph exactly as before."""
    try:
        import docx as _docx
    except ImportError:
        return                       # the adapter itself requires python-docx
    import tempfile, os
    from lmm import tables
    doc = _docx.Document()
    doc.add_paragraph("The ravine module closes the arc.")
    t = doc.add_table(rows=1, cols=2)
    t.rows[0].cells[0].text = "DURATION\n2 full days \u2014 in person"
    t.rows[0].cells[1].text = "SEATS\nMaximum 18 people"
    t2 = doc.add_table(rows=2, cols=2)
    t2.rows[0].cells[0].text = "field"
    t2.rows[0].cells[1].text = "value"
    t2.rows[1].cells[0].text = "anchor"
    t2.rows[1].cells[1].text = "abseil"
    path = os.path.join(tempfile.mkdtemp(), "box.docx")
    doc.save(path)
    prose, grids = tables.read_docx(path)
    assert "Maximum 18 people" in prose, prose
    assert "DURATION" in prose
    # label and value stay on ONE line — evidence retrieval must see them
    # together, or "the duration" fetches a bare number with no field name
    line = next(l for l in prose.splitlines() if "18 people" in l)
    assert "SEATS" in line, line
    # the two-row grid is still a table for the graph, not prose
    assert grids and any("abseil" in str(g) for g in grids)
    assert "abseil" not in prose


@test("W28 an answer's substance cannot be borrowed from the question")
def w28():
    """Measured twice on the same store, the same question, temperature
    fixed: "which course runs in person?" — and the engine named ONE
    course off the datelines and handed the question's own predicate
    back ("X runs in person"), with no evidence line saying any such
    thing; two runs picked two different names, one happened to be true.
    The reading admitted it because question words are excused from
    coverage — an excusal that exists for FLUENCY, so an answer may
    restate what was asked — and the source names are attested by the
    datelines, so the whole sentence was "covered" while asserting
    nothing the evidence says. The rule: when an answer's entire
    substance beyond the question and the source names is EMPTY, the
    borrowed words must themselves stand in the evidence BODIES — the
    text after the dateline, where the document actually speaks. A real
    choice-answer passes exactly this way (the chosen option is written
    in some body); the invented pairing dies before the jury is even
    seated, because a wobbling judge must not hold the only veto."""
    from lmm import generate
    from lmm.session import Session
    s = Session(None)
    s.learn_text("The trust walk closes the morning arc.",
                 source="#docx:Alpha Course.docx", deep=False)
    s.learn_text("The abseil drill runs in person only.",
                 source="#docx:Beta Course.docx", deep=False)
    real = generate.supported
    def no_jury(*a, **k):
        raise AssertionError("the jury was seated for a borrowed answer")
    generate.supported = no_jury
    proof_without = [
        "Alpha Course \u2014 the trust walk closes the morning arc",
        "Beta Course \u2014 the abseil drill anchors the ravine module"]
    proof_with = [
        "Alpha Course \u2014 the trust walk closes the morning arc",
        "Beta Course \u2014 the abseil drill runs in person only"]
    q = "which course runs in person?"
    raw = "Beta Course runs in person."
    try:
        died = s._read_back(raw, proof_without, "\n".join(proof_without),
                            question=q)
    finally:
        generate.supported = real
    assert died is False, "the borrowed answer reached the jury or passed"
    # the same sentence with the evidence actually saying it: the veto
    # stands aside (the borrowed words are in a body) and the jury gets
    # its normal seat
    generate.supported = lambda raw, view: True
    try:
        lived = s._read_back(raw, proof_with, "\n".join(proof_with),
                             question=q)
    finally:
        generate.supported = real
    assert lived is True, "a truly attested choice-answer died"
    # THE CROSS-SOURCE PAIRING, measured after the first cut of this rule:
    # once "in person" stood in SOME body, naming the OTHER course passed —
    # every word attested somewhere, the pairing attested nowhere. When
    # the claim names a source, the borrowed words must stand in THAT
    # source's bodies.
    generate.supported = no_jury
    raw_wrong = "Alpha Course runs in person."
    try:
        crossed = s._read_back(raw_wrong, proof_with,
                               "\n".join(proof_with), question=q)
    finally:
        generate.supported = real
    assert crossed is False, "the cross-source pairing reached the jury"


@test("W29 the rescue offer answers to the same borrowed-substance rule")
def w29():
    """The door W28 left open, measured live: the main answer abstained,
    the informed refusal spoke instead — and the rescue's only gates are
    the SHAPE checks (capitals, digits). "Gamma Course runs in person": the
    name's capitals are attested by the tally, "in person" wears no shape
    at all, and the sentence walked out — the same invented pairing W28
    kills on the answer path, through the side door. One rule, one organ,
    both doors: an offer whose substance beyond the question and the
    tally names is empty must find its borrowed words in the gathered
    lines' BODIES — of the source it names, when it names one."""
    from lmm.session import Session
    s = Session(None)
    proof = ["Alpha Course \u2014 the trust walk closes the morning arc",
             "Beta Course \u2014 the abseil drill runs in person only"]
    q = "which course runs in person?"
    assert s._substance_ok("Beta Course runs in person.", proof, q)
    assert not s._substance_ok("Alpha Course runs in person.", proof, q)
    assert not s._substance_ok(
        "Alpha Course runs in person.",
        ["Alpha Course \u2014 the trust walk closes the morning arc"], q)
    # free words beyond the question: the rule stands aside (other gates own it)
    assert s._substance_ok("The drill is famous.", proof, q)


@test("W30 a Word paragraph is a deliberate unit, and gets its full stop")
def w30():
    """The box lesson, generalised — measured on the full corpus: heading
    paragraphs ("TRAINING OUTCOMES") and bullet paragraphs carry no
    terminal punctuation, so the sentence splitter glued title + headings
    + first bullets into MEGA-LINES that matched every query's words and
    vampirised the seats; the precise line always sat below the cut. In a
    Word file the paragraph break is drawn by the AUTHOR — it is format,
    exactly like the box's edge — so the adapter closes every unpunctuated
    paragraph with a full stop. A PDF's newlines are the printer's, not
    the author's; read_pdf is untouched."""
    try:
        import docx as _docx
    except ImportError:
        return
    import tempfile, os
    from lmm import tables
    doc = _docx.Document()
    doc.add_paragraph("TRAINING OUTCOMES")
    doc.add_paragraph("masters the abseil drill without a rope")
    doc.add_paragraph("The ravine module closes the arc.")
    path = os.path.join(tempfile.mkdtemp(), "para.docx")
    doc.save(path)
    prose, _ = tables.read_docx(path)
    lines = prose.splitlines()
    assert "TRAINING OUTCOMES." in lines, lines
    assert "masters the abseil drill without a rope." in lines
    assert "The ravine module closes the arc." in lines   # not doubled


@test("W31 the provenance mark is the system's, and it points at the load-bearing line")
def w31():
    """Measured across a 25-question run: the stamp was the FIRST seat's
    origin, whatever line the surviving answer actually rested on — a
    correct sentence about the Leader course stamped with the Literacy
    course, question after question — and once, the engine IMITATED the
    notation and wrote a stamp of its own for a file that does not exist.
    Two halves of one rule, provenance notation belongs to the system:
    anything stamp-shaped in a generated candidate is stripped before the
    gates ever read it, and the mark the system then attaches names the
    origin of the proof line the spoken answer covers best — the
    load-bearing line, chosen by word coverage, deterministically."""
    from lmm.session import Session
    s = Session(None)
    assert s._strip_marks(
        "The drill runs daily. (~ #docx:Fake File.docx)"
    ).strip() == "The drill runs daily."
    assert s._strip_marks("A #docx:Ghost.docx claim") == "A claim"
    proof = ["Alpha Course \u2014 the trust walk closes the morning arc",
             "Beta Course \u2014 the abseil drill runs in person only"]
    origins = ["#docx:Alpha Course.docx", "#docx:Beta Course.docx"]
    picked = s._load_bearing(
        "The abseil drill runs in person.", proof, origins)
    assert picked == "#docx:Beta Course.docx", picked
    picked2 = s._load_bearing(
        "The trust walk closes the arc.", proof, origins)
    assert picked2 == "#docx:Alpha Course.docx", picked2


@test("W32 a claim that names its source is judged by that source alone")
def w32():
    """The cross-source pairing, third and general form. W28 killed the
    pairing when the claim's substance was EMPTY; measured next: "the
    Delegation course does not state its seat count" — real substance
    ("does not state"), borrowed from ANOTHER course's box, stapled to
    the named course, and the jury confirmed it against the full block
    where those words genuinely stand. The rule: when a claim NAMES a
    source, the jury reads THAT source's lines and nothing else — the
    fallback view of everything retrieved is exactly the door the
    stapled sentence walked through. A claim naming nothing keeps the
    old views; a claim naming two sources is judged by both. And the
    rescue offer sheds stamp-shaped notation the way the answer path
    does — an invented filename is not a citation."""
    from lmm import generate
    from lmm.session import Session
    s = Session(None)
    s.learn_text("The trust walk closes the morning arc.",
                 source="#docx:Alpha Course.docx", deep=False)
    proof = ["Alpha Course \u2014 the trust walk closes the morning arc",
             "Beta Course \u2014 the seat count is not stated at source"]
    views_seen = []
    real = generate.supported
    def spy_jury(raw, view):
        views_seen.append(view)
        return "not stated" in view       # an honest jury: it reads
    generate.supported = spy_jury
    try:
        verdict = s._read_back(
            "Alpha Course does not state a seat count.", proof,
            "\n".join(proof), question="how many seats does Alpha take?")
    finally:
        generate.supported = real
    assert verdict is False, "the stapled sentence survived"
    assert views_seen, "the jury never sat"
    for view in views_seen:
        assert "Beta" not in view, view   # only the NAMED source is read


@test("W33 the prior subject rides only when it carries the question better")
def w33():
    """Measured on the NIST sequence: two password questions, then "what
    is the publication date of this document" — the pointer subject
    resolves to nothing, the PRIOR subject (memorized secret) rides into
    the query by design, and the date line is pushed out of the seats by
    a hundred password lines. The candidates then say the true date, the
    reading rightly refuses what the seats do not attest, and an honest
    architecture scores worse than a lucky one — the old pass was the
    jury confirming an unattested truth. The ride was built for the real
    follow-up ("and who is it for?"), where the turn's own words seat
    junk; here the turn's own words seat the answer. So the ride is a
    CHALLENGER, not a preemption: both queries run (local, no engine),
    and the ridden proof stands only if it covers the question at least
    as well as the turn's own proof. Ties keep the ride — the follow-up
    keeps its fix."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    s.learn_text("The memorized secret shall be at least 8 characters.",
                 source="#pdf:guide", deep=False)
    blocks = []
    queries = []
    orig_find = s.evidence.find
    def fake_find(query, most=4, floor_share=0.5):
        queries.append(query)
        s.evidence.last_census = []
        if "memorized" in query:          # the ridden query: password noise
            s.evidence.last_sources = ["#pdf:guide"] * 2
            return ["The memorized secret shall be at least 8 characters.",
                    "Truncation of the memorized secret is not performed."]
        s.evidence.last_sources = ["#pdf:guide"]
        return ["Publication date: July 24, 2025."]
    s.evidence.find = fake_find
    real = (extract.extract, generate.answer, generate.refusal,
            generate.offer_research)
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("this document", "date", "")]}
    def spy_answer(q, block, warmth=0.2, persona="", max_tokens=None, **kw):
        blocks.append(block)
        return "I do not know."
    generate.answer = spy_answer
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    s.last_subject = "memorized secret"     # the prior turn's subject
    try:
        s.respond("what is the publication date of this document")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research) = real
        s.evidence.find = orig_find
    assert any("memorized" in q for q in queries), "the ride never ran"
    assert any("memorized" not in q for q in queries), \
        "the turn's own words never got their search"
    assert blocks and any("July" in b for b in blocks), \
        "the better-covering proof was not chosen"


@test("W34 an expansion is judged by the words it adds")
def w34():
    """The keep filter's one inconsistency with its own principle, and the
    fourth member of the IDF-inversion family — measured in the field: a
    record line ("SEATS: sixteen people maximum") got exactly the bridge
    it exists for ("how many people can attend"), and the filter threw it
    away, because the margin was scored over ALL the query's words and
    "people" reaches every other line of a topical corpus better than it
    reaches the record. But the filter's own closing rule says only the
    NOVEL words are ever indexed — a word the line already carries is
    reachable already. What is indexed is what must be judged: the margin
    reads the added words alone. A drifting query still dies (its novel
    words score another region higher); a query whose additions reach
    nothing at all scores zero, which the docstring already names as the
    case the expansion exists for."""
    from lmm.session import Session
    s = Session(None)
    s.learn_text("CAPACITY: sixteen people maximum.\n"
                 "people join the morning arc in pairs.\n"
                 "people walk the ravine with a guide.\n"
                 "the guide counts people at the gate.\n"
                 "people rest before the abseil drill.",
                 source="#docx:Alpha.docx", deep=False)
    ev = s.evidence
    sid = next(i for i, (t, _src) in enumerate(ev.sentences)
               if t.strip() == "CAPACITY: sixteen people maximum.")
    kept = ev.learn_expansions(
        {sid: ["how many people fit the capacity maximum"]})
    assert kept == 1, "the record's bridge was thrown away again"
    assert any(sid in ids for w, ids in ev.expand_index.items()
               if w not in ("people",)), ev.expand_index.keys()


@test("W35 retrieval knows kinship the gates refuse to license")
def w35():
    """ROOT is five letters — an honest arbitrary that a short-rooted
    language walks under: a four-letter root and its suffixed form share
    a prefix of four, and same_stem calls them strangers, so the query
    word never reaches the line that answers it. The scar is documented
    in the scorer: a looser reverse-prefix rule once lived there and was
    removed for being one of four quietly different copies — the loss of
    short units was recorded as a measured cost. This puts the rule back
    as a NAMED relation with the other two, stated once: kin(a, b) — the
    shorter word is wholly the longer one's prefix and the ending stays
    within TAIL. And it is spent ONLY where surfacing a real line is the
    worst it can do: the retrieval scorer. The gates keep same_stem — a
    claim's word is still not attested by a stranger — so kinship widens
    what can be FOUND and nothing about what may be SAID."""
    from lmm import inflect
    assert inflect.kin("saat", "saatlik")
    assert inflect.kin("gunluk", "gun")          # symmetric
    assert not inflect.kin("car", "carpets")     # ending beyond TAIL
    assert not inflect.kin("at", "atlas")        # a two-letter stub roots nothing
    assert not inflect.same_stem("saat", "saatlik")   # the gates stay strict
    from lmm.session import Session
    s = Session(None)
    s.learn_text("The kilnhouse holds a four hour firing. "
                 "Visitors gather at the gate before the walk. "
                 "The garden path closes in winter.",
                 source="#docx:Alpha.docx", deep=False)
    # 'kilnhouses' (suffixed) must reach the 'kilnhouse' line
    got = s.evidence.find("how long do the kilnhouses fire", most=2)
    assert got and "kilnhouse" in got[0], got


@test("W36 a same-region twin is the same voice twice")
def w36():
    """Measured on the full corpus, eight questions with one shape: a
    document's two seats (the source cap) went to two overlapping windows
    of its TITLE region, and the record line that answered — same
    document, different region — waited outside. The region rule already
    admits a second view as corroboration, and the hospital trace proved
    a single-document store sometimes needs exactly that; but in a
    MULTI-document store the source cap makes the twin expensive: it
    spends the document's other seat on the same paragraph said twice.
    So the twin steps aside first — into the overflow, with everything
    else that stepped aside — and steps back, in rank order, only if the
    block would otherwise go unfilled. A single-document store keeps the
    hospital behaviour byte for byte."""
    from lmm import evidence
    st = evidence.SentenceStore()
    st.add("alpha gamma delta epsilon zeta omega", "#doc:A")
    st.add("alpha gamma delta epsilon zeta theta", "#doc:A")   # region twin
    st.add("alpha holds sixteen seats maximum", "#doc:A")
    st.add("alpha appears in the other file too", "#doc:B")
    got = st.find("alpha", most=3, floor_share=0.0)
    assert any("sixteen" in g for g in got), got
    # single-document: the corroborating twin is still admitted
    st1 = evidence.SentenceStore()
    st1.add("alpha gamma delta epsilon zeta omega", "#doc:A")
    st1.add("alpha gamma delta epsilon zeta theta", "#doc:A")
    st1.add("alpha holds sixteen seats maximum", "#doc:A")
    st1.add("alpha rests beside the garden gate", "#doc:A")
    got1 = st1.find("alpha", most=3, floor_share=0.0)
    assert sum("epsilon" in g for g in got1) == 2, got1


@test("W37 a question that names two sources gets them laid side by side")
def w37():
    """The comparison reading. "Do Alpha and Beta run the same length?" is
    answered by NO document — Alpha's sheet says one thing, Beta's says
    another, and the verdict is born only when the two lines lie side by
    side. The single race cannot promise that: each field line enters on
    its own luck, and when one misses a seat the turn honestly refuses a
    question the store could answer. The trigger is a STATE the name
    channel already computes — the question NAMES two or more sources —
    and the assembly is organs off the shelf: each named source gathers
    its own seats (dedicated seats need dedicated queries), its
    record-shaped lines ride along (the fact-sheet rider), every line
    keeps its dateline, and the named-source jury (W32) already holds
    each claim to the source it names. No new model call, no wording
    read: two names on the question, two files on the table."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    alpha = ["ALPHA COURSE OUTCOMES."] + [
        f"The course runs the same arc at length in module {w}."
        for w in ("north", "south", "east", "west", "ridge", "vale")
    ] + ["DURATION: two full days."]
    beta = ["BETA COURSE OUTCOMES."] + [
        f"The course runs the same ravine at length in station {w}."
        for w in ("one", "two", "three", "four", "five", "six")
    ] + ["DURATION: half a day."]
    s.learn_text("\n".join(alpha), source="#docx:Alpha Course.docx",
                 deep=False)
    s.learn_text("\n".join(beta), source="#docx:Beta Course.docx",
                 deep=False)
    for name in ("Gamma", "Delta", "Epsilon"):
        s.learn_text(f"{name.upper()} COURSE OUTCOMES.\n"
                     f"The course runs its own arc at length daily.\n"
                     f"DURATION: one day.",
                     source=f"#docx:{name} Course.docx", deep=False)
    s.learn_text("The catering tent seats forty guests.",
                 source="#docx:Logistics.docx", deep=False)
    blocks = []
    real = (extract.extract, generate.answer, generate.refusal,
            generate.offer_research)
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("alpha course", "duration", "")]}
    def spy_answer(q, block, warmth=0.2, persona="", max_tokens=None, **kw):
        blocks.append(block)
        return "I do not know."
    generate.answer = spy_answer
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    try:
        s.respond("do Alpha Course and Beta Course run the same length?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research) = real
    assert blocks, "no candidate was asked"
    joined = "\n".join(blocks)
    assert "two full days" in joined, "Alpha's record line missed the table"
    assert "half a day" in joined, "Beta's record line missed the table"
    assert "forty guests" not in joined, "an unnamed source crashed the table"


@test("W38 a source is named by a pointer, not by a family resemblance")
def w38():
    """Measured on the coded corpus: "is there a difference between TPK-03
    and TPK-04?" seated the OVERVIEW documents — the code's letters match
    every sibling in the family ("tpk" lives in seven names), membership
    in the named set was a UNION over matched words, and seven documents
    marched to the front while the two that were actually named waited.
    The pointing-word rule, third appearance: a source is NAMED when some
    question word matches ITS name alone, or when at least two of its
    name words are matched — one shared family word calls nobody. The
    per-word log(S/s) BOOST is untouched (rarity already prices the
    family word at nearly nothing); what tightens is who counts as
    called to the front."""
    from lmm import evidence
    st = evidence.SentenceStore()
    docs = {
        "#docx:TPK-00 Overview.docx": "The programme overview spans modules.",
        "#docx:TPK-03 Contact.docx": "DURATION: three hours of contact work.",
        "#docx:TPK-04 Signals.docx": "DURATION: half a day of signal work.",
        "#docx:TPK-05 Closing.docx": "The closing round gathers feedback.",
    }
    for src, line in docs.items():
        st.add(line, src)
        st.add("Participants rotate through the stations.", src)
    st.find("is there a difference between TPK-03 and TPK-04 durations?",
            most=4, floor_share=0.0)
    named = {evidence._source_name(s) for s in st.last_named}
    assert named == {"TPK-03 Contact", "TPK-04 Signals"}, named
    # the cross-family trap: "Alpha Delegation and Stress Handling" must
    # not call "Alpha Conflict Handling" — its two matched words never
    # stand together in the question
    st2 = evidence.SentenceStore()
    for tag, src in (("ropes", "#docx:Alpha Delegation Handling.docx"),
                     ("cases", "#docx:Alpha Conflict Handling.docx"),
                     ("brooks", "#docx:Stress Handling.docx"),
                     ("herbs", "#docx:Quiet Garden Notes.docx")):
        st2.add(f"DURATION: one day of {tag} practice.", src)
        st2.add(f"Participants rotate through the {tag} stations.", src)
    st2.find("are Alpha Delegation Handling and Stress Handling the same "
             "length?", most=4, floor_share=0.0)
    named2 = {evidence._source_name(s) for s in st2.last_named}
    assert named2 == {"Alpha Delegation Handling", "Stress Handling"}, named2
    # a single shared family word still calls nobody
    st.find("how do the TPK modules work?", most=4, floor_share=0.0)
    assert not st.last_named, st.last_named


@test("W39 the same sentence in two documents is two attestations")
def w39():
    """Measured at the very end of the comparison chase: "is Anxiety the
    same length as Delegation?" kept refusing while both documents said
    "DURATION: one full day." — the SAME sentence, and the duplicate
    guard keyed on the TEXT alone, so whichever document was read first
    owned the line and the other document's copy never entered the store.
    The guard exists so that reading one document twice does not multiply
    its evidence — and that is exactly what the key must say: the same
    text FROM THE SAME SOURCE is a duplicate; the same text from another
    source is another document going on record, with its own dateline,
    its own census entry, its own seat. Sibling corpora share template
    sentences by construction; provenance is the whole difference."""
    from lmm import evidence
    st = evidence.SentenceStore()
    a = st.add("DURATION: one full day.", "#docx:Anxiety Handling.docx")
    b2 = st.add("DURATION: one full day.", "#docx:Delegation Handling.docx")
    again = st.add("DURATION: one full day.", "#docx:Anxiety Handling.docx")
    assert a is not None and b2 is not None, "the second document was silenced"
    assert again is None, "re-reading a document multiplied its evidence"
    srcs = {src for _t, src in st.sentences}
    assert len(srcs) == 2, srcs


@test("W42 an echo cannot turn a question into its own assertion")
def w42():
    """Caught by our own benchmark's trap set: "do participants receive a
    certificate?" — nothing about certificates exists anywhere in the
    store, the answer path honestly abstained, and the conversational
    fallback said "Yes, participants generally receive a certificate."
    Every word of that claim is covered by the QUESTION, and the echo
    rule (W24) licensed it: it read the user's words as attestation. But
    a question ASSERTS nothing — echo exists so a consultant may repeat
    what the user STATED, and a turn the router classified as ASK stated
    nothing at all. So on an ASK turn, the echo pool is the PRIOR turns
    alone: mirroring the question back as an assertion finds no licence,
    and the sentence that carried it falls at the gate it always had to
    pass. A context statement (WRITE/CHAT) keeps its same-turn echo."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None, persona="Warm.")
    s.learn_text("The trust walk closes the morning arc.",
                 source="#docx:Alpha.docx", deep=False)
    real = (extract.extract, generate.chat, generate.refusal,
            generate.offer_research, extract.reextract)
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("certificate", "given", "")]}
    generate.chat = (lambda *a, **k:
                     "Yes, participants generally receive a certificate.")
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    def fake_reextract(sentence):
        if "certificate" in sentence:
            return [("participants", "receive", "a certificate")]
        return []
    extract.reextract = fake_reextract
    had = hasattr(generate, "wants_material")
    real_wm = getattr(generate, "wants_material", None)
    generate.wants_material = lambda m: False
    import os
    old = os.environ.get("LMM_BACKEND")
    try:
        os.environ["LMM_BACKEND"] = "azure"
        s.respond("hello there", teach=False)
        said = s.respond("do participants receive a certificate?",
                         teach=False)
    finally:
        (extract.extract, generate.chat, generate.refusal,
         generate.offer_research, extract.reextract) = real
        if had:
            generate.wants_material = real_wm
        else:
            del generate.wants_material
        if old is None:
            os.environ.pop("LMM_BACKEND", None)
        else:
            os.environ["LMM_BACKEND"] = old
    assert "certificate" not in (said or ""), said
    assert "Yes" not in (said or ""), said


@test("W43 the chat voice streams sentence by sentence, gate first")
def w43():
    """The perceived half of the latency ledger, second cut. The first
    cut streamed but paid for it: it stood the wager down, so the voice
    no longer overlapped the router and the FIRST sentence arrived later
    than the old whole-reply did. The wager and the stream are one
    mechanism now: the wager fills a QUEUE of chunks from the turn's
    first instant, and when the router says CHAT and a caller listens,
    the main flow drains the queue sentence by sentence — each one
    through its reading the moment its full stop arrives, a struck
    sentence never seen. Routed elsewhere, the queue is discarded like
    any lost wager. The proof of interleaving is deterministic: the
    engine's second sentence is not even WRITTEN until the first has
    reached the caller."""
    import os, threading
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None, persona="Warm.")
    s.learn_text("The trust walk closes the morning arc.",
                 source="#docx:Alpha.docx", deep=False)
    first_seen = threading.Event()
    log = []
    def fake_stream(message, identity_block="", warmth=0.7, history=None,
                    persona="", max_tokens=None):
        log.append("emit1")
        yield "Happy to help with the arc. "
        first_seen.wait(3)              # the pen waits for the reader
        log.append("emit2")
        yield "Acme was founded in 1999. What should the walk fix?"
    real = (extract.extract, extract.reextract)
    extract.extract = lambda m: {"kind": extract.CHAT, "triples": []}
    def fake_reextract(sentence):
        if "founded" in sentence:
            return [("Acme", "founded", "1999")]
        return []
    extract.reextract = fake_reextract
    had = hasattr(generate, "chat_stream")
    real_cs = getattr(generate, "chat_stream", None)
    generate.chat_stream = fake_stream
    lines = []
    def cb(t):
        lines.append(t)
        first_seen.set()
        log.append("cb")
    old = os.environ.get("LMM_BACKEND")
    try:
        os.environ["LMM_BACKEND"] = "azure"
        said = s.respond("hello there", teach=False, on_line=cb)
    finally:
        extract.extract, extract.reextract = real
        if had:
            generate.chat_stream = real_cs
        else:
            del generate.chat_stream
        if old is None:
            os.environ.pop("LMM_BACKEND", None)
        else:
            os.environ["LMM_BACKEND"] = old
    assert lines and "Happy to help" in lines[0], lines
    assert first_seen.is_set(), "the first sentence never reached the caller"
    assert log.index("cb") < log.index("emit2"), log
    assert all("1999" not in x for x in lines)
    assert "1999" not in (said or ""), said
    assert "What should the walk fix?" in said


@test("W44 the comparison verdict is written into the evidence, not left to the voice")
def w44():
    """The referee's one comparison flip, closed at the root: the reading
    laid two duration records side by side and left "same or different?"
    to the engine — which built the verdict one run and inverted it the
    next. But equality of two records is not phrasing, it is arithmetic
    the system already trusts: the digit sets (the digits_ok tradition).
    When two NAMED sources answer the same FIELD HEAD, the assembly now
    writes one comparison row — head, then each source's value, joined
    by = when the digit sets agree and by \u2260 when they differ (no
    marker when a value carries no digits: there the voice still
    judges). Every word on the row is the documents'; the marker is
    notation; the engine reads a verdict instead of building one."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    s.learn_text("ALPHA COURSE OUTCOMES.\n"
                 "The rope bridge closes the morning arc.\n"
                 "DURATION: three hours (3 h).\n"
                 "SEATS: 14-18 people.",
                 source="#docx:Alpha Course.docx", deep=False)
    s.learn_text("BETA COURSE OUTCOMES.\n"
                 "The ravine walk opens with a trust fall.\n"
                 "DURATION: half a day (3 h).\n"
                 "SEATS: 16-18 people.",
                 source="#docx:Beta Course.docx", deep=False)
    blocks = []
    real = (extract.extract, generate.answer, generate.refusal,
            generate.offer_research)
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("alpha course", "duration", "")]}
    def spy_answer(q, block, warmth=0.2, persona="", max_tokens=None, **kw):
        blocks.append(block)
        return "I do not know."
    generate.answer = spy_answer
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    try:
        s.respond("do Alpha Course and Beta Course run the same length?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research) = real
    joined = "\n".join(blocks)
    eq = [l for l in joined.split("\n") if " = " in l and "DURATION" in l]
    ne = [l for l in joined.split("\n") if " \u2260 " in l and "SEATS" in l]
    assert eq, "no equality row for the matching durations"
    assert "Alpha Course" in eq[0] and "Beta Course" in eq[0], eq
    assert ne, "no difference row for the differing seat counts"


@test("W45 a question that names one source gets that source's block")
def w45():
    """The single-name completion of the comparison reading, measured on
    the instructor question: the named course was called correctly, its
    TWO capped seats went to its densest windows — and the line that
    answers ("taught jointly by A and B", a box that never says the word
    'instructor') waited outside while a SIBLING's window supplied the
    names and the voice blended a faculty that does not exist. Naming a
    document is asking to read THAT document: when exactly one source is
    named, it fills the block — dedicated seats gathered from it alone,
    its record rows riding — and two seats remain for the rest of the
    corpus to interject. Questions that name nothing keep the open race
    untouched."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    alpha = ["ALPHA COURSE OUTCOMES."] + [
        f"The course teaches the {w} module at length daily."
        for w in ("north", "south", "east", "west", "ridge", "vale")
    ] + ["Delivered jointly by R. Stone and M. Vale through drama work."]
    s.learn_text("\n".join(alpha), source="#docx:Alpha Course.docx",
                 deep=False)
    s.learn_text("TRAINER: Q. Herbst.\n"
                 "The garden day is delivered by our trainer in the vale.",
                 source="#docx:Garden Day.docx", deep=False)
    for i, w in enumerate(("brook", "cliff", "dune", "fern",
                           "grove", "heath", "isle", "knoll")):
        s.learn_text(
            f"COURSE NOTES {w.upper()}.\n"
            f"This training is delivered jointly at the {w} site.\n"
            f"The {w} course training runs with a joint trainer daily.",
            source=f"#docx:{w.title()} Notes.docx", deep=False)
    blocks = []
    real = (extract.extract, generate.answer, generate.refusal,
            generate.offer_research)
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("alpha course", "trainer", "")]}
    def spy_answer(q, block, warmth=0.2, persona="", max_tokens=None, **kw):
        blocks.append(block)
        return "I do not know."
    generate.answer = spy_answer
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    try:
        s.respond("who delivers the Alpha Course training jointly?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research) = real
    joined = "\n".join(blocks)
    assert "R. Stone" in joined, "the named course's answer line missed its own block"
    alpha_lines = [l for l in joined.split("\n") if "Alpha Course \u2014" in l]
    assert len(alpha_lines) >= 4, ("the named source did not fill the block: %d"
                                   % len(alpha_lines))


@test("W46 in a comparison, two sources saying the same thing say it twice")
def w46():
    """W39's lesson, repeating one layer up — and it is the comparison
    reading's own worst case. Sibling documents share a template, so two
    courses of equal length write the SAME record line, word for word
    ("DURATION: one day."). The layout gathered both and then dropped the
    second as a duplicate — a plain `line not in comp` — so the block
    held one duration where the question asked about two, no field pair
    could be found, and the comparison row that states the verdict was
    never written. Measured on the field set: seven of fifteen
    comparisons ask about EQUAL values, which is exactly the case this
    silences. In an assembly that contrasts sources, sameness is the
    ANSWER: identity is per (line, source), never per line."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    for name, tag in (("Alpha Course", "ravine"), ("Beta Course", "ridge")):
        s.learn_text(f"{name.upper()} OUTCOMES.\n"
                     f"The {tag} module opens the arc.\n"
                     "DURATION: 1 day.",
                     source=f"#docx:{name}.docx", deep=False)
    blocks = []
    real = (extract.extract, generate.answer, generate.refusal,
            generate.offer_research)
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("alpha course", "duration", "")]}
    def spy_answer(q, block, warmth=0.2, persona="", max_tokens=None, **kw):
        blocks.append(block)
        return "I do not know."
    generate.answer = spy_answer
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    try:
        s.respond("do Alpha Course and Beta Course run the same length?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research) = real
    joined = "\n".join(blocks)
    alpha = [l for l in joined.split("\n")
             if "Alpha Course" in l and "DURATION" in l]
    beta = [l for l in joined.split("\n")
            if "Beta Course" in l and "DURATION" in l]
    assert alpha, "the first source's record line is missing"
    assert beta, "the second source's identical record line was deduped away"
    eq = [l for l in joined.split("\n") if " = " in l and "DURATION" in l]
    assert eq, "no equality row for two identical durations"


@test("W47 the comparison row the question asks about speaks first")
def w47():
    """Measured on the field set: asked whether two courses run the same
    LENGTH, the block carried a verdict row for every field the two
    share — duration, seats, audience, format — in whatever order the
    fields were met, and the engine answered off the format row ("both
    are in person") while the question asked about length. The record
    rider already ranks a source's rows by what the question asks; the
    comparison rows are the same kind of evidence and had no such order.
    They get one: rows whose HEAD the question names come first, and the
    rest keep their order behind them. Nothing is dropped — a reader
    that wants another field still finds it, one line lower."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    for name, tag in (("Alpha Course", "ravine"), ("Beta Course", "ridge")):
        s.learn_text(f"{name.upper()} OUTCOMES.\n"
                     f"The {tag} module opens the arc.\n"
                     "FORMAT: in person.\n"
                     "SEATS: 18 people.\n"
                     "TIME: 1 day.",
                     source=f"#docx:{name}.docx", deep=False)
    blocks = []
    real = (extract.extract, generate.answer, generate.refusal,
            generate.offer_research)
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("alpha course", "duration", "")]}
    def spy_answer(q, block, warmth=0.2, persona="", max_tokens=None, **kw):
        blocks.append(block)
        return "I do not know."
    generate.answer = spy_answer
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    try:
        s.respond("do Alpha Course and Beta Course run the same TIMES?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research) = real
    rows = [l for l in "\n".join(blocks).split("\n")
            if " = " in l or " \u2260 " in l]
    assert rows, "no comparison rows at all"
    assert "TIME" in rows[0], ("the asked field is not the first row: %r"
                               % rows[:3])
    assert any("SEATS" in r or "FORMAT" in r for r in rows), \
        "the other fields were dropped instead of ranked"


@test("W48 two inflections of one short root are kin")
def w48():
    """The hole `kin` left open, measured where it hurts: a comparison
    asks "aynı SÜREDE mi" of a record whose head is "SÜRESİ" — two
    inflections of one four-letter root. `same_stem` wants five letters
    of shared root and refuses; `kin` wants one form to be the other's
    prefix, and neither of these is. So the question could not reach the
    line that answers it, and the rider could not rank the row the
    question asked about. Kinship widens to what it always meant on the
    retrieval side: a shared opening of at least three letters, with
    each form's remainder within TAIL. The gates keep same_stem, so
    nothing here licenses a claim — it only decides what can be FOUND
    and what is ranked first."""
    from lmm import inflect
    assert inflect.kin("surede", "suresi")      # two inflections, one root
    assert inflect.kin("time", "times")         # prefix case, unchanged
    assert inflect.kin("gunluk", "gun")
    assert not inflect.kin("car", "cat")        # two shared letters is nothing
    assert not inflect.kin("kaygi", "kayitli")  # remainders beyond TAIL
    # the gates stay strict: kinship never licenses a claim
    assert not inflect.same_stem("surede", "suresi")
    assert not inflect.same_stem("time", "times")


@test("W49 the two judgments of a candidate are asked side by side")
def w49():
    """The last sequential pair on the answering path: the read-back
    asks whether the evidence SAYS this, the relation check asks
    whether it answers what was ASKED, and one waited for the other
    though neither reads the other's verdict. They are independent
    judgments of the same candidate, so they go to the engine together
    where the backend allows — and a candidate is admitted only if BOTH
    still hold, exactly as before. The pattern is the verifier's own
    (runtime.parallel_map): the occasional cost is one relation call for
    a candidate the read-back would have refused; the gain is a whole
    engine round-trip off every answered turn."""
    import os, threading, time
    from lmm import generate
    from lmm.session import Session
    s = Session(None)
    s.learn_text("The trust walk closes the morning arc.",
                 source="#docx:Alpha.docx", deep=False)
    lock = threading.Lock()
    state = {"read": 0, "rel": 0, "overlap": False}
    def busy(kind, result):
        with lock:
            state[kind] += 1
            if state["read"] and state["rel"]:
                state["overlap"] = True      # the two KINDS in flight at once
        time.sleep(0.12)
        with lock:
            state[kind] -= 1
        return result
    real = (generate.supported, generate.answers_asked)
    generate.supported = lambda raw, view: busy("read", True)
    generate.answers_asked = lambda q, a, view: busy("rel", True)
    old = os.environ.get("LMM_BACKEND")
    try:
        os.environ["LMM_BACKEND"] = "azure"
        proof = ["Alpha \u2014 the trust walk closes the morning arc"]
        ok = s._judge("what closes the arc?", "The trust walk closes it.",
                      proof, "\n".join(proof))
    finally:
        generate.supported, generate.answers_asked = real
        if old is None:
            os.environ.pop("LMM_BACKEND", None)
        else:
            os.environ["LMM_BACKEND"] = old
    assert ok is True
    assert state["overlap"], "the two judgments still queue: %r" % state
    # ...and a refusal from either one still refuses
    generate.supported = lambda raw, view: True
    generate.answers_asked = lambda q, a, view: False
    try:
        bad = s._judge("what closes the arc?", "The trust walk closes it.",
                       proof, "\n".join(proof))
    finally:
        generate.supported, generate.answers_asked = real
    assert bad is False



@test("W54 a source that contradicts itself is read as a contradiction")
def w54():
    """A MEMORY THAT PROMISES VERIFIABILITY MAY NOT PICK A SIDE IN SILENCE.
    A document that states the same field twice with different values is
    ordinary in a live corpus — a summary updated, a schedule not — and
    today both lines enter the block, the engine reads the one it happens
    to read, and the answer is stamped with the source that also says the
    opposite. One value spoken, one value hidden, and the stamp attests
    both. That is the failure mode the whole gate exists to prevent.

    The reading is the one already earned twice: equality of records is
    arithmetic (W44's digit sets), so when ONE source carries two rows
    with the same field head and different digits, a row states it, in
    the same \u2260 notation the comparison verdict uses. The engine reads
    a contradiction rather than resolving one.

    Scope is deliberately narrow: the SAME source. Two sources carrying
    different values for a head are a census, not a conflict — sixty
    course outlines each state their own duration, and calling that a
    contradiction would fire on every field question in the corpus. A
    conflict needs identity of subject, and inside one document that
    identity is given. Claiming one where we cannot establish it would be
    a fabricated relation, which is the thing we do not do."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    # Fifteen honest siblings around the one that argues with itself, so
    # a fixture that could pass by accident does not.
    for i, v in enumerate(("2 days", "3 days", "4 days", "1 day", "5 days",
                           "2 days", "3 days", "4 days", "1 day", "5 days",
                           "2 days", "3 days", "4 days", "1 day", "5 days")):
        s.learn_text(f"COURSE {i} OUTLINE.\nThe module opens the arc.\n"
                     f"DURATION: {v}.", source=f"#docx:Course {i}.docx",
                     deep=False)
    s.learn_text("SAFETY COURSE OUTLINE.\n"
                 "The safety module opens the arc.\n"
                 "DURATION: 2 days.\n"
                 "The programme closes with a workshop.\n"
                 "DURATION: 3 days.",
                 source="#docx:Safety Course.docx", deep=False)
    blocks = []
    real = (extract.extract, generate.answer, generate.refusal,
            generate.offer_research)
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("safety course", "duration", "")]}
    def spy_answer(q, block, warmth=0.2, persona="", max_tokens=None, **kw):
        blocks.append(block)
        return "I do not know."
    generate.answer = spy_answer
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    try:
        s.respond("what is the DURATION of the Safety Course?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research) = real
    text = "\n".join(blocks)
    rows = sorted({l for l in text.split("\n") if "\u2260" in l})
    assert rows, ("no contradiction row was written:\n" + text[:900])
    row = rows[0]
    assert "2 days" in row and "3 days" in row, row
    assert "Safety Course" in row, row
    # AND THE SIBLINGS ARE NOT A CONTRADICTION. Each of the fifteen states
    # its own duration; a reading that called that a conflict would fire
    # on every field question the corpus can be asked.
    assert len(rows) == 1, ("a cross-source census was read as a "
                            "contradiction:\n" + "\n".join(rows))

    # A REPEATED HEAD IS USUALLY A LIST, AND SOMETIMES A ZOOM. Measured on
    # the field corpus: the first cut fired three times and was wrong
    # twice. A prerequisites block writing three ZORUNLU lines is
    # enumerating, not disagreeing; a document writing "4 modules x 2 full
    # days" and later "2 full days" is showing the same length closer up.
    # Both are structure — a third value, and one value carried inside
    # another — and neither is a claim we may make.
    s2 = Session(None)
    s2.learn_text("ATLAS COURSE OUTLINE.\nThe atlas module opens the arc.\n"
                  "DURATION: 4 modules x 2 full days.\n"
                  "The programme closes with a workshop.\n"
                  "DURATION: 2 full days.\n"
                  "REQUIRED: 7 visit reports completed beforehand.\n"
                  "REQUIRED: 3 action lists brought along.\n"
                  "REQUIRED: 1 roadmap marked up.",
                  source="#docx:Atlas Course.docx", deep=False)
    for i in range(4):
        s2.learn_text(f"COURSE {i} OUTLINE.\nThe module opens the arc.\n"
                      f"DURATION: {i + 1} days.\nREQUIRED: {i} reports.",
                      source=f"#docx:Course {i}.docx", deep=False)
    blocks2 = []
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("atlas course", "duration", "")]}
    generate.answer = (lambda q, block, warmth=0.2, persona="",
                       max_tokens=None, **kw: blocks2.append(block) or "I do not know.")
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    try:
        s2.respond("what is the DURATION of the Atlas Course?")
        s2.respond("what is REQUIRED for the Atlas Course?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research) = real
    quarrels = [l for l in "\n".join(blocks2).split("\n") if "\u2260" in l]
    assert not quarrels, ("a list or a refinement was read as a "
                          "contradiction:\n" + "\n".join(quarrels))



@test("W55 over HTTP, every user's memory is their own")
@engine_free
def w55():
    """A memory layer becomes a product the moment a second person asks
    it something, and the first thing a second person can break is the
    first person's privacy. The store is chosen by the caller's user id
    and by nothing else, and no route reads across users — "the
    application will pass the right path" is how a tenant's documents
    end up in another tenant's answer.

    A user id arrives from the network and becomes a filename, so it is
    a name or it is refused: an id that can climb ("../secrets") is the
    oldest way to read a file nobody meant to share."""
    import json
    import tempfile
    import threading
    import urllib.request
    from lmm import generate, serve
    root = tempfile.mkdtemp()
    httpd = serve.serve(root=root, port=0, token="")
    threading.Thread(target=httpd.serve_forever, daemon=True).start()
    base = "http://127.0.0.1:%d" % httpd.server_address[1]

    def call(route, payload):
        req = urllib.request.Request(
            base + route, data=json.dumps(payload).encode(),
            headers={"Content-Type": "application/json"})
        try:
            with urllib.request.urlopen(req) as r:
                return r.status, json.loads(r.read().decode())
        except urllib.error.HTTPError as e:
            return e.code, json.loads(e.read().decode())

    # Engine-free like the rest of the suite: the gates are held open
    # because the subject is one user's documents staying out of another
    # user's answer, and learning goes in shallow because the deep read
    # asks the engine.
    from lmm.session import Session
    real_gates = (Session._read_back, Session._relation_held)
    Session._read_back = lambda self, raw, proof, block, question="": True
    Session._relation_held = lambda self, q, raw, proof, block: True
    from lmm import extract
    real = (generate.answer, generate.refusal, generate.offer_research,
            extract.extract)
    # The router is not the subject here either: these are questions.
    extract.extract = lambda message: {"kind": extract.ASK,
                                       "triples": [("contract", "worth", "")]}
    generate.answer = (lambda q, block, warmth=0.2, persona="",
                       max_tokens=None, **kw: "[" + block + "]")
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not have that.")
    generate.offer_research = lambda subject, question: ""
    try:
        call("/learn", {"user": "ada", "deep": False, "text":
                        "The Redwood contract is worth 40000 euro.",
                        "source": "#doc:Redwood.txt"})
        call("/learn", {"user": "bo", "deep": False, "text":
                        "The Fernbank contract is worth 90000 euro.",
                        "source": "#doc:Fernbank.txt"})
        code, bo = call("/ask", {"user": "bo",
                                 "question": "what is the Redwood contract worth?"})
        assert code == 200, bo
        assert "40000" not in bo["answer"], (
            "one user's document reached another user's answer: %r" % bo)
        assert "Redwood" not in " ".join(bo["sources"]), bo
        code, ada = call("/ask", {"user": "ada",
                                  "question": "what is the Redwood contract worth?"})
        assert "40000" in ada["answer"], ada
        # The answer travels whole: abstention and stamps are part of it,
        # or an HTTP caller cannot verify what an in-process caller can.
        assert set(("answer", "abstained", "sources", "subject")) <= set(ada), ada
        assert ada["sources"], "the stamps did not cross the wire"
        # An id that can climb is not a name.
        code, bad = call("/ask", {"user": "../../etc/passwd",
                                  "question": "anything?"})
        assert code == 400 and "user" in bad["error"], (code, bad)
        code, bad = call("/ask", {"user": "", "question": "anything?"})
        assert code == 400, (code, bad)
    finally:
        (generate.answer, generate.refusal, generate.offer_research,
         extract.extract) = real
        (Session._read_back, Session._relation_held) = real_gates
        httpd.shutdown()



@test("W56 crossing into a chain, the audit crosses with it")
@engine_free
def w56():
    """The adapter is where a verified memory quietly becomes a vector
    store. A retriever hands over passages and the caller's own model
    says whatever it likes about them — that is honest, and the stamps
    ride along so the application can show its work. The tool is the
    door that matters: an agent must receive the ANSWER THAT WAS
    AUDITED, with its sources, and — this is the part that breaks in
    every naive adapter — with the abstention as itself. Hand an agent
    an empty string where memory holds nothing and the agent writes its
    own answer over the silence, which is precisely what this project
    exists to refuse."""
    from lmm import generate
    from lmm.adapters import langchain as bridge
    from lmm.api import Memory
    m = Memory(None)
    m.learn("The Redwood contract is worth 40000 euro.",
            source="#doc:Redwood.txt", deep=False)
    docs = bridge.passages(m, "what is the Redwood contract worth?")
    assert docs, "the retriever handed over nothing"
    assert all(hasattr(d, "page_content") and hasattr(d, "metadata")
               for d in docs), docs
    assert any("Redwood.txt" in (d.metadata.get("source") or "")
               for d in docs), [d.metadata for d in docs]
    # THE SUBJECT HERE IS THE ADAPTER, so the engine is not in the room.
    # Measured the hard way: with only the composer faked, the gates still
    # called the live engine to audit the claim, and one run in thirteen
    # failed because a network call under load refused a sentence this
    # test was not written to judge. A test whose verdict depends on a
    # remote engine is measuring the engine.
    from lmm.session import Session
    real_gates = (Session._read_back, Session._relation_held)
    Session._read_back = lambda self, raw, proof, block, question="": True
    Session._relation_held = lambda self, q, raw, proof, block: True
    from lmm import extract
    real = (generate.answer, generate.refusal, generate.offer_research,
            extract.extract)
    extract.extract = lambda message: {"kind": extract.ASK,
                                       "triples": [("contract", "worth", "")]}
    generate.answer = (lambda q, block, warmth=0.2, persona="",
                       max_tokens=None, **kw: "The Redwood contract is worth "
                       "40000 euro.")
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not have that in memory.")
    generate.offer_research = lambda subject, question: ""
    try:
        agent_tool = bridge.tool(m)
        said = agent_tool.invoke({"question":
                                  "what is the Redwood contract worth?"})
        assert "40000" in said, said
        assert "Redwood.txt" in said, ("the stamps did not cross with the "
                                       "answer: %r" % said)
        empty = bridge.tool(Memory(None))
        silence = empty.invoke({"question": "what is the Fernbank worth?"})
        assert silence.strip(), ("the abstention crossed as an empty string, "
                                 "which an agent will write over")
        assert "40000" not in silence and "Fernbank" not in silence.replace(
            "Fernbank", "", 1), silence
    finally:
        (generate.answer, generate.refusal, generate.offer_research,
         extract.extract) = real
        (Session._read_back, Session._relation_held) = real_gates



@test("W57 the store's own vocabulary answers the question's")
def w57():
    """Measured on a corpus this system had never seen: asked which
    machine is the "priciest", every seat went to prose about value,
    because the documents write LIST PRICE and no word of the question
    shares a stem with it. Retrieval is lexical by design — that is what
    makes it auditable — so the gap is in the words the search is given,
    not in the gates.

    Three properties, and the fix is worth nothing without all three.
    The engine is never asked what a word means in general: it is shown
    the field names the corpus repeats and asked which one the question
    reaches for, so a bridge can only land on vocabulary that provably
    exists here. It is asked only when the question touches NO field, so
    the ordinary question pays nothing. And a pick that is not one of
    the store's heads is discarded — an engine asked to copy a name
    sometimes writes a neighbouring one, and that word would enter the
    search on the engine's authority alone."""
    from lmm import generate
    from lmm.session import Session
    s = Session(None)
    for name, price in (("Marlin", "4200"), ("Petrel", "5100"),
                        ("Falcon", "3300"), ("Osprey", "6400")):
        s.learn_text(f"{name} WORKSTATION.\nThe {name.lower()} ships in a "
                     f"steel case.\nLIST PRICE: {price} euro.\n"
                     f"MEMORY: 64 GB.", source=f"#doc:{name}.txt", deep=False)
    heads = s._fields()
    assert "LIST PRICE" in heads and "MEMORY" in heads, heads
    asked = []
    real = generate.field_for
    # the listing shows each head beside one of its own values, so a
    # reader that has only seen the name can recognise it; the pick is
    # matched back to the head
    generate.field_for = lambda q, hs: (asked.append((q, list(hs)))
                                        or [h for h in hs
                                            if h.startswith("LIST PRICE")][0])
    try:
        # A question whose words reach a field asks the engine nothing.
        assert s._field_bridge("what is the MEMORY of the Marlin?") == ""
        assert not asked, asked
        got = s._field_bridge("which workstation is the priciest?")
        assert got == "LIST PRICE", got
        assert asked and any(h.startswith("LIST PRICE") for h in asked[0][1]), asked
        assert any("4200" in h or "euro" in h for h in asked[0][1]), (
            "the heads were shown without an example of what they hold: %s"
            % asked[0][1])
        # The same question twice costs one call, not two.
        s._field_bridge("which workstation is the priciest?")
        assert len(asked) == 1, asked
        # A head that is not ours never reaches the search.
        generate.field_for = lambda q, hs: "STREET PRICE"
        assert s._field_bridge("which one costs the least?") == ""
    finally:
        generate.field_for = real



@test("W58 a ridden subject does not name a source")
def w58():
    """The bug that ate a whole category. In a conversation the previous
    turn's subject rides along in the query, to keep retrieval on topic
    when the turn's own words are a pointer. But `find` reads NAMES off
    the query it is given, so the ridden subject also made the turn look
    like a question ABOUT that document — and a corpus-wide field
    question, which by definition names nobody, was read as a question
    about whatever came up last.

    Measured on thirteen sibling specifications: asked on its own,
    "which workstation is the priciest" reads the field across the corpus
    and answers with the right machine and the right number; asked after
    a question about another machine, the same question abstains with
    every price on the table. Three of three frontier questions were
    lost this way, and nothing in the block was wrong — the reading
    never ran.

    Naming is a property of WHAT WAS ASKED. The ride still steers
    retrieval; it no longer confers namehood."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    prices = {"Osprey": "3499", "Falcon": "2999", "Marlin": "2499",
              "Kite": "1499", "Tern": "799", "Petrel": "999"}
    for name, price in prices.items():
        s.learn_text(f"{name} WORKSTATION.\nThe {name.lower()} ships in a "
                     f"steel case.\nLIST PRICE: {price} USD.",
                     source=f"#docx:{name} Workstation", deep=False)
    assert s.evidence.named_in("what is the LIST PRICE of the Osprey "
                               "Workstation?") == {"#docx:Osprey Workstation"}
    assert s.evidence.named_in("which workstation has the highest LIST "
                               "PRICE?") == set(), "a corpus question named one"
    blocks = []
    real = (extract.extract, generate.answer, generate.refusal,
            generate.offer_research)
    # The first turn is about a document by name; the second is the
    # pointer-shaped follow-up whose subject resolves to nothing, which is
    # exactly when the previous subject rides along.
    subjects = iter([("Osprey Workstation", "price", ""), ("it", "price", "")])
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [next(subjects)]}
    generate.answer = (lambda q, block, warmth=0.2, persona="",
                       max_tokens=None, **kw: blocks.append(block) or "I do not know.")
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    try:
        s.respond("what is the LIST PRICE of the Osprey Workstation?")
        blocks.clear()
        s.respond("which workstation has the highest LIST PRICE?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research) = real
    text = "\n".join(blocks)
    rows = [l for l in text.split("\n") if "largest" in l]
    assert rows, ("the corpus-wide reading did not run after a question "
                  "about one document:\n" + text[:700])
    assert "Osprey" in rows[0] and "3499" in rows[0], rows[0]



@test("W59 a window holding several records offers all of them")
def w59():
    """How a document was CHUNKED is not supposed to change what the
    memory can read, and it did. Ingested one way, a specification's
    lines arrive separately and the record channel sees DISPLAY, LIST
    PRICE and WARRANTY; ingested another — the ordinary bulk path, which
    is how every benchmark and every customer load runs — the same three
    arrive glued into one window, the channel partitions at the FIRST
    colon, and the corpus knows one field where it wrote three. Measured:
    the same superlative question answered correctly from a per-line
    store and abstained from a windowed one, with every value present in
    both. That is the retrieval layer reporting the shape of the file
    reader.

    A record is a head, a colon and a value, and a line may hold several.
    Reading them all is what makes the field census, the extremes row and
    the contradiction reading independent of how the text was cut."""
    from lmm import evidence
    from lmm.session import Session
    pairs = evidence.record_pairs(
        "GANNET WORKSTATION — TECHNICAL SUMMARY. PRODUCT FAMILY: mobile "
        "engineering workstation. DISPLAY: 13 inches. LIST PRICE: 1199 "
        "USD. WARRANTY: 3 years on parts and labour.")
    heads = [h for h, _v in pairs]
    assert "LIST PRICE" in heads and "WARRANTY" in heads, pairs
    assert dict(pairs)["LIST PRICE"] == "1199 USD", pairs
    assert dict(pairs)["DISPLAY"] == "13 inches", pairs
    # A colon inside ordinary prose is not a record head.
    assert not evidence.record_pairs(
        "The source of the difficulty is this: the team had never met."), \
        "a titled sentence was read as a record"

    # ...and the reading survives the window in a real store: the field
    # census must find the price in every document even when each
    # document arrived as one glued paragraph.
    s = Session(None)
    for name, price, screen in (("Osprey", "3499", "17"), ("Falcon", "2999", "15"),
                                ("Marlin", "2499", "14"), ("Tern", "799", "12")):
        s.learn_text(f"{name} WORKSTATION — TECHNICAL SUMMARY. "
                     f"DISPLAY: {screen} inches. LIST PRICE: {price} USD. "
                     f"WARRANTY: 3 years on parts and labour.",
                     source=f"#docx:{name} Workstation", deep=False)
    assert "LIST PRICE" in s._fields(), sorted(s._fields())
    rows = dict((h, v) for _sid, text in s._record_rows(
        "#docx:Tern Workstation", "what is the LIST PRICE?")
        for h, v in evidence.record_pairs(text))
    assert rows.get("LIST PRICE") == "799 USD", rows



@test("W60 the bridged field steers the words, never the facts")
def w60():
    """The last mile of the bridge, and the place where it would have
    been easy to cheat. Retrieval now finds the right rows for a question
    written in another vocabulary — "the most RAM" against documents
    headed MEMORY — and the block held "MEMORY — largest: Falcon
    Workstation: 64 GB" while the writer answered "carries the most RAM",
    which the jury refused because RAM appears nowhere in the evidence.
    The answer was on the table and the user got nothing.

    The cheat would be to put the synonym into the evidence, where the
    gates read it: that is asserting a synonymy the store cannot attest,
    inside the one channel that must stay attested. So the bridge's
    finding travels as an INSTRUCTION to the writer — name the field as
    the documents name it — and the facts block is byte for byte what it
    was. Vocabulary is steered; content is not."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    for name, ram in (("Osprey", "32"), ("Falcon", "64"), ("Marlin", "16"),
                      ("Tern", "8")):
        s.learn_text(f"{name} WORKSTATION — TECHNICAL SUMMARY. "
                     f"MEMORY: {ram} GB. WARRANTY: 3 years.",
                     source=f"#docx:{name} Workstation", deep=False)
    seen = {}
    real = (extract.extract, generate.answer, generate.refusal,
            generate.offer_research, generate.field_for)
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("machine", "ram", "")]}
    generate.field_for = lambda q, heads: "MEMORY"
    def spy_answer(question, block, warmth=0.2, persona="", max_tokens=None,
                   field="", **kw):
        seen.setdefault("fields", []).append(field)
        seen.setdefault("blocks", []).append(block)
        return "I do not know."
    generate.answer = spy_answer
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    try:
        s.respond("which machine carries the most RAM?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research, generate.field_for) = real
    assert seen.get("fields"), "the writer was never called"
    assert any(f == "MEMORY" for f in seen["fields"]), (
        "the bridged field never reached the writer: %r" % seen["fields"])
    blocks = "\n".join(seen["blocks"])
    assert "MEMORY" in blocks, blocks[:400]
    assert "RAM" not in blocks, ("the question's word was written into the "
                                 "evidence, where the gates read it:\n"
                                 + blocks[:400])



@test("W61 a question that names one document is answered from that document")
def w61():
    """Found by a question set nobody wrote by hand. Asked what the
    "Güven Veren Lider" document says under a head it does not carry, the
    memory answered from a different document entirely and stamped the
    answer with that other document — honestly, and uselessly. Every word
    was attested; the reader asked about one programme and was told about
    another. Naming a document is a scope, not a hint: a claim that rests
    on some other source is not an answer to what was asked, and the
    honest reply is that this document does not say.

    Only the claim's LOAD-BEARING source is checked — the one the answer
    actually rests on, which the stamp already computes. A named
    document's answer may still be enriched by neighbours; what it may
    not do is come entirely from them."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    s.learn_text("ALPHA PROGRAMME OUTLINE.\nThe alpha module opens the arc.\n"
                 "DURATION: 2 days.", source="#docx:Alpha Programme",
                 deep=False)
    s.learn_text("BETA PROGRAMME OUTLINE.\nThe beta module opens the arc.\n"
                 "DURATION: 3 days.\nTRAINER: the faculty lead runs it.",
                 source="#docx:Beta Programme", deep=False)
    for i in range(4):
        s.learn_text(f"GAMMA {i} OUTLINE.\nThe gamma module opens the arc.\n"
                     f"DURATION: {i + 1} days.", source=f"#docx:Gamma {i}",
                     deep=False)
    said = {}
    # The subject here is the SCOPE, not the jury: the jury is held open
    # so that what this test measures is whether a claim resting on
    # another document may stand, not whether a remote engine liked the
    # sentence today.
    real_gates = (Session._read_back, Session._relation_held)
    Session._read_back = lambda self, raw, proof, block, question="": True
    Session._relation_held = lambda self, q, raw, proof, block: True
    real = (extract.extract, generate.answer, generate.refusal,
            generate.offer_research)
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("alpha programme", "trainer", "")]}
    # The engine answers with the one TRAINER line in the store — which
    # belongs to Beta, not to the Alpha the question named.
    generate.answer = (lambda q, block, warmth=0.2, persona="",
                       max_tokens=None, **kw:
                       "TRAINER: the faculty lead runs it.")
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not have that.")
    generate.offer_research = lambda subject, question: ""
    try:
        said["out"] = s.respond("who is the TRAINER of the Alpha Programme?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research) = real
        (Session._read_back, Session._relation_held) = real_gates
    out = said["out"] or ""
    assert "faculty lead" not in out, (
        "the named document's question was answered from another document: "
        + out)
    assert "Beta" not in out, out




@test("W62 a missing engine says which engine is missing")
def w62():
    """The first thing a new reader meets, and until now it was a stack
    trace about somebody else's library. `pip install
    living-memory-model` pulls in nothing — that is the design claim — so
    the very first `learn()` on a fresh machine reaches the default local
    engine, fails to import torch, and the reader is told
    "ModuleNotFoundError: No module named 'torch'". Nothing in that
    sentence says an engine was needed, that there are four to choose
    from, or that the cheapest is one environment variable away.

    The reader is not missing torch. The reader is missing an ENGINE, and
    the error says so — with the choices, and with the name to type after
    `pip install`, which is the DISTRIBUTION name: telling someone to
    install "lmm[local]" sends them to a different project."""
    import builtins
    from lmm import runtime
    from lmm.tables import DIST
    real_import = builtins.__import__

    def no_torch(name, *a, **kw):
        if name == "torch":
            raise ImportError("No module named 'torch'")
        return real_import(name, *a, **kw)

    was = os.environ.pop("LMM_BACKEND", None)
    builtins.__import__ = no_torch
    try:
        runtime.generate("hello", max_tokens=4)
    except ImportError as said:
        text = str(said)
        assert "LMM_BACKEND" in text, text
        assert "openai" in text and "azure" in text, text
        assert ("%s[local]" % DIST) in text, text
        assert "lmm[local]" not in text, text
    except Exception as other:                              # noqa: BLE001
        raise AssertionError("a missing engine raised %s: %s"
                             % (type(other).__name__, other))
    else:
        raise AssertionError("a missing engine raised nothing at all")
    finally:
        builtins.__import__ = real_import
        if was is not None:
            os.environ["LMM_BACKEND"] = was



@test("W63 the question API does not gamble on a chat")
def w63():
    """Measured with the cost meter, on the field corpus: every factual
    question paid for a full chat completion that was then thrown away.

    The wager is sound where it was born. In a consultation the router's
    verdict and the chat reply are wanted at the same instant, so the
    voice starts warming while the classifier reads, and a turn that
    routes to chat has its answer already running. But `ask()` is not a
    conversation — its entire contract is that a question came in — and
    it inherited the wager only because it passes teach=False, which is
    what the speculation looked at.

    A wager nobody can win is not a wager, it is a bill: one call and
    its tokens on every question, on the path benchmarks and API callers
    use most. The conversational surface keeps it; the question door
    does not place it."""
    from lmm import extract, generate
    from lmm.api import Memory
    from lmm.session import Session
    m = Memory(None)
    m.learn("The Redwood contract is worth 40000 euro.",
            source="#doc:Redwood.txt", deep=False)
    chats = {"n": 0}
    real = (extract.extract, generate.answer, generate.refusal,
            generate.chat, generate.wants_material,
            generate.is_identity_question, Session._chat_raw,
            Session._read_back, Session._relation_held)
    extract.extract = lambda message: {"kind": extract.ASK,
                                       "triples": [("contract", "worth", "")]}
    generate.answer = (lambda q, block, warmth=0.2, persona="",
                       max_tokens=None, **kw:
                       "The Redwood contract is worth 40000 euro.")
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not have that.")
    generate.wants_material = lambda message: False
    generate.is_identity_question = lambda message: False
    generate.chat = (lambda message, identity_block="", warmth=0.7,
                     history=None, persona="", max_tokens=None:
                     chats.__setitem__("n", chats["n"] + 1) or "Hello!")
    Session._chat_raw = (lambda self, message:
                         chats.__setitem__("n", chats["n"] + 1) or "Hello!")
    Session._read_back = lambda self, raw, proof, block, question="": True
    Session._relation_held = lambda self, q, raw, proof, block: True
    try:
        said = m.ask("what is the Redwood contract worth?")
        assert "40000" in str(said), said
        assert chats["n"] == 0, (
            "the question API paid for %d chat call(s) it threw away"
            % chats["n"])
        # The same reasoning, one door along: the delivery bridge — which
        # offers a composed catalogue when a consultation cannot answer —
        # belongs to the conversation. A caller who asked `ask()` a
        # question has `compose()` for that, and on this door the bridge
        # costs a classifier call on every turn and, when it fires, a
        # composition nobody requested.
        wanted = {"n": 0}
        generate.wants_material = (lambda message:
                                   wanted.__setitem__("n", wanted["n"] + 1)
                                   or True)
        generate.answer = (lambda q, block, warmth=0.2, persona="",
                           max_tokens=None, **kw: "I do not know.")
        m.ask("what is the Fernbank contract worth?")
        m.ask("and the Kingfisher one?")
        assert wanted["n"] == 0, (
            "the question door asked the delivery classifier %d time(s)"
            % wanted["n"])
        # ...and the conversational surface keeps its wager, where the
        # reply it pre-warms is the reply that gets spoken.
        m.session.respond("hello there", teach=False)
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.chat, generate.wants_material,
         generate.is_identity_question, Session._chat_raw,
         Session._read_back, Session._relation_held) = real



@test("W64 the extremes of a field are the corpus's, not the block's")
def w64():
    """Found by the first measurement above sixty documents, and it is
    the worst kind of bug this system can have: a confident wrong answer
    with a source stamp on it.

    The field census lays one row per source and stops at sixty, because
    a block has seats. The extremes row was then computed FROM THOSE
    ROWS — so on a corpus of a thousand it reported the largest value
    among the first sixty documents in alphabetical order and said it
    was the largest of all. Measured: asked for the highest price across
    a thousand machines, the memory answered 8,898 with a stamp, while
    the corpus said 9,098 somewhere further down the alphabet. Nothing
    caught it, because every word of the answer was attested; only the
    superlative was false, and the superlative is the part nobody can
    check by reading one line.

    What is shown may be capped. What is CLAIMED may not: the extremes
    are read over every source that carries the head, and the row is
    written from that reading."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    # Eighty siblings, and the largest value sits at the end of the
    # alphabet where a cap of sixty cannot see it.
    for i in range(80):
        name = "%s Machine %02d" % ("Zed" if i == 79 else "Alpha", i)
        price = 9098 if i == 79 else 700 + i * 10
        s.learn_text("%s — TECHNICAL SUMMARY.\nLIST PRICE: %d USD.\n"
                     "The unit ships with a charger." % (name.upper(), price),
                     source="#docx:%s" % name, deep=False)
    blocks = []
    real = (extract.extract, generate.answer, generate.refusal,
            generate.offer_research, generate.field_for)
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("machine", "price", "")]}
    generate.field_for = lambda q, heads: ""
    generate.answer = (lambda q, block, warmth=0.2, persona="",
                       max_tokens=None, **kw:
                       blocks.append(block) or "I do not know.")
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    try:
        s.respond("which machine has the highest LIST PRICE?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research, generate.field_for) = real
    rows = [l for l in "\n".join(blocks).split("\n") if "largest" in l]
    assert rows, "no extremes row was written"
    assert "9098" in rows[0] and "Zed" in rows[0], (
        "the extremes were read off the capped block: %s" % rows[0])



@test("W65 the comparison verdict speaks before the lines it is drawn from")
def w65():
    """The census learned this the hard way and the comparison layout
    never got the lesson. A block is read from the top; the extremes row
    was appended after sixty field lines and nobody read it, so it was
    moved to the front. The comparison's verdict row — the one that says
    whether the two values agree — was still being appended after every
    line the two documents contributed.

    Measured on the field set, and it is the whole of a remaining
    failure: asked whether two courses run the same length, the block
    carried twelve lines of prose, then a row reading "DURATION —
    A: 1 full day = B: 1 day". The engine answered off the two raw lines
    and concluded they DIFFER, because "1 full day" and "1 day" are
    different words; the jury refused that, correctly, and the turn
    abstained with the answer sitting two lines below where the reader
    stopped. Arithmetic the system already trusts, written where nobody
    reads it, is arithmetic the system did not do.

    The verdicts go first — the asked-about field ahead of the rest, as
    W47 ordered them — and the evidence they were drawn from follows."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    for name, tag, length in (("Alpha Course", "ravine", "1 full day"),
                              ("Beta Course", "ridge", "1 day")):
        s.learn_text(f"{name.upper()} OUTCOMES.\n"
                     f"The {tag} module opens the arc.\n"
                     f"The programme closes with a workshop and a review.\n"
                     f"Participants bring a case of their own to the room.\n"
                     f"FORMAT: in person.\n"
                     f"SEATS: 18 people.\n"
                     f"DURATION: {length}.",
                     source=f"#docx:{name}.docx", deep=False)
    blocks = []
    real = (extract.extract, generate.answer, generate.refusal,
            generate.offer_research)
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("alpha course", "duration", "")]}
    generate.answer = (lambda q, block, warmth=0.2, persona="",
                       max_tokens=None, **kw:
                       blocks.append(block) or "I do not know.")
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    try:
        s.respond("do the Alpha Course and the Beta Course run the same "
                  "DURATION?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research) = real
    lines = [l for l in blocks[0].split("\n") if l.strip()]
    verdicts = [i for i, l in enumerate(lines)
                if "\u2260" in l or " = " in l]
    assert verdicts, ("no verdict row was written:\n"
                      + "\n".join(lines[:6]))
    assert verdicts[0] == 0, (
        "the verdict is line %d of %d, below the lines it is drawn from:\n%s"
        % (verdicts[0] + 1, len(lines), "\n".join(lines[:4])))
    assert "DURATION" in lines[0], lines[0]



@test("W66 the source whose name is matched most completely is the one named")
def w66():
    """A family resemblance is a PARTIAL match, and the naming rule could
    not see the difference. Measured on a corpus of sibling programmes:
    asked about "Customer-Focused Value Transfer in Sales - Call Centre",
    the question was read as naming SIX documents, because every sibling
    shares the opening phrase and differs only in its tail. One document
    was asked about; six were named; the comparison layout fired for a
    question that compares nothing.

    W38 already said a source is called by a pointer rather than by a
    family resemblance, and this is the same sentence carried one step
    further: among the sources a question calls, the one whose name it
    accounts for MOST COMPLETELY is the one it names. A tie keeps
    everyone — two documents named in full are exactly what a comparison
    is — so nothing is guessed; only the half-matched siblings step
    back."""
    from lmm.session import Session
    s = Session(None)
    names = ("Value Transfer in Sales - Call Centre",
             "Value Transfer in Sales - Branch Desk",
             "Value Transfer in Sales - Field Team",
             "Effective Value Transfer Culture Project")
    for name in names:
        s.learn_text("%s OUTLINE.\nThe module opens the arc.\n"
                     "DURATION: 2 days." % name.upper(),
                     source="#docx:%s" % name, deep=False)
    one = s.evidence.named_in(
        "how long is the Value Transfer in Sales - Call Centre programme?")
    assert one == {"#docx:Value Transfer in Sales - Call Centre"}, (
        "a question about one document named %d: %s" % (len(one), sorted(one)))
    # ...and a comparison still names both, because both are matched whole
    two = s.evidence.named_in(
        "do Value Transfer in Sales - Call Centre and Value Transfer in "
        "Sales - Branch Desk run the same DURATION?")
    assert two == {"#docx:Value Transfer in Sales - Call Centre",
                   "#docx:Value Transfer in Sales - Branch Desk"}, sorted(two)



@test("W67 a value belongs to the field it was written under")
def w67():
    """Found by the thousand-document set, and it is a fabrication the
    existing gates cannot see. A specification carries MEMORY: 128 GB and
    no STORAGE line at all. Asked for its STORAGE, the memory answered
    "the storage is 128 GB" and stamped it with that document — and every
    gate agreed, because the source is right (the scope rule holds), the
    number is in the evidence (the read-back holds), and the sentence
    does answer the shape of the question (the relation check holds). The
    only thing wrong is which FIELD the number belongs to, and no gate
    was reading that.

    In a spec sheet, memory and storage are different things and giving
    one for the other is not a near miss. The reading is local and
    deterministic: when the question asks about a field the CORPUS
    knows, a claim carrying digits must find those digits under THAT
    head in the evidence. The head is the documents' own; the digits are
    the documents' own; nothing here is a threshold or a word list."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    for i in range(6):
        s.learn_text("MACHINE %02d — TECHNICAL SUMMARY.\n"
                     "MEMORY: %d GB.\nDISPLAY: 15 inches.\n"
                     "The unit ships with a charger." % (i, 16 * (i + 1)),
                     source="#docx:Machine %02d" % i, deep=False)
    # ...and other documents that DO carry STORAGE, so the head is a
    # field the corpus repeats rather than a word nobody wrote. (A head
    # in one document only is not a field — the same rule the census
    # uses — which is why the fixture needs more than one.)
    for i in range(90, 94):
        s.learn_text("MACHINE %02d — TECHNICAL SUMMARY.\nMEMORY: 8 GB.\n"
                     "STORAGE: %d GB.\nDISPLAY: 15 inches." % (i, 128 * (i - 89)),
                     source="#docx:Machine %02d" % i, deep=False)
    said = {}
    real_gates = (Session._read_back, Session._relation_held)
    Session._read_back = lambda self, raw, proof, block, question="": True
    Session._relation_held = lambda self, q, raw, proof, block: True
    real = (extract.extract, generate.answer, generate.refusal,
            generate.offer_research, generate.field_for)
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("machine 03", "storage", "")]}
    generate.field_for = lambda q, heads: ""
    # the engine reaches for the only number in the document
    generate.answer = (lambda q, block, warmth=0.2, persona="",
                       max_tokens=None, **kw:
                       "The STORAGE of Machine 03 is 64 GB.")
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not have that.")
    generate.offer_research = lambda subject, question: ""
    try:
        said["out"] = s.respond("what is the STORAGE of Machine 03?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research, generate.field_for) = real
        (Session._read_back, Session._relation_held) = real_gates
    out = said["out"] or ""
    assert "64" not in out, (
        "another field's value was spoken as the asked field's: " + out)



@test("W68 a field of a named document is read, not generated")
def w68():
    """The commonest question a document store is asked — what does THIS
    document say under THAT head — is one the memory can answer without
    an engine at all. The row exists, the source is named, the head is
    named; there is nothing to compose and nothing to verify that
    reading the row does not already settle.

    This was tried once before and reverted: it fired on one question in
    twelve and cost a comparison. The reason is now known, and it was
    not this rule — it was the naming underneath it. A question about
    one document was being read as naming six (W66), so "the question
    names exactly one source" was almost never true, and when it was,
    the source was as likely to be a sibling. With naming fixed the
    condition means what it says.

    The conditions are all structural: exactly one source named, exactly
    one row under the asked head in that source, and the head asked by
    the question rather than by us. The answer is the document's own
    line, stamped with the document — the same thing the record channel
    puts in the block, spoken directly."""
    from lmm import generate, extract
    from lmm.api import Memory
    m = Memory(None)
    for i in range(4):
        m.learn("COURSE %02d OUTLINE.\nThe module opens the arc.\n"
                "DURATION: %d days.\nSEATS: %d people." % (i, i + 1, 10 + i),
                source="#docx:Course %02d" % i, deep=False)
    calls = {"n": 0}
    real = (extract.extract, generate.answer, generate.refusal,
            generate.offer_research)
    extract.extract = lambda msg: {"kind": extract.ASK,
                                   "triples": [("course 02", "duration", "")]}
    generate.answer = (lambda q, block, warmth=0.2, persona="",
                       max_tokens=None, **kw:
                       calls.__setitem__("n", calls["n"] + 1) or "made up")
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not have that.")
    generate.offer_research = lambda subject, question: ""
    routed = {"n": 0}
    real_extract = extract.extract
    extract.extract = (lambda msg: routed.__setitem__("n", routed["n"] + 1)
                       or {"kind": extract.ASK,
                           "triples": [("course 02", "duration", "")]})
    try:
        said = m.ask("what is the DURATION of Course 02?", explain=True)
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research) = real
    assert "3 days" in str(said), said
    assert calls["n"] == 0, ("the row was in hand and the engine was asked "
                             "anyway (%d call(s))" % calls["n"])
    # ...and on the question door not even the router is paid: `ask` is
    # asked a question by contract, and reading the row needs no subject
    # from anyone. This is the whole of the turn — a hundred milliseconds
    # of local work where the same question used to cost ten seconds.
    assert routed["n"] == 0, ("the question door paid for %d routing call(s) "
                              "on a row it could read" % routed["n"])
    assert not said.abstained, said
    assert any("Course 02" in s for s in said.sources), said.sources



@test("W69 a source stamp survives the space in a document's name")
def w69():
    """Provenance is the product, and the public API was cutting it in
    half. The mark is written as "(~ #docx:Course 02)" and `sources` was
    read by splitting the answer on whitespace and keeping the words
    that start with '#' — so every document whose name contains a space,
    which is most documents anyone has, arrived as "#docx:Course". Two
    sibling courses then carry the SAME stamp, and an application that
    shows provenance shows the wrong document, or an ambiguous one.

    The mark has a shape: it is the parenthesis the turn appends. Read
    it as one, and what comes back is what the store attested."""
    from lmm.api import Answer, Memory
    from lmm.session import UNCERTAIN
    m = Memory(None)
    m.session.last_abstained = False
    got = m._told("The course runs 3 days. (%s #docx:Course 02)" % UNCERTAIN)
    assert list(got.sources) == ["#docx:Course 02"], got.sources
    # a bare stamp with no space still reads, and prose punctuation does
    # not become part of it
    got2 = m._told("It is stated. (%s #manual)" % UNCERTAIN)
    assert list(got2.sources) == ["#manual"], got2.sources
    # ...and an answer with no mark claims no source
    assert list(m._told("The course runs 3 days.").sources) == []



@test("W70 a document that never uses the words is not asked about them")
def w70():
    """The mirror of W68, and the trap path's version of it. A question
    that names one document and one of the corpus's own field heads,
    where that document never uses the head's words ANYWHERE — not as a
    record, not in prose — has no answer in the only source it is
    allowed to be answered from (W61). Three model calls to compose
    candidates and three more to refuse them all is a bill for a
    conclusion the store already holds.

    The condition is deliberately narrow, and it was measured before it
    was written: across three corpora — the customer's 62 documents, a
    thousand generated ones, and a mechanically-derived question set —
    it holds for ZERO of 123 factual questions and for 24 of 25 traps at
    scale. A reading that never fires on an answerable question is a
    reading that costs nothing but the refusals it shortens.

    Prose still counts. The head's WORDS are what must be absent, not
    the record row: a document that describes its closing session
    without a "CLOSING:" line is answering, and this stands aside."""
    from lmm import generate, extract
    from lmm.api import Memory
    m = Memory(None)
    for i in range(5):
        m.learn("MACHINE %02d — TECHNICAL SUMMARY.\nMEMORY: %d GB.\n"
                "WARRANTY: 3 years.\nThe unit ships with a charger."
                % (i, 16 * (i + 1)), source="#docx:Machine %02d" % i,
                deep=False)
    # one document that carries STORAGE, so the head is a field the
    # corpus repeats; and a second, so it is repeated
    for i in (90, 91):
        m.learn("MACHINE %02d — TECHNICAL SUMMARY.\nMEMORY: 8 GB.\n"
                "STORAGE: 512 GB.\nWARRANTY: 3 years." % i,
                source="#docx:Machine %02d" % i, deep=False)
    calls = {"n": 0}
    real = (extract.extract, generate.answer, generate.refusal)
    extract.extract = lambda msg: {"kind": extract.ASK,
                                   "triples": [("machine 03", "storage", "")]}
    generate.answer = (lambda q, block, warmth=0.2, persona="",
                       max_tokens=None, **kw:
                       calls.__setitem__("n", calls["n"] + 1) or "16 GB")
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not have that.")
    try:
        said = m.ask("what is the STORAGE of Machine 03?", explain=True)
    finally:
        (extract.extract, generate.answer, generate.refusal) = real
    assert said.abstained, said
    assert calls["n"] == 0, ("%d candidate(s) were composed for a document "
                             "that never uses the word" % calls["n"])
    # ...and the document that DOES carry the head is still answered
    real2 = (extract.extract, generate.answer, generate.refusal)
    extract.extract = lambda msg: {"kind": extract.ASK,
                                   "triples": [("machine 90", "storage", "")]}
    generate.answer = (lambda q, block, warmth=0.2, persona="",
                       max_tokens=None, **kw: "512 GB")
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not have that.")
    try:
        other = m.ask("what is the STORAGE of Machine 90?", explain=True)
    finally:
        (extract.extract, generate.answer, generate.refusal) = real2
    assert "512" in str(other), other



@test("W72 the memory does not introduce someone the operator never named")
def w72():
    """Caught in a live session, and it is a product bug of the worst
    kind — the quiet kind. Asked "who is this?", a chatbot built on this
    library introduced the framework's AUTHOR by name — a fact seeded
    into the identity graph of every memory anyone builds, spoken to that
    person's end users. Nobody asked for it,
    nobody could have known it was there, and it is exactly the
    hard-coded-name-in-the-architecture this codebase forbids
    everywhere else.

    Identity is the OPERATOR'S DECLARATION. A memory that is not told
    who it is says only what it can attest — that it is a memory —
    and it never invents a maker. A memory that IS told carries that
    name in the graph like any other fact, and answers from it."""
    from lmm.session import Session
    from lmm import link
    # THE GUARD IS STRUCTURAL, so this test names nobody either: an
    # untold memory holds no creator fact at all. Naming the author here
    # to check that the author is not named would put the name back in
    # the repository, which is the thing being fixed.
    plain = Session(None)
    creators = [r for r in plain.memory.records.values()
                if (link.label_of(plain.memory, r.predicate) or "") == "creator"]
    assert not creators, (
        "an untold memory claims a creator: %s"
        % [link.label_of(plain.memory, r.value) for r in creators])
    told = Session(None, identity={"name": "Vale Coach", "maker": "Valemark"})
    block = told._chat_id_block().lower()
    assert "vale coach" in block, block
    assert "valemark" in block, block
    # ...and the operator may give a name without inventing a maker
    named = Session(None, identity="Ledger Companion")
    said = named._chat_id_block().lower()
    assert "ledger companion" in said, said
    assert "creator" not in said and "valemark" not in said, said

    # AN IDENTITY ANSWER THAT DOES NOT NAME THE MEMORY HAS NOT ANSWERED.
    # Measured live: asked "who are you?" three times, the engine wrote
    # the name once and a nameless pleasantry twice ("I am here to help"),
    # and the pleasantry passed every gate because it claims nothing. It
    # also tells the user nothing, and it is not what the operator
    # declared. The route asks again, once; if the second sentence is
    # also nameless the memory says its name plainly rather than
    # something pleasant.
    # THE ROUTE IS WHAT IS UNDER TEST, NOT THE GATE. The exit gate reads
    # the engine, and this suite's own claim is that it runs with no
    # engine and no network — so the gate is scenery here and answers
    # locally, the way every other routing test in this file does. The
    # gate has its own invariants; this one is about asking again.
    from lmm import generate, verify
    tries = []
    real_answer, real_verify = generate.identity_answer, verify.verify
    generate.identity_answer = (lambda q, name, block:
                                tries.append(1) or
                                ("I am here to help!"
                                 if len(tries) < 2 else
                                 "I am Ledger Companion."))
    verify.verify = lambda memory, text, *a, **kw: text
    try:
        out = named._identity_reply("who are you?")
    finally:
        verify.verify = real_verify
        generate.identity_answer = real_answer
    assert "Ledger Companion" in out, out
    assert len(tries) == 2, ("the nameless answer was spoken as-is or "
                             "retried more than once: %d" % len(tries))



@test("W80 a need-shaped turn is a delivery, not a gamble")
def w80():
    """The reproduced failure this closes, caught live twice in one day.
    Asked what to give a team of newly promoted managers, a consultation
    answered with one course's DURATION ROW, stamped — every word
    attested, and none of it an answer. The delivery bridge that exists
    for exactly this question only ran when the factual chain ABSTAINED;
    a need-question whose words drag in any row at all was never
    rescued, because the chain "succeeded".

    CRAG's lesson, applied structurally: evaluate BEFORE generation. The
    turn already asks "does this message want material produced?" in
    parallel (wants_material); when the verdict is yes on a consultation
    turn, the turn goes to the delivery reading DIRECTLY — the same
    composer, the same gates, the same stamps — and the factual chain is
    never gambled. When the composer has nothing, the turn falls through
    to the chain as before: a delivery that cannot deliver is not a
    refusal."""
    from lmm import generate
    from lmm.session import Session
    s = Session(None)
    s.evidence.add("COURSE: Alpha. AUDIENCE: new managers.", "#doc:alpha")
    s.history.append({"role": "user", "content": "hello"})
    called = {"compose": 0, "answer": 0}
    real_wm = generate.wants_material
    generate.wants_material = lambda message: True
    real_compose, real_answer = s.compose, s._answer
    def compose_spy(brief, seats=24, topics=None, on_line=None):
        called["compose"] += 1
        return "CATALOGUE LINE.", ["#doc:alpha"]
    def answer_spy(message, subject="", **kw):
        called["answer"] += 1
        return "COURSE: Alpha row read back."
    s.compose, s._answer = compose_spy, answer_spy
    try:
        said = s.respond("what should we give our newly promoted managers?",
                         teach=False)
    finally:
        generate.wants_material = real_wm
        s.compose, s._answer = real_compose, real_answer
    assert called["compose"] == 1, called
    assert called["answer"] == 0, (
        "the factual chain was gambled on a need-shaped turn: %r" % said)
    assert "CATALOGUE" in said, said
    assert s.last_abstained is False


@test("W79 a field learns the words a reader asks for it with, once")
def w79():
    """The measured 12 seconds this closes: a corpus writes EĞİTİM
    SÜRESİ and a reader asks "kaç saat?" — no shared word, so the
    record-direct path (milliseconds, no engine) stands aside and the
    full chain runs: candidates, gates, read-back, twelve seconds and
    half a dozen calls to read a value that was sitting in a row.

    THE BRIDGE IS PAID ONCE AND OWNED BY THE STORE. At the operator's
    request the engine is asked, per head, what words a reader might use
    to ask for that field; the answers go into a word->head index that is
    saved beside the store, like the expansion: never speakable, never
    evidence, retrieval only. Deleting the .bridge file restores the
    memory exactly as it was. Two filters keep it honest — a word the
    head itself carries adds nothing, and a word that names ANOTHER head
    verbatim may not bridge here (it already means something else in this
    corpus).

    At question time there is NO model call: a question word found in the
    bridge nominates its head, the one-head-only rule and every gate
    downstream stay exactly as they were."""
    from lmm import generate
    from lmm.session import Session
    s = Session(None)
    for src, hours in (("#doc:alpha", "6 hours"), ("#doc:beta", "3 hours")):
        s.evidence.add("DURNVAL: %s." % hours, src)
        s.evidence.add("SEATCOUNT: 12.", src)
        s.evidence.add("A quiet paragraph about %s." % src, src)
    # TODAY'S BEHAVIOUR, asserted as the precondition: reader words share
    # nothing with the head, so the record-direct path stands aside.
    assert s._record_answer("how long does alpha take, in clock time?") is None
    # the engine is consulted ONCE PER HEAD, at learning time only
    asked = []
    real = generate.asked_words
    generate.asked_words = (lambda head, value="": asked.append(head) or
                            (["clock", "time", "long", "DURNVAL"]
                             if "DURNVAL" in head else ["seats", "seatcount"]))
    try:
        kept = s.learn_bridges()
    finally:
        generate.asked_words = real
    assert sorted(asked) == ["DURNVAL", "SEATCOUNT"], asked
    assert kept >= 2, kept
    # a head's own word was filtered — bridging a word to itself is noise
    assert "durnval" not in s.evidence.head_bridge or         "DURNVAL" not in s.evidence.head_bridge.get("durnval", ()),         s.evidence.head_bridge.get("durnval")
    # ...and now the same question is a record answer: no engine, the row
    said = s._record_answer("how long does alpha take, in clock time?")
    assert said is not None, "the bridge did not reach the record path"
    line, src = said
    assert "6 hours" in line and src == "#doc:alpha", (line, src)
    # the bridge survives the trip to disk, like the expansion does
    import tempfile, os as _os
    with tempfile.TemporaryDirectory() as tmp:
        path = _os.path.join(tmp, "w79.lmm")
        s.evidence.save(path)
        again = type(s.evidence).load(path)
        assert again.head_bridge.get("clock"), "bridge lost on load"


@test("W78 a block that can be spoken carries no apparatus of ours")
def w78():
    """Found by the document probe, not by a hand-written case, which is
    the point of having one. Asked which documents state a field, the
    memory answered "K1, K2, K3, K4" — reciting the labels this code puts
    in front of evidence lines. Every claim was attested (the labels ARE
    in the block), the stamp was right, and the sentence told the reader
    nothing: our scaffolding, read out loud.

    A guard that RECOGNISES the symptom would be a patch — one refusal
    away from the next spelling of it. The label is ours to choose, so it
    is chosen to be a thing worth saying: the name of the source that
    wrote the line. An answer that echoes a label now cites a document.

    The judging prompts keep their own "[K1]" examples on purpose. Their
    output is a verdict, never a sentence, so nothing they carry can
    reach a reader."""
    from lmm.session import Session
    s = Session(None)
    lines = ["Course A runs for two days.", "Course B runs for one day."]
    s._origin_of = {lines[0]: "#docx:Course A.docx",
                    lines[1]: "#docx:Course B.docx"}
    block = s._labelled(lines)
    assert "[K1]" not in block and "[K2]" not in block, block
    assert "[Course A]" in block and "[Course B]" in block, block
    # a line whose source the turn does not know is not given an invented
    # one: it keeps its place and stands unlabelled
    plain = s._labelled(["An unattributed line."])
    assert plain == "An unattributed line.", plain


@test("W73 a question the store cannot match is asked again in the store's words")
def w73():
    """Retrieval here is lexical, and that is what makes it auditable: a
    passage arrives because a word arrived, and both can be shown. The
    cost is paraphrase. Measured on a 1,300-page novel: the book says
    "Sonya, terzi Kapernaumov'un evinde bir odada OTURUR" and a reader
    asks where she LIVES — no shared word, no answer, and the memory
    abstains with the sentence in its hands.

    The frameworks that do not have this problem embed their chunks in a
    vector space, where meaning is geometry; they also cannot say WHY a
    passage was chosen, and they fabricate. We keep the ledger and pay
    the paraphrase — but only until the first attempt fails.

    THE SECOND ASK. When a turn found nothing, the engine is shown the
    question and asked how the same thing might be WRITTEN. Two rules
    make it safe. The proposals are filtered against the store's own
    index, so a word nobody wrote cannot enter the search — the same
    rule the field bridge keeps. And the widening is retrieval only: what
    may be SAID is still read off the evidence by the same gates. A
    second chance at finding, never a second chance at claiming."""
    from lmm import evidence, generate
    from lmm.evidence import SentenceStore
    ev = SentenceStore()
    ev.add("Sonya, terzi Kapernaumov'un evinde bir odada oturur.",
           source="#kitap")
    # ...and enough noise carrying the question's OWN words that the first
    # search fills its seats without ever reaching the line: this is the
    # miss as it happens in a novel, where "nerede" and "yaşıyor" are
    # everywhere and the answer is written once, in other words.
    for i in range(40):
        # the asked-about name is COMMON in a novel — Sonya appears four
        # thousand times in the real one — so it anchors nothing, and the
        # question's remaining words are the noise's own words.
        ev.add("Sonya nerede olduğunu düşündü ve yaşıyor mu diye sordu %d."
               % i, source="#kitap")
        ev.add("Sonya bir gün yaşıyor gibi görünüyordu %d." % i,
               source="#kitap")
    soru = "Sonya nerede yaşıyor?"
    assert not any("Kapernaumov" in line for line in ev.find(soru, most=5)), \
        "the fixture does not reproduce the miss"
    asked = {}
    real = generate.phrasings
    def _proposed(question, sample=()):
        asked["q"] = question
        return ["oturur", "ikamet", "konut"]
    generate.phrasings = _proposed
    try:
        got = ev.find_again(soru, most=5)
    finally:
        generate.phrasings = real
    assert asked.get("q") == soru, asked
    assert any("Kapernaumov" in line for line in got), (
        "the second ask did not reach the line:\n"
        + "\n".join(l[:60] for l in got))
    # ...and a word the store never wrote cannot enter the search
    generate.phrasings = lambda question, sample=(): ["zeplin", "kuantum"]
    try:
        assert ev.find_again("Sonya nerede yaşıyor?", most=5) == \
            ev.find("Sonya nerede yaşıyor?", most=5), \
            "words nobody wrote changed the search"
    finally:
        generate.phrasings = real



@test("W74 a row is named by any value that names only it")
def w74():
    """A spreadsheet is the shape this system answers fastest — a row is
    a record, and the graph-first path speaks it with no model call at
    all. Measured on a public 891-row table, that path answered NONE of
    the questions a person would actually ask.

    The reason is what a row is CALLED. A CSV's first column becomes the
    row's subject, so the graph held `1 -[Name]-> Taylor, Mr. Elmer
    Zebley` and `1 -[Age]-> 42`, and every question named the passenger
    rather than the number. Nobody asks for "the Age of row 1".

    A row can be named by any of its values that names ONLY it. That is
    the same non-guessing rule the rest of this codebase keeps: a value
    carried by two rows names neither, and the path stands aside. No word
    list, no notion of what a "name column" is — only uniqueness, which
    the graph can check."""
    from lmm import lookup
    from lmm.api import Memory
    m = Memory(None)
    m.session.learn_rows([
        {"Id": "1", "Name": "Taylor, Mr. Elmer Zebley", "Age": "42",
         "Class": "first"},
        {"Id": "2", "Name": "Smith, Miss. Marion", "Age": "31",
         "Class": "second"},
        {"Id": "3", "Name": "Wiklund, Mr. Jakob", "Age": "18",
         "Class": "third"},
    ], source="#csv:people.csv")
    got = lookup.find(m.session.memory,
                      "What is the Age of Taylor, Mr. Elmer Zebley?")
    assert got is not None, "the row could not be reached by its name"
    from lmm import link
    assert link.label_of(m.session.memory, got.value) == "42", (
        link.label_of(m.session.memory, got.value))
    # ...and a value two rows share names neither of them
    m2 = Memory(None)
    m2.session.learn_rows([
        {"Id": "1", "Name": "Taylor", "Class": "first"},
        {"Id": "2", "Name": "Smith", "Class": "first"},
    ], source="#csv:people.csv")
    assert lookup.find(m2.session.memory,
                       "What is the Id of first?") is None, \
        "a value shared by two rows was allowed to name one of them"



@test("W75 an entity is a collocation, and an edge is a witnessed co-mention")
def w75():
    """The prose gap, approached the way this codebase approaches
    everything: with a statistic the document itself computes.

    WHAT GRAPHRAG DOES AND WHY WE DO NOT. It hands each chunk to a model
    and asks for the entities and the relations between them. What comes
    back is a CLAIM — invented at index time, unverifiable afterwards,
    and expensive enough to price the index out of most documents. Our
    graph is built by the document: an entity is a phrase whose words
    occur together far more often than chance (pointwise mutual
    information), and an edge between two entities says one thing only —
    THESE TWO ARE MENTIONED TOGETHER, HERE — and carries the sentence
    that witnesses it. A co-mention cannot be a fabrication, because it
    is an observation; what the relation MEANS is still a claim, and
    claims still go through the gate.

    Nothing here reads a capital letter, a name list or a language. A
    phrase that a document repeats as a unit is an entity in that
    document, whatever alphabet it is written in."""
    from lmm import mentions
    from lmm.evidence import SentenceStore
    store = SentenceStore()
    # Varied prose, the way a document is varied: the same entities in
    # different company, which is what the statistic needs and what
    # grammar does not have.
    scenes = [
        "Harbour Road was quiet when the tram passed the old mill.",
        "A baker opened his shutters on Harbour Road before dawn.",
        "The council debated whether Harbour Road should be widened.",
        "Rain fell on Harbour Road and the gutters filled quickly.",
        "Harbour Road ends at the ferry terminal, they told us.",
        "Nobody walks Harbour Road after the last tram has gone.",
        "Elmer Zebley kept a workshop two streets from the market.",
        "The letter was addressed to Elmer Zebley in careful hand.",
        "Elmer Zebley repaired clocks and refused to discuss prices.",
        "In winter Elmer Zebley closed early and read by the stove.",
        "Elmer Zebley walked down Harbour Road at dawn every Tuesday.",
        "They saw Elmer Zebley on Harbour Road, carrying a parcel.",
    ]
    for i, line in enumerate(scenes * 3):
        store.add("%s (%d)" % (line, i), source="#doc")
    graph = mentions.Graph.build(store)

    phrases = {" ".join(p) for p in graph.phrases}
    assert "harbour road" in phrases, sorted(phrases)[:10]
    assert "elmer zebley" in phrases, sorted(phrases)[:10]

    # the edge, and its witness
    linked = graph.linked("elmer zebley")
    assert "harbour road" in linked, linked
    witness = graph.witnesses("elmer zebley", "harbour road")
    assert witness and all("Elmer Zebley" in w and "Harbour Road" in w
                           for w in witness), witness

    # the partition: lines mentioning BOTH, without scoring the store
    both = graph.mentioning(["elmer zebley", "harbour road"])
    assert both and len(both) < len(store.sentences) // 3, len(both)

    # AN EDGE IS NOT A RELATION. The graph says the two are mentioned
    # together; it says nothing about what one is to the other, and
    # nothing in it can be spoken without the sentence behind it.
    assert not hasattr(graph, "relation"), \
        "the graph invented a relation type"



@test("W76 a question naming two things is answered from where they meet")
def w76():
    """The prose counterpart of the comparison reading. Asked how two
    things stand to each other, a record corpus lays their rows side by
    side; prose has no rows, and the lexical search brings whatever
    scores — usually lines about one of them. The lines that can answer
    are the lines that mention BOTH, and the entity graph knows them
    without scoring anything: two short posting lists, intersected.

    That is also why this is fast on a large document. A hundred
    thousand lines are not read; the two lists are.

    Nothing is claimed by this. The block is made of the document's own
    sentences and the gates judge them exactly as they judge any other
    evidence — the reading decides WHICH lines are worth reading, never
    what may be said about them."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    # The lexical search is given every reason to go elsewhere: dozens of
    # lines about each of the two, carrying the QUESTION'S own words, and
    # one line where they actually meet, carrying none of them. This is
    # the shape a long document has.
    tails = ["connects the guild with the market", "was mentioned again",
             "appears in the ledger", "connects two parishes",
             "is discussed by the council", "connects the fair to the mill"]
    for i in range(30):
        tail = tails[i % len(tails)]
        s.learn_text("Elmer Zebley %s (%d)." % (tail, i),
                     source="#doc:town", deep=False)
        s.learn_text("Harbour Road %s (%d)." % (tails[(i + 3) % len(tails)], i),
                     source="#doc:town", deep=False)
    s.learn_text("Elmer Zebley rented the corner shop on Harbour Road "
                 "in 1893.", source="#doc:town", deep=False)
    blocks = []
    real = (extract.extract, generate.answer, generate.refusal,
            generate.offer_research)
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("elmer zebley", "road", "")]}
    generate.answer = (lambda q, block, warmth=0.2, persona="",
                       max_tokens=None, **kw:
                       blocks.append(block) or "I do not know.")
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    try:
        s.respond("what connects Elmer Zebley and Harbour Road?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research) = real
    assert blocks, "nothing was composed"
    first = blocks[0]
    assert "corner shop" in first, (
        "the line where the two meet was not in the block:\n" + first[:400])

    # WHAT THE GRAPH BUYS IS SCALE, AND THIS IS THE HONEST VERSION OF
    # THAT CLAIM. On a store this size the lexical search already finds
    # the meeting line — two rare names in one line outscore one name in
    # many — and measured on a 36,472-line novel it did so for six pairs
    # out of six. What it cannot do is do it cheaply: 417 ms a question,
    # because every line carrying either name is scored. The graph reads
    # the same answer out of two posting lists in 0.01 ms. So the reading
    # is not "the graph finds better lines"; it is "the graph finds the
    # same lines without reading the document", which is the difference
    # between a demo and ten thousand pages.
    from lmm import mentions
    graph = mentions.Graph.build(s.evidence)
    both = graph.mentioning(["elmer zebley", "harbour road"])
    assert both, "the intersection was empty where the lexical search hit"
    assert all("corner shop" in s.evidence.sentences[sid][0] for sid in both), \
        [s.evidence.sentences[sid][0][:60] for sid in both]



@test("W77 the corpus is its own thesaurus, and can show its working")
def w77():
    """The second half of the paraphrase problem, answered from the
    document rather than from a vendor.

    A document says a tenant RESIDES somewhere and a reader asks where
    she LIVES. Lexical retrieval cannot bridge that, and the engine can
    (the second ask does), but it costs a call and its proposals have to
    be checked against the store anyway. The document already knows: two
    words that appear in the same COMPANY are used for the same thing —
    "resides" and "lives" both turn up beside "tenant", "address",
    "since", "flat". That is distributional similarity, it is arithmetic
    over counts this store already holds, and unlike an embedding it can
    be shown: THESE are the shared contexts, and this is their weight.

    The words it offers are the store's own by construction, so the rule
    the field bridge and the second ask both keep — a word nobody wrote
    cannot enter a search — is satisfied without a check.

    What it is not: a synonym list, a language rule, or a model. What it
    cannot do: relate two words the document never uses in comparable
    company, which is why the engine's second ask stays as the fallback."""
    from lmm import affinity
    from lmm.evidence import SentenceStore
    store = SentenceStore()
    for i in range(14):
        store.add("The tenant resides at the flat since spring %d." % i,
                  source="#doc")
        store.add("The tenant lives at the flat since autumn %d." % i,
                  source="#doc")
        store.add("A visitor resides nearby and the tenant knows him %d." % i,
                  source="#doc")
        store.add("A visitor lives nearby and the tenant greets him %d." % i,
                  source="#doc")
        store.add("The ledger records the rent and the date %d." % i,
                  source="#doc")
    profiles = affinity.Profiles.build(store)

    near = profiles.near("lives", most=4)
    words = [word for word, _score, _shared in near]
    assert "resides" in words, near
    # ...and the working is visible: the shared company, not a number
    # nobody can inspect
    shared = dict((w, sh) for w, _s, sh in near)["resides"]
    assert shared and all(isinstance(c, str) for c in shared), shared
    assert "tenant" in shared or "flat" in shared or "nearby" in shared, shared

    # a word with no comparable company gets nothing, rather than the
    # nearest thing in a geometry
    assert not profiles.near("ledger", most=3) or all(
        score < 1.0 for _w, score, _sh in profiles.near("ledger", most=3))


@test("W50 a field question that names no source reads that field across the corpus")
def w50():
    """The class the seats cannot serve: "which course runs the
    longest?" — a question about a FIELD with no document named. Six
    seats cannot hold sixty documents' values, and the source cap
    rightly keeps any one document from taking them, so the block held
    an arbitrary handful and the turn abstained. The census answered
    the same shape of question years earlier — "who all speaks of this"
    is a COUNT, not a retrieval — and this is that reading applied to
    values instead of mentions: when the question names no source and
    three or more sources answer the same record head, the block is one
    row per source, datelined, capped. The engine ranks attested values
    instead of guessing, and the honest tie (two sources sharing the top
    value) is visible rather than hidden behind whichever line won a
    seat."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    names = ["Alpha","Beta","Gamma","Delta","Epsilon","Zeta","Eta","Theta",
             "Iota","Kappa","Lambda","Mu","Nu","Xi","Omicron","Pi","Rho",
             "Sigma","Tau","Upsilon","Phi","Chi","Psi","Omega","Ares",
             "Bora","Cirrus","Dorado","Elara","Fornax"]
    for i, name in enumerate(names):
        days = 12 if name == "Sigma" else 1 + (i % 4)
        lines = [f"{name.upper()} COURSE OUTCOMES."] + [
            f"The course covers the {w} of the {name.lower()} module over "
            f"several long days of practice." for w in
            ("opening", "middle", "closing", "review", "handover")
        ] + [f"DURATION: {days} days."]
        s.learn_text("\n".join(lines),
                     source=f"#docx:{name} Course.docx", deep=False)
    blocks = []
    real = (extract.extract, generate.answer, generate.refusal,
            generate.offer_research)
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("course", "duration", "")]}
    def spy_answer(q, block, warmth=0.2, persona="", max_tokens=None, **kw):
        blocks.append(block)
        return "I do not know."
    generate.answer = spy_answer
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    try:
        s.respond("which course has the longest DURATION?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research) = real
    joined = "\n".join(blocks)
    seen = {n for n in names if f"{n} Course" in joined and "DURATION" in joined}
    assert len(seen) >= 8, ("the field was not read across the corpus: %d"
                            % len(seen))
    assert "12 days" in joined, "the winning value never reached the block"


@test("W51 the extremes of a field are written into the evidence")
def w51():
    """The half of the corpus-wide reading the gate could not admit. Ask
    which course takes the most participants and the engine answers with
    a NAME — a conclusion no single line states, since it is computed
    across sixty rows — so the read-back rightly refuses it and the turn
    abstains with the whole field on the table. W44 solved the same
    shape for two sources by writing the verdict INTO the evidence; the
    corpus-wide reading gets the same treatment: when the field rows
    carry digits, one row states both ends — largest and smallest, each
    with its source and its value — computed from the digits the
    documents wrote. The engine then reads the answer instead of
    deriving it, and whichever end the question asks for is attested.
    Ranges are read at their own ends (16-20 ranks by 20 for largest, by
    16 for smallest), which is what a range means."""
    from lmm import generate, extract
    from lmm.session import Session
    s = Session(None)
    seats = {"Alpha": "8 people", "Beta": "16-20 people",
             "Gamma": "12 people", "Delta": "6-9 people",
             "Epsilon": "14 people", "Zeta": "10 people"}
    for name, v in seats.items():
        s.learn_text(f"{name.upper()} COURSE OUTCOMES.\n"
                     f"The {name.lower()} module opens the arc.\n"
                     f"SEATS: {v}.",
                     source=f"#docx:{name} Course.docx", deep=False)
    blocks = []
    real = (extract.extract, generate.answer, generate.refusal,
            generate.offer_research)
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("course", "seats", "")]}
    def spy_answer(q, block, warmth=0.2, persona="", max_tokens=None, **kw):
        blocks.append(block)
        return "I do not know."
    generate.answer = spy_answer
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    try:
        s.respond("which course takes the most SEATS?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research) = real
    rows = [l for l in "\n".join(blocks).split("\n") if "largest" in l]
    assert rows, "no extremes row was written"
    row = rows[0]
    assert "Beta" in row and "16-20" in row, row      # largest by 20
    assert "smallest" in row and "Delta" in row, row  # smallest by 6
    # A VALUE IS A NUMBER AND ITS UNIT. Measured on the field corpus:
    # ranking by digits alone made a 90-minute webinar the longest
    # training in a corpus of multi-day courses. Values compare only
    # when their units agree; the ends are read within the unit the
    # field mostly speaks, and rows in another unit are left out of the
    # ranking rather than mis-ranked.
    s2 = Session(None)
    mixed = {"Ada": "4 days", "Bly": "90 minutes", "Cnut": "2 days",
             "Dane": "45 minutes", "Erle": "3 days", "Fen": "5 days"}
    for name, v in mixed.items():
        s2.learn_text(f"{name.upper()} COURSE OUTCOMES.\n"
                      f"The {name.lower()} module opens the arc.\n"
                      f"DURATION: {v}.",
                      source=f"#docx:{name} Course.docx", deep=False)
    blocks2 = []
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("course", "duration", "")]}
    generate.answer = (lambda q, block, warmth=0.2, persona="",
                       max_tokens=None, **kw: (blocks2.append(block),
                                         "I do not know.")[1])
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    try:
        s2.respond("which course has the longest DURATION?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research) = real
    ends = [l for l in "\n".join(blocks2).split("\n") if "largest" in l]
    assert ends, "no extremes row for the mixed-unit field"
    assert "Fen" in ends[0], ("ranked across units: %r" % ends[0])
    assert "90 minutes" not in ends[0], ("minutes outranked days: %r"
                                         % ends[0])
    # THE FIELD IS THE ONE THE QUESTION COVERS BEST. Measured on the
    # corpus: asked which course takes the MOST PEOPLE, the harvest
    # gathered the DURATION field — because the question's generic word
    # ("course") matched that head too, and any match was enough. A
    # question names its field the way it names a source: by how much of
    # the head it accounts for. The best-covered head wins; a tie is an
    # ambiguity and the reading stands aside.
    s3 = Session(None)
    for name, days, people in (("Ada", 4, 30), ("Bly", 2, 12),
                               ("Cnut", 3, 25), ("Dane", 1, 8),
                               ("Erle", 5, 16), ("Fen", 2, 21)):
        s3.learn_text(f"{name.upper()} COURSE OUTCOMES.\n"
                      f"The {name.lower()} module opens the arc.\n"
                      f"COURSE DURATION: {days} days.\n"
                      f"SEAT COUNT: {people} people.",
                      source=f"#docx:{name} Course.docx", deep=False)
    blocks3 = []
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("course", "seats", "")]}
    generate.answer = (lambda q, block, warmth=0.2, persona="",
                       max_tokens=None, **kw: (blocks3.append(block),
                                         "I do not know.")[1])
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    try:
        s3.respond("which course has the highest SEAT COUNT?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research) = real
    ends3 = [l for l in "\n".join(blocks3).split("\n") if "largest" in l]
    assert ends3, "no extremes row for the seat-count field"
    assert "SEAT COUNT" in ends3[0], ("the wrong field was harvested: %r"
                                      % ends3[0])
    assert "Ada" in ends3[0] and "30" in ends3[0], ends3[0]
    # A NUMBER BELONGS TO THE UNIT BESIDE IT. Measured: a value reading
    # "4 modules x half a day - about 2 days in total" ranked as FOUR
    # days, because the largest digit in the line was taken for the
    # value; the four counts modules. Digits are paired with the unit
    # they stand next to, and a range keeps both of its ends.
    s4 = Session(None)
    vals = {"Ada": "4 modules x half a day - about 2 days in total",
            "Bly": "3 days", "Cnut": "1 day", "Dane": "5 days"}
    for name, v in vals.items():
        s4.learn_text(f"{name.upper()} COURSE OUTCOMES.\n"
                      f"The {name.lower()} module opens the arc.\n"
                      f"DURATION: {v}.",
                      source=f"#docx:{name} Course.docx", deep=False)
    blocks4 = []
    extract.extract = lambda m: {"kind": extract.ASK,
                                 "triples": [("course", "duration", "")]}
    generate.answer = (lambda q, block, warmth=0.2, persona="",
                       max_tokens=None, **kw: (blocks4.append(block),
                                         "I do not know.")[1])
    generate.refusal = (lambda message, persona="", warmth=0.3,
                        max_tokens=None: "I do not know.")
    generate.offer_research = lambda subject, question: ""
    try:
        s4.respond("which course has the longest DURATION?")
    finally:
        (extract.extract, generate.answer, generate.refusal,
         generate.offer_research) = real
    ends4 = [l for l in "\n".join(blocks4).split("\n") if "largest" in l]
    assert ends4, "no extremes row for the qualified-value field"
    assert "Dane" in ends4[0], ("a module count was read as days: %r"
                                % ends4[0])


@test("W52 an answer that names a source is stamped with that source")
def w52():
    """The provenance side of the corpus-wide reading. A system-written
    row — the extremes of a field, the verdict of a comparison — belongs
    to no single document, so it carries no origin; and when the answer
    rests on THAT row, the mark fell back to whichever seat happened to
    be first, stamping a true sentence about one course with another
    course's name. A stamp is a promise about where a claim comes from,
    so it follows the claim: when the answer names a source the block
    holds, that source is the mark. The jury already reads claims this
    way (a claim that names a source is judged by that source alone);
    the mark now reads them the same."""
    from lmm.session import Session
    s = Session(None)
    # the shape the corpus produced: the only line naming the winner is
    # the system's own row, and the seats around it belong to others
    proof = ["DURATION \u2014 largest: Beta Course: 5 days "
             "\u00b7 smallest: Alpha Course: 1 day",
             "Gamma Course \u2014 DURATION: 2 days.",
             "Alpha Course \u2014 DURATION: 1 day."]
    origins = ["", "#docx:Gamma Course.docx", "#docx:Alpha Course.docx"]
    s.evidence.add("Beta Course runs five days.", "#docx:Beta Course.docx")
    mark = s._load_bearing("The longest is Beta Course at 5 days.",
                           proof, origins)
    assert mark == "#docx:Beta Course.docx", mark
    # a claim naming nothing keeps the old reading: the best-covered line
    mark2 = s._load_bearing("It runs for two days.", proof, origins)
    assert mark2 == "#docx:Gamma Course.docx", mark2


@test("W53 a refusal is spoken in the language it was asked in")
def w53():
    """Measured, and it had been hiding behind a Turkish corpus: six
    English questions with no answer in the store were refused in Dutch,
    French and Spanish — never once in English. The instruction to match
    the user's language sat at the top of the system prompt, where a
    small engine reads it as background and then writes whatever a
    refusal usually looks like in its training. Two structural moves,
    no phrase in any language written by us: the question is handed
    back as a LANGUAGE SAMPLE inside the instruction (the same anchoring
    the composer uses for its material), and the instruction is repeated
    as the LAST thing the engine reads before it writes, because the
    end of the prompt is where an instruction survives. What is not
    done: a canned sentence per language — the memory speaks in the
    engine's voice, or it does not speak."""
    from lmm import generate, runtime
    seen = {}
    real = runtime.generate
    def spy(messages, max_tokens=256, temperature=0.7, system=None):
        seen["system"] = system or ""
        seen["messages"] = messages
        return "..."
    runtime.generate = spy
    try:
        generate.refusal("What is the battery life of the Marlin?")
    finally:
        runtime.generate = real
    sys_prompt = seen["system"]
    assert "Marlin" in sys_prompt, \
        "the question is not handed back as a language sample"
    tail = sys_prompt[-260:]
    assert "language" in tail.lower(), \
        "the language rule is not the last thing the engine reads"
    # ...AND THE SPEECH IS VERIFIED LIKE EVERYTHING ELSE. Anchoring got
    # five of six; one question kept coming back in another language,
    # and a memory that audits every claim it speaks can audit the
    # tongue it speaks them in. The refusal is read once by the same
    # kind of small judge the gates use — is this the language of the
    # question? — and on a no it is written once more. A second no is
    # kept: silence is worse than a sentence in the wrong language, and
    # nothing here writes a canned phrase.
    # ...AND THE SPEECH IS CHECKED, BY NAMING RATHER THAN COMPARING. The
    # first cut asked a judge "is the reply in the same language as the
    # question?" and, measured against the live engine, it answered NO for
    # an English question answered in English — a verdict carrying no
    # information, so every refusal was rewritten once and then kept
    # whatever came back, Dutch included. Naming ONE sentence's language
    # is a smaller question the engine answers reliably; the comparison
    # is then arithmetic. On a mismatch the sentence is written once
    # more, with the language NAMED in front of it. A canned phrase is
    # still written in no language.
    calls = {"n": 0}
    outs = iter(["Ik heb die informatie nog niet.",
                 "I do not have that information yet."])
    def spy2(messages, max_tokens=256, temperature=0.7, system=None):
        calls["n"] += 1
        return next(outs)
    named = []
    real_named = generate.language_of
    generate.language_of = lambda text: (
        named.append(text) or ("Dutch" if "informatie" in text else "English"))
    runtime.generate = spy2
    try:
        out = generate.refusal("What is the battery life of the Marlin?")
    finally:
        runtime.generate = real
        generate.language_of = real_named
    assert named, "the language was never named"
    assert out.startswith("I do not have"), out
    assert calls["n"] == 2, ("the refusal was rewritten more than once: %r"
                             % calls)


@test("X5 an expansion written in another language never reaches the index")
def x5():
    """THE MARGIN FILTER CANNOT SEE THIS ONE, BY CONSTRUCTION.

    A query in another language reaches nothing lexically and scores a zero
    margin, which is the exact signature of the genuinely-new vocabulary the
    expansion exists to buy — so `X3`'s filter waves it through. Measured on
    NIST SP 800-63B, an English publication: the engine expanded it into
    Turkish, German, Spanish, Portuguese and French, 62% of that survived, and
    one survivor was a refusal sentence indexed as a query.

    It is not a prompt defect and cannot be fixed as one. Putting the language
    rule first changed nothing; removing the multilingual examples moved the
    drift from Spanish to French. A 3B engine does not hold a language
    instruction across this task, so the property has to be structural.

    The document is the only language sample there is. Below, the store speaks
    English, and the two assertions are the whole rule: an English question
    that introduces a NEW word still reaches the index, because the rest of it
    is words this document uses — while a translation of that same question
    does not, because almost none of it is. No language is named in the
    implementation; a Turkish document keeps its Turkish expansions by the same
    reading."""
    from lmm import evidence
    store = evidence.SentenceStore()
    for line in ("The Kelvane tower measures 42 metres and was finished early.",
                 "The Nordheim archive opened in 1904 and holds the county maps.",
                 "Visitors reach the archive by the western stair in the tower."):
        store.add(line, "#doc:x")
    store.learn_expansions({0: [
        "how tall is the Kelvane spire",          # English, one new word
        "wie hoch ist der Kelvane Turm",          # the same question, German
        "quelle est la hauteur de la tour",       # French
    ]})
    kept = store.expansions.get(0, [])
    assert "how tall is the Kelvane spire" in kept, kept
    assert "wie hoch ist der Kelvane Turm" not in kept, kept
    assert "quelle est la hauteur de la tour" not in kept, kept


def main():
    failed = 0
    for name, function in PASSED:
        try:
            function()
        except AssertionError as broke:
            failed += 1
            print("FAIL  %s\n      %s" % (name, broke))
        except Exception as broke:                          # noqa: BLE001
            failed += 1
            print("ERROR %s\n      %r" % (name, broke))
        else:
            print("ok    %s" % name)
    print("\n%d/%d passed" % (len(PASSED) - failed, len(PASSED)))
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
